from datetime import date
from pathlib import Path

import pytest
import xlsxwriter
from openpyxl import Workbook
from engines.onboarding.onboarding_common import load_profile
from engines.onboarding.onboarding_input import validate_headers, _xlsx_headers, _mixed_numeric_date_overrides
from engines.onboarding.utils import ReportError


def test_duplicate_headers_are_not_silently_renamed():
    with pytest.raises(ReportError,match='duplicate headers'):
        validate_headers(['Name',' NAME ','ID'])


def test_whitespace_case_header_normalization():
    assert validate_headers([' Registered   Facility ','ID'])=={'registered facility':0,'id':1}


@pytest.mark.parametrize('writer',['openpyxl','xlsxwriter'])
def test_header_metadata_shared_and_inline_strings(tmp_path,writer):
    profile=load_profile();path=tmp_path/'headers.xlsx'
    if writer=='openpyxl':
        w=Workbook();s=w.active;s.title='Export';s.append(profile['headers']);w.save(path)
    else:
        with xlsxwriter.Workbook(path) as w:
            w.add_worksheet('Export').write_row(0,0,profile['headers'])
    result=_xlsx_headers(path)
    assert result[0]['name']=='Export'
    assert result[0]['headers']==profile['headers']


def test_multiple_worksheets_empty_first_sheet(tmp_path):
    p=tmp_path/'sheets.xlsx';w=Workbook();w.active.title='Empty';w.create_sheet('Export').append(['ID','Name']);w.save(p)
    result=_xlsx_headers(p)
    assert [s['name'] for s in result]==['Empty','Export']
    assert result[0]['headers']==[]


def test_native_serial_vs_literal_numeric_text(tmp_path):
    profile=load_profile();p=tmp_path/'dates.xlsx';w=Workbook();s=w.active;s.append(profile['headers'])
    col=profile['headers'].index('Onboarded Date')+1
    s.cell(2,col).value=61
    s.cell(3,col).value='61'
    w.save(p)
    info=_xlsx_headers(p)[0]
    result=_mixed_numeric_date_overrides(p,info,['61','61'])
    assert result=={0:date(1900,3,1)}


@pytest.mark.parametrize("change", ["reordered", "missing"])
def test_layout_preserves_fixed_formula_column_positions(tmp_path, change):
    import json
    from engines.onboarding.onboarding_common import load_profile
    from engines.onboarding.utils import ReportError
    profile = load_profile()
    if change == "reordered":
        profile["headers"][7], profile["headers"][8] = profile["headers"][8], profile["headers"][7]
    else:
        profile["headers"] = profile["headers"][:10]
    path = tmp_path / "onboarding_layout.json"
    path.write_text(json.dumps(profile), encoding="utf-8")
    with pytest.raises(ReportError, match="preserved formulas"):
        load_profile(path)
