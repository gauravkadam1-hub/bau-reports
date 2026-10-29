from pathlib import Path
import csv

import pytest
from openpyxl import load_workbook

from portal.file_detector import classify_filename, detect_files
from tests.attendance_target_fixtures import make_inputs, configure_engine, DC, XD


def _floater_csv(path: Path, ids=('test-0001',)):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w=csv.writer(f); w.writerow(['Valinor ID','NAME','BRANCH CODE'])
        for ident in ids:
            w.writerow([ident,'Synthetic Floater','Test Alpha DC'])
    return path


@pytest.fixture
def attendance(monkeypatch):
    from engines.attendance import engine
    monkeypatch.setattr(engine, 'pl', None)
    return engine


@pytest.mark.parametrize('name', ['Floater.csv','Floaters.xlsx','floater base.xlsm','FLOATERS-list.xls'])
def test_optional_floater_filename_detection(name):
    assert classify_filename(name) == 'floater'


def test_floater_is_optional_for_attendance_readiness():
    names=['all_user_data_all_a.csv','all_user_data_all_b.csv','attendance_report_x.csv']
    d=detect_files(names)
    assert d.attendance_ready and d.floater == []
    d2=detect_files(names+['Floater Base.csv'])
    assert d2.attendance_ready and d2.floater == ['Floater Base.csv']


def test_formula_contract_with_and_without_uploaded_floater(attendance):
    e=attendance
    with_file=e._all_user_floater_formula('C3',True,True)
    assert with_file == '=IF(COUNTIF(Floater!$A:$A,C3)>0,TRUE,INDEX($1:$1048576,ROW(),MATCH("Floater",$1:$1,0))="Yes")'
    no_file=e._all_user_floater_formula('C3',True,False)
    assert no_file == '=INDEX($1:$1048576,ROW(),MATCH("Floater",$1:$1,0))="Yes"'
    att_with=e._attendance_floater_formula('C3','C',True,True)
    assert att_with == '=IF(COUNTIF(Floater!$A:$A,C3)>0,TRUE,IFERROR(INDEX(\'All User data\'!$1:$1048576,MATCH(C3,\'All User data\'!$C:$C,0),MATCH("Floater",\'All User data\'!$1:$1,0))="Yes",FALSE()))'
    att_without=e._attendance_floater_formula('C3','C',True,False)
    assert att_without == '=IFERROR(INDEX(\'All User data\'!$1:$1048576,MATCH(C3,\'All User data\'!$C:$C,0),MATCH("Floater",\'All User data\'!$1:$1,0))="Yes",FALSE())'


@pytest.mark.parametrize(
    ('raw','listed','expected'),
    [
        ('Yes', False, True),
        ('No', False, False),
        ('-', False, False),
        ('', False, False),
        ('No', True, True),
        ('-', True, True),
    ],
)
def test_main_floater_priority_is_uploaded_list_then_all_user_yes(attendance,raw,listed,expected):
    ids={'test-0001'} if listed else set()
    assert attendance._floater_from_user(raw,'test-0001',ids) is expected


def _generate(e, paths, out, reports, floater=None):
    return e.generate_reports(e.ReportInputs(
        str(paths['users1']), str(paths['users2']), str(paths['attendance']), str(out),
        paths['dates'][0], paths['dates'][-1], DC[:1], XD[:1], reports,
        floater_path=str(floater) if floater else None,
    ))


def test_main_without_floater_has_no_sheet_and_no_sheet_references(attendance,tmp_path):
    e=attendance; paths=make_inputs(tmp_path/'inputs',date_count=2); configure_engine(e,paths)
    output,=_generate(e,paths,tmp_path/'out',[e.MAIN_REPORT_TOKEN])
    wb=load_workbook(output,read_only=True,data_only=False)
    try:
        assert 'Floater' not in wb.sheetnames
        au=wb['All User data']; ab=wb['Attendance Base']
        assert au['A1'].value=='To be considered'
        assert au['B1'].value=='To be considered(Values - TRUE/FALSE)'
        assert 'Floater!$A:$A' not in au['B2'].value
        assert 'Floater!$A:$A' not in ab['B3'].value
        assert 'MATCH("Floater"' in au['B2'].value and 'MATCH("Floater"' in ab['B3'].value
    finally:
        wb.close()


def test_main_with_floater_copies_sheet_and_uses_supplied_formulas(attendance,tmp_path):
    e=attendance; paths=make_inputs(tmp_path/'inputs',date_count=2); configure_engine(e,paths)
    floater=_floater_csv(tmp_path/'Floater Base.csv')
    output,=_generate(e,paths,tmp_path/'out',[e.MAIN_REPORT_TOKEN],floater)
    wb=load_workbook(output,read_only=True,data_only=False)
    try:
        assert 'Floater' in wb.sheetnames
        assert wb['Floater']['A1'].value=='Valinor ID' and wb['Floater']['A2'].value=='test-0001'
        assert 'COUNTIF(Floater!$A:$A,C2)>0' in wb['All User data']['B2'].value
        assert 'COUNTIF(Floater!$A:$A,C3)>0' in wb['Attendance Base']['B3'].value
    finally:
        wb.close()


def test_housekeeping_security_keep_only_to_be_considered_helper_and_ignore_floater(attendance,tmp_path):
    e=attendance; paths=make_inputs(tmp_path/'inputs',date_count=2); configure_engine(e,paths)
    floater=_floater_csv(tmp_path/'Floaters.csv',('test-0006','test-0011'))
    outputs=_generate(e,paths,tmp_path/'out',[e.MAIN_REPORT_TOKEN,'Housekeeping','Security'],floater)
    assert len(outputs)==3
    for output in outputs:
        wb=load_workbook(output,read_only=True,data_only=False)
        try:
            if 'Housekeeping' not in output.name and 'Security' not in output.name:
                continue
            assert 'Floater' not in wb.sheetnames
            for sheet,row in [('All User data',1),('Attendance Base',1)]:
                ws=wb[sheet]
                headers=[ws.cell(row,c).value for c in range(1,ws.max_column+1)]
                assert headers[0]=='To be considered'
                assert 'To be considered(Values - TRUE/FALSE)' not in headers
            summary=wb['Attendance Adherence - DC']
            formulas=[summary.cell(4,c).value for c in (2,3,5)]
            assert all('To be considered(Values - TRUE/FALSE)' not in (f or '') for f in formulas)
            assert all('Floater!' not in (f or '') for f in formulas)
        finally:
            wb.close()


def test_dedicated_cached_metrics_do_not_change_when_floater_list_contains_ground_workers(attendance,tmp_path):
    e=attendance; paths=make_inputs(tmp_path/'inputs',date_count=2); configure_engine(e,paths)
    without=_generate(e,paths,tmp_path/'without',['Housekeeping','Security'])
    # If only dedicated reports are selected, even an invalid/irrelevant optional Floater file must be ignored.
    bogus=tmp_path/'Floater Base.csv'; bogus.write_text('not even needed\n',encoding='utf-8')
    with_file=_generate(e,paths,tmp_path/'with',['Housekeeping','Security'],bogus)
    for a,b in zip(without,with_file):
        wa=load_workbook(a,read_only=True,data_only=True); wb=load_workbook(b,read_only=True,data_only=True)
        try:
            for sheet in ['Attendance Adherence - DC','Attendance Adherence - XD']:
                va=[[cell.value for cell in row] for row in wa[sheet].iter_rows()]
                vb=[[cell.value for cell in row] for row in wb[sheet].iter_rows()]
                assert va==vb
        finally:
            wa.close(); wb.close()

def test_uploaded_floater_excel_is_read_and_preserved(attendance,tmp_path):
    from openpyxl import Workbook
    e=attendance
    path=tmp_path/'Floaters.xlsx'
    wb=Workbook(); ws=wb.active; ws.title='Floaters'
    ws.append(['Valinor ID','NAME','BRANCH CODE'])
    ws.append(['SFXTEST1','Test Rider','Test Alpha DC'])
    wb.save(path)
    rows,ids=e._read_floater_input(str(path))
    assert rows[0]==['Valinor ID','NAME','BRANCH CODE']
    assert rows[1]==['SFXTEST1','Test Rider','Test Alpha DC']
    assert 'sfxtest1' in ids
