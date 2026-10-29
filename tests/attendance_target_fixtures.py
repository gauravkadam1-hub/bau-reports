"""Small, entirely synthetic inputs/masters for Attendance target-row tests."""
from datetime import date, timedelta
from pathlib import Path
import csv
from openpyxl import Workbook

DC = ['Test Alpha DC', 'Test Bravo DC']
XD = ['Test Delta XD', 'Test Echo XD']


def write_csv(path, headers, rows, second_header=None):
    with Path(path).open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        if second_header is not None:
            writer.writerow(second_header)
        writer.writerows(rows)


def make_inputs(folder, date_count=5):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    dates = [date(2026, 10, 1) + timedelta(days=i) for i in range(date_count)]
    users, attendance = [], []
    for si, site in enumerate(DC + XD + ['Excluded HUB']):
        for role in ['GC', 'Housekeeping', 'Security', 'Maintenance']:
            for k, punch in enumerate(['SP', 'MIS', 'HF', 'P', 'P']):
                ident = f'test-{len(users)+1:04}'
                users.append([ident, 'Synthetic Employee', role, 'No',
                              'Onboarded - Active', site, '2026-09-30'])
                row = [ident, 'Synthetic Employee', role, site, 'test-stint', '']
                for di in range(date_count):
                    code = ['SP', 'MIS', 'HF', 'P', 'P'][(k+di+si)%5]
                    hours = {'SP':'0', 'MIS':'4', 'HF':'6', 'P':'9'}[code]
                    row += ['P1', hours, hours, code]
                attendance.append(row)
    user_headers = ['ID', 'Name', 'Role', 'Floater', 'Status', 'Registered Facility', 'Onboarded Date']
    paths = {name:folder/f'{name}.csv' for name in ['users1','users2','attendance']}
    half = len(users)//2
    write_csv(paths['users1'],user_headers,users[:half])
    write_csv(paths['users2'],user_headers,users[half:])
    top = ['user_id','user_name','role','facility_name','stint_id','last_present_day']
    sub = ['']*6
    for day in dates:
        top += [day.isoformat(),'','','']
        sub += ['shift','total_hours','working_hours','punch_status']
    write_csv(paths['attendance'],top,attendance,sub)
    master = Workbook(); master.remove(master.active)
    for role in ['Housekeeping','Security','Maintenance']:
        ws = master.create_sheet(f'{role} - Deployed HC')
        ws.append(['Site','Deployed HC' if role=='Housekeeping' else 'Required HC'])
        for si,site in enumerate(DC+XD):
            # Omit one facility to exercise the existing user-count fallback.
            if si!=3:
                ws.append([site,8+si])
    paths['master'] = folder/'synthetic_deployed.xlsx';master.save(paths['master'])
    ref = Workbook();ref.active.title='Ref';ref.active.append(['GC'])
    paths['ref'] = folder/'synthetic_ref.xlsx';ref.save(paths['ref'])
    assets = folder/'assets';assets.mkdir(exist_ok=True)
    write_csv(assets/'floater.csv',['User ID'],[])
    paths['assets']=assets;paths['dates']=dates
    return paths


def configure_engine(engine, paths):
    """Call only on an isolated module, or under pytest monkeypatch."""
    engine.MASTER_HC_PATH=paths['master']
    engine.REF_MASTER_PATH=paths['ref']
    engine.ASSET_DIR=paths['assets']
