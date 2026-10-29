"""Plain Base, positive-only RED scales, and lossless compact XLSX regressions."""
from __future__ import annotations

from datetime import datetime, time
from hashlib import sha256
from pathlib import Path
import xml.etree.ElementTree as ET
from zipfile import ZipFile, ZIP_STORED

import pytest
from openpyxl import load_workbook

from engines.onboarding import onboarding_writer as writer
from engines.onboarding.onboarding_common import load_profile, Counts
from engines.onboarding.onboarding_packaging import compact_workbook
from engines.onboarding.utils import ReportError
from tests.onboarding_fixtures import rows_for, reference_prepare, END

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def generate(tmp_path, *, zero=False, dc=None, xd=None):
    profile = load_profile()
    records = rows_for(profile)
    if zero:
        for row in records:
            row[20] = 'green'
    kwargs = {}
    if dc is not None: kwargs['dc'] = dc
    if xd is not None: kwargs['xd'] = xd
    report = reference_prepare(profile, records, **kwargs)
    path = tmp_path / 'compact.xlsx'
    writer.write_report(path, profile, report)
    return path, report


def red_scales(sheet):
    return [(key, rule) for key, rules in sheet.conditional_formatting._cf_rules.items()
            for rule in rules if rule.type == 'colorScale' and len(rule.colorScale.color) == 2]


@pytest.mark.parametrize('name', ['Aadhar - DC', 'Aadhar - XD', 'Hyperverge - DC', 'Hyperverge - XD'])
def test_all_zero_columns_have_no_red_scale(tmp_path, name):
    path, _ = generate(tmp_path, zero=True)
    workbook = load_workbook(path)
    assert not red_scales(workbook[name])
    if name.startswith('Hyperverge'):
        assert not workbook[name].conditional_formatting
    workbook.close()


@pytest.mark.parametrize('name,positive_columns', [
    ('Aadhar - DC', [19]), ('Aadhar - XD', [19]),
    ('Hyperverge - DC', [3, 9]), ('Hyperverge - XD', [3, 5, 9]),
])
def test_scale_columns_exactly_match_positive_counts(tmp_path, name, positive_columns):
    path, _ = generate(tmp_path)
    workbook = load_workbook(path)
    scales = red_scales(workbook[name])
    assert sorted(next(iter(key.sqref.ranges)).min_col for key, _ in scales) == positive_columns
    workbook.close()


@pytest.mark.parametrize('name', ['Aadhar - DC', 'Aadhar - XD', 'Hyperverge - DC', 'Hyperverge - XD'])
def test_scale_zero_is_white_positive_ceiling_and_no_total(tmp_path, name):
    path, _ = generate(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook[name]
    for key, scale in red_scales(sheet):
        area = next(iter(key.sqref.ranges))
        assert area.min_row == 3
        assert area.max_row == sheet.max_row - 1
        assert area.min_col == area.max_col
        assert [c.rgb for c in scale.colorScale.color] == ['FFFFFFFF', 'FFF8696B']
        assert scale.colorScale.cfvo[0].type == 'num'
        assert scale.colorScale.cfvo[0].val == 0
        ceiling = scale.colorScale.cfvo[1]
        assert ceiling.type == 'num' and ceiling.val > 0
        cached = load_workbook(path, data_only=True)
        values = [cached[name].cell(row, area.min_col).value
                  for row in range(area.min_row, area.max_row+1)]
        assert ceiling.val == max(values)
        cached.close()
    workbook.close()


def test_single_positive_site_has_scale_and_total_is_excluded(tmp_path):
    path, _ = generate(tmp_path, dc=['Beta DC'], xd=[])
    workbook = load_workbook(path)
    for name in ['Aadhar - DC', 'Hyperverge - DC']:
        assert red_scales(workbook[name])
        for key, _ in red_scales(workbook[name]):
            area = next(iter(key.sqref.ranges))
            assert area.min_row == area.max_row == 3
    for name in ['Aadhar - XD', 'Hyperverge - XD']:
        assert not workbook[name].conditional_formatting
    workbook.close()


def test_no_red_scale_when_only_unselected_facility_is_positive(tmp_path):
    profile = load_profile()
    rows = rows_for(profile)
    for row in rows:
        row[20] = 'red' if row[7] == 'Unselected DC' else 'green'
    report = reference_prepare(profile, rows)
    path = tmp_path / 'unselected.xlsx'
    writer.write_report(path, profile, report)
    workbook = load_workbook(path)
    for name in ['Aadhar - DC', 'Aadhar - XD', 'Hyperverge - DC', 'Hyperverge - XD']:
        assert not red_scales(workbook[name])
    workbook.close()


def test_base_omits_blank_cells_and_per_row_style_records(tmp_path):
    path, report = generate(tmp_path)
    with ZipFile(path) as archive:
        sheet = ET.fromstring(archive.read('xl/worksheets/sheet7.xml'))
        cells = sheet.findall('.//s:sheetData/s:row/s:c', NS)
        expected = 27 + sum(v is not None and v != '' for row in report.base_rows for v in row)
        assert len(cells) == expected
        assert all(c.find('s:v', NS) is not None for c in cells)
        for row in sheet.findall('.//s:sheetData/s:row', NS):
            assert not {'s', 'ht', 'customFormat', 'customHeight'}.intersection(row.attrib)
    workbook = load_workbook(path)
    base = workbook['Onboarding Base']
    assert base.freeze_panes == 'A2'
    assert base.sheet_view.showGridLines is not False
    for row in base:
        for cell in row:
            if cell.value is None: continue
            if isinstance(cell.value, datetime):
                assert cell.number_format == 'dd-mmm-yyyy'
            else:
                assert not cell.has_style, cell.coordinate
            assert not cell.font.bold and cell.fill.patternType is None
            assert all(getattr(cell.border, side).style is None for side in ['top', 'left', 'right', 'bottom'])
    workbook.close()


def test_raw_types_text_and_unicode_survive_plain_base(tmp_path):
    profile = load_profile()
    rows = rows_for(profile)
    rows[0][1] = '=1+1'
    rows[0][3] = False
    rows[0][5] = 0
    rows[0][6] = 12.5
    rows[0][11] = 'literal <xml> & text\nsecond line'
    rows[0][12] = 'https://example.invalid/not-a-link'
    rows[0][21] = '\u092a\u0930\u0940\u0915\u094d\u0937\u0923 \U0001f680'
    rows[0][18] = datetime(1990, 5, 2, 8, 30)
    path = tmp_path / 'types.xlsx'
    writer.write_report(path, profile, reference_prepare(profile, rows))
    workbook = load_workbook(path)
    base = workbook['Onboarding Base']
    for index in [0, 1, 3, 5, 6, 11, 12, 18, 21]:
        assert base.cell(2, index+1).value == rows[0][index]
    assert base['B2'].data_type == 's'
    assert base['M2'].hyperlink is None
    workbook.close()


def part_hashes(path):
    with ZipFile(path) as archive:
        return [(n, sha256(archive.read(n)).hexdigest()) for n in archive.namelist()]


def test_compaction_preserves_every_uncompressed_part_byte_for_byte(tmp_path):
    path, _ = generate(tmp_path)
    before = part_hashes(path)
    result = compact_workbook(path)
    assert part_hashes(path) == before
    assert result['after_bytes'] <= result['before_bytes']
    assert result['saved_bytes'] == result['before_bytes'] - result['after_bytes']
    assert not list(tmp_path.glob('.onboarding-compact-*'))


def test_compaction_reduces_compressible_zip_and_preserves_comment(tmp_path):
    path = tmp_path / 'stored.xlsx'
    with ZipFile(path, 'w', compression=ZIP_STORED) as archive:
        archive.comment = b'preserve me'
        archive.writestr('xl/worksheets/sheet1.xml', b'<cell>same value</cell>' * 10000)
    before = part_hashes(path)
    result = compact_workbook(path)
    assert result['after_bytes'] < result['before_bytes'] / 10
    assert part_hashes(path) == before
    with ZipFile(path) as archive:
        assert archive.comment == b'preserve me'


def test_compression_failure_never_publishes_partial_output(tmp_path, monkeypatch):
    profile = load_profile()
    path = tmp_path / 'previous.xlsx'
    path.write_bytes(b'EXISTING REPORT')
    def fail(_): raise OSError('synthetic disk failure')
    monkeypatch.setattr(writer, 'compact_workbook', fail)
    with pytest.raises(ReportError, match='synthetic disk failure'):
        writer.write_report(path, profile, reference_prepare(profile, rows_for(profile)))
    assert path.read_bytes() == b'EXISTING REPORT'
    assert not list(tmp_path.glob('.onboarding-*'))


def test_bad_archive_is_not_replaced_and_temporary_files_are_removed(tmp_path):
    path = tmp_path / 'invalid.xlsx'
    path.write_bytes(b'NOT A ZIP')
    with pytest.raises(Exception):
        compact_workbook(path)
    assert path.read_bytes() == b'NOT A ZIP'
    assert not list(tmp_path.glob('.onboarding-compact-*'))
