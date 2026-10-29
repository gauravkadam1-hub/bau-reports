"""Target row is presentation-only and specific to Housekeeping/Security."""
from datetime import date, datetime, timedelta
from pathlib import Path
import hashlib
import xml.etree.ElementTree as ET
from zipfile import ZipFile
import pytest
import xlsxwriter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from engines.attendance import engine as e
from tests.attendance_target_fixtures import make_inputs, DC, XD

EXPECTED = ('-', '-', .75, '-', '-', '-', '-', 1., '-', '-', .80)
SHEETS = ('Attendance Adherence - DC', 'Attendance Adherence - XD')


def spec_for(role):
    main=role=='Main Attendance'
    return e.ReportSpec('main' if main else f'role:{role.strip().casefold()}', role,
                        ('GC',) if main else (role,),None if main else 'Test - Deployed HC',
                        'Active \nHC' if main else ('Deployed \nHC' if role.strip().casefold()=='housekeeping' else 'Required HC'),main)


def make_summary(path,role,date_count,site_count):
    dates=[date(2026,10,1)+timedelta(days=i) for i in range(date_count)]
    sites=[f'Test {i} DC' for i in range(site_count)]
    values={d:{s:[8,5,.625,1,1,1,2,.8,.2,.2,.4] for s in sites} for d in dates}
    with xlsxwriter.Workbook(path,{'constant_memory':True,'use_future_functions':True}) as wb:
        for sn in SHEETS:
            e._write_summary_sheet(wb,sn,spec_for(role),sites,dates,values,
                                   {'facility':'H','status':'G','onboarded_date':'I'},'B')
    e._inject_summary_date_merges(Path(path),len(dates))
    return dates, sites


def check_target_sheet(ws, date_count, site_count):
    assert ws['A3'].value=='Target'
    assert ws.freeze_panes=='B4'
    assert ws.max_row==site_count+4
    assert ws.cell(ws.max_row,1).value=='Grand Total'
    assert str(ws.merged_cells).count('A1:A2')==1
    for di in range(date_count):
        col=2+di*11
        assert tuple(ws.cell(3,col+j).value for j in range(11))==EXPECTED
        for j in range(11):
            cell=ws.cell(3,col+j)
            assert cell.font.bold and cell.font.sz==10
            assert cell.fill.fgColor.rgb[-6:]=='DBE9F7'
            assert cell.alignment.horizontal=='center'
            assert cell.alignment.vertical=='center'
            assert all(getattr(cell.border,side).style=='thin' for side in ('left','right','top','bottom'))
            assert cell.number_format==('0%' if j in (2,7,10) else 'General')
            assert cell.data_type==('n' if j in (2,7,10) else 's')
        assert f'{get_column_letter(col)}1:{get_column_letter(col+10)}1' in ws.merged_cells
        for offset in (0,1,3,4,5,6):
            letter=get_column_letter(col+offset)
            expected=f'=SUM({letter}4:{letter}{3+site_count})' if site_count else '=0'
            assert ws.cell(ws.max_row,col+offset).value==expected
        if site_count:
            assert '$A4' in ws.cell(4,col).value
            assert '$A4' in ws.cell(4,col+1).value
            assert ws.cell(4,col+2).value==f'=IFERROR({get_column_letter(col+1)}4/{get_column_letter(col)}4,0)'
    assert len(ws.conditional_formatting)==(date_count*5 if site_count else 0)
    for conditional in ws.conditional_formatting:
        for span in conditional.sqref.ranges:
            assert span.min_row==4 and span.max_row==site_count+3
            assert span.min_row>3 and span.max_row<ws.max_row
    for di in range(date_count):
        col=2+di*11
        dims=[d for d in ws.column_dimensions.values() if (d.min or 1)<=col<=(d.max or d.min or 1)]
        assert dims and dims[0].hidden==(di<max(0,date_count-2))


@pytest.mark.parametrize('role',['Housekeeping','Security'])
@pytest.mark.parametrize('date_count',[1,2,5,23,31])
@pytest.mark.parametrize('site_count',[0,1,3])
def test_all_date_targets_styles_formulas_and_exclusions(tmp_path,role,date_count,site_count):
    path=tmp_path/'targets.xlsx'
    make_summary(path,role,date_count,site_count)
    wb=load_workbook(path,data_only=False)
    cached=load_workbook(path,data_only=True)
    try:
        for sn in SHEETS:
            check_target_sheet(wb[sn],date_count,site_count)
            assert cached[sn]['A3'].value=='Target'
            assert cached[sn].cell(cached[sn].max_row,2).value==8*site_count
            assert cached[sn].cell(cached[sn].max_row,4).value==(.625 if site_count else 0)
    finally:
        wb.close();cached.close()


@pytest.mark.parametrize('role',['Main Attendance','Maintenance','Housekeeping Supervisor','Security Lead'])
def test_main_and_other_roles_have_no_target_row(tmp_path,role):
    path=tmp_path/'no_targets.xlsx';make_summary(path,role,5,2)
    wb=load_workbook(path,data_only=False)
    try:
        for sn in SHEETS:
            ws=wb[sn]
            assert ws['A3'].value=='Test 0 DC'
            assert ws.freeze_panes=='B3'
            assert all(c.value!='Target' for row in ws for c in row)
            assert ws['B5'].value=='=SUM(B3:B4)'
            for cf in ws.conditional_formatting:
                assert all(r.min_row==3 and r.max_row==4 for r in cf.sqref.ranges)
    finally:wb.close()


@pytest.mark.parametrize('role',[' HOUSEKEEPING ','security','SeCuRiTy'])
def test_exact_role_matching_is_case_and_whitespace_tolerant(tmp_path,role):
    path=tmp_path/'case.xlsx';make_summary(path,role,2,1)
    wb=load_workbook(path)
    try:
        for sn in SHEETS:check_target_sheet(wb[sn],2,1)
    finally:wb.close()


def assert_package(path):
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith(('.xml','.rels')):ET.fromstring(z.read(name))
        ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        styles=ET.fromstring(z.read('xl/styles.xml'))
        style_count=len(styles.find('m:cellXfs',ns))
        for family in styles.findall('.//m:family',ns):int(family.get('val'))
        for name in z.namelist():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                root=ET.fromstring(z.read(name))
                for cell in root.findall('.//m:c',ns):
                    assert int(cell.get('s','0'))<style_count
                    assert cell.get('t')!='e'
        assert ET.fromstring(z.read('docProps/core.xml')).find('{http://purl.org/dc/elements/1.1/}title') is None


def generate_and_check(tmp_path,monkeypatch,date_count,backend,dc_only=False):
    if backend=='polars':
        module=pytest.importorskip('polars',reason='Real Polars is required for this backend test.')
        monkeypatch.setattr(e,'pl',module)
    else:monkeypatch.setattr(e,'pl',None)
    paths=make_inputs(tmp_path/'inputs',date_count)
    monkeypatch.setattr(e,'MASTER_HC_PATH',paths['master'])
    monkeypatch.setattr(e,'REF_MASTER_PATH',paths['ref'])
    monkeypatch.setattr(e,'ASSET_DIR',paths['assets'])
    before={k:hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items() if isinstance(p,Path) and p.is_file()}
    selection=[e.MAIN_REPORT_TOKEN,'Housekeeping','Security','Maintenance']
    outputs=e.generate_reports(e.ReportInputs(str(paths['users1']),str(paths['users2']),str(paths['attendance']),
                                str(tmp_path/'outputs'),paths['dates'][0],paths['dates'][-1],DC,[] if dc_only else XD,selection))
    assert len(outputs)==4
    for output in outputs:
        assert_package(output)
        wb=load_workbook(output,data_only=False);cached=load_workbook(output,data_only=True)
        try:
            targeted=any(role in output.name for role in ['Housekeeping','Security'])
            for sn,site_count in [(SHEETS[0],2),(SHEETS[1],0 if dc_only else 2)]:
                if targeted:check_target_sheet(wb[sn],date_count,site_count)
                else:
                    assert wb[sn]['A3'].value!='Target'
                    assert wb[sn].freeze_panes=='B3'
                for row in cached[sn]:
                    assert all(c.data_type!='e' for c in row)
            for ws in wb:
                if ws.title not in SHEETS:
                    assert not any(cell.value=='Target' for row in ws for cell in row)
            assert 'Excluded HUB' not in [row[7] for row in cached['All User data'].iter_rows(min_row=2,values_only=True)]
            assert 'Excluded HUB' not in [row[5] for row in cached['Attendance Base'].iter_rows(min_row=3,values_only=True)]
        finally:wb.close();cached.close()
    assert before=={k:hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items() if isinstance(p,Path) and p.is_file()}
    return outputs


@pytest.mark.parametrize('date_count',[1,2,5,23])
def test_full_generation_existing_pandas_fallback(tmp_path,monkeypatch,date_count):
    generate_and_check(tmp_path,monkeypatch,date_count,'pandas')


def test_dedicated_dc_only_keeps_empty_xd_targets_safe(tmp_path,monkeypatch):
    generate_and_check(tmp_path,monkeypatch,3,'pandas',dc_only=True)


@pytest.mark.parametrize('date_count',[1,5,23])
def test_full_generation_real_polars_backend(tmp_path,monkeypatch,date_count):
    generate_and_check(tmp_path,monkeypatch,date_count,'polars')
