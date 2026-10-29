from __future__ import annotations

import ast
import json
import math
import re
import shutil
import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from engines.onboarding.onboarding_common import Counts, load_profile, REQUIRED_TEMPLATE_SHEETS, excel_text_matches
from engines.onboarding.onboarding_writer import write_report
from engines.onboarding import report_generator as engine
from engines.onboarding.utils import parse_date_value, output_filename, ReportError
from tests.onboarding_fixtures import rows_for, reference_prepare, START, END, DC, XD

ROOT=Path(__file__).resolve().parents[1]
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

@pytest.fixture
def profile(): return load_profile()

@pytest.fixture
def generated(tmp_path, profile):
    report=reference_prepare(profile,rows_for(profile))
    path=tmp_path/'synthetic.xlsx'
    write_report(path,profile,report)
    return path,report


def test_sheet_order(generated):
    with ZipFile(generated[0]) as archive:
        w=ET.fromstring(archive.read('xl/workbook.xml'))
        assert [s.attrib['name'] for s in w.find('s:sheets',NS)]==list(REQUIRED_TEMPLATE_SHEETS)


def test_every_xml_part_is_well_formed(generated):
    with ZipFile(generated[0]) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith(('.xml','.rels')): ET.fromstring(archive.read(name))
        assert not any('externalLinks/' in n or 'calcChain' in n for n in archive.namelist())


def test_freeze_panes_all_seven_sheets(generated):
    w=load_workbook(generated[0])
    for sheet in w:
        assert sheet.freeze_panes==('B3' if sheet.title.startswith(('Aadhar','Hyperverge')) else 'A2')
        assert sheet.sheet_view.pane.state=='frozen'
        assert len(sheet.sheet_view.selection)==(3 if sheet.freeze_panes=='B3' else 1)
    w.close()


def test_hidden_dates_and_visible_latest_two(generated):
    w=load_workbook(generated[0])
    for sheet in w:
        if not sheet.title.startswith(('Aadhar','Hyperverge')): continue
        width=4 if sheet.title.startswith('Aadhar') else 2
        def hidden(col):
            return any(dim.hidden and dim.min<=col<=dim.max for dim in sheet.column_dimensions.values())
        for index in range(4):
            for col in range(2+index*width,2+(index+1)*width):
                assert hidden(col)==(index<2)
        assert not hidden(1)
        assert not any(d.hidden for d in sheet.row_dimensions.values())
    w.close()


def test_red_white_scale_only_summary_data_rows(generated):
    w=load_workbook(generated[0])
    for name in ['Aadhar - DC','Aadhar - XD']:
        s=w[name]; total=s.max_row
        scales=[]
        for key,rules in s.conditional_formatting._cf_rules.items():
            for rule in rules:
                if rule.type=='colorScale' and len(rule.colorScale.color)==2:
                    scales.append((key,rule))
        assert len(scales)==1  # The preceding date is all zero and has no RED scale.
        for key,rule in scales:
            cell_range=list(key.sqref.ranges)[0]
            assert cell_range.min_row==3 and cell_range.max_row==total-1
            assert cell_range.min_col in [18,19] and cell_range.max_col==cell_range.min_col
            assert [c.rgb[-6:] for c in rule.colorScale.color]==['FFFFFF','F8696B']
            assert rule.colorScale.cfvo[0].type=='num' and rule.colorScale.cfvo[0].val==0
    w.close()


def test_original_manual_and_compliance_rules_remain(generated):
    w=load_workbook(generated[0]); s=w['Aadhar - DC']
    rules=[r for v in s.conditional_formatting._cf_rules.values() for r in v]
    assert sum(r.type=='cellIs' and r.operator=='greaterThan' for r in rules)==4
    assert sum(r.type=='colorScale' and len(r.colorScale.color)==3 for r in rules)==4
    w.close()


def test_cached_values_and_known_counts(generated):
    w=load_workbook(generated[0],data_only=True)
    dc=w['Aadhar - DC']
    assert dc['A3'].value=='Beta DC'
    assert [dc.cell(3,c).value for c in range(14,18)]==[3,0,3,0]
    assert dc['S3'].value==3
    assert w['Hyperverge - DC']['I3'].value==3
    assert w['Aadhar - DC']['N6'].value==5  # 3 Beta + 2 exact/case-insensitive Alpha; padded site excluded.
    w.close()


def test_formula_cache_matches_independent_formula_evaluation(generated):
    w=load_workbook(generated[0],data_only=False)
    cached=load_workbook(generated[0],data_only=True)
    memo={}
    def evaluate(sheet,coordinate):
        key=(sheet,coordinate.replace('$',''))
        if key in memo: return memo[key]
        value=w[sheet][key[1]].value
        if not isinstance(value,str) or not value.startswith('='): return value
        expression=value[1:]
        if expression=='0': result=0
        elif expression.startswith('COUNTIFS('):
            args=[a.strip() for a in expression[9:-1].split(',')]
            conditions=[]
            for i in range(0,len(args),2):
                col=re.search(r'!\$([A-Z]+):',args[i]).group(1)
                criterion=args[i+1]
                criterion=criterion[1:-1] if criterion.startswith('"') else evaluate(sheet,criterion)
                conditions.append((column_index_from_string(col),criterion))
            def match(raw,criterion):
                if isinstance(criterion,datetime): return raw==criterion
                if isinstance(criterion,str): return excel_text_matches(raw,criterion)
                return raw==criterion
            base=w['Onboarding Base']
            result=sum(all(match(base.cell(row,col).value,criterion) for col,criterion in conditions) for row in range(2,base.max_row+1))
        elif expression.startswith('SUM('):
            start,end=expression[4:-1].split(':')
            result=sum(evaluate(sheet,c.coordinate) or 0 for row in w[sheet][start+':'+end] for c in row)
        elif expression.startswith('IFERROR('):
            numerator,denominator=expression[8:].split(',',1)[0].split('/')
            den=evaluate(sheet,denominator)
            result=evaluate(sheet,numerator)/den if den else 1
        elif '!' in expression:
            other,ref=expression.split('!',1); result=evaluate(other.strip("'"),ref)
        else:
            result=evaluate(sheet,expression)
        memo[key]=result
        return result
    formula_count=0
    for sheet in w:
        for row in sheet:
            for cell in row:
                if cell.data_type=='f':
                    expected=evaluate(sheet.title,cell.coordinate)
                    actual=cached[sheet.title][cell.coordinate].value
                    assert actual is not None,(sheet.title,cell.coordinate)
                    assert actual==expected or (isinstance(expected,float) and math.isclose(actual,expected)),(sheet.title,cell.coordinate,actual,expected)
                    formula_count+=1
    assert formula_count>100
    w.close();cached.close()


@pytest.mark.parametrize('dc,xd',[([],[]),(DC,[]),([],XD)])
def test_empty_groups_no_circular_totals(tmp_path,profile,dc,xd):
    path=tmp_path/'empty.xlsx'; report=reference_prepare(profile,rows_for(profile),dc=dc,xd=xd)
    write_report(path,profile,report)
    w=load_workbook(path);c=load_workbook(path,data_only=True)
    for kind,sites in [('DC',dc),('XD',xd)]:
        if sites: continue
        for prefix in ['Aadhar','Hyperverge']:
            sheet=w[f'{prefix} - {kind}']
            assert sheet['A3'].value=='Grand Total'
            assert sheet['B3'].value=='=0'
            assert c[sheet.title]['B3'].value==0
            assert not sheet.conditional_formatting
            assert not any(':B2' in str(cell.value) for row in sheet for cell in row)
    w.close();c.close()


def test_one_day_summary(tmp_path,profile):
    path=tmp_path/'day.xlsx'; report=reference_prepare(profile,rows_for(profile),first=END,last=END)
    write_report(path,profile,report)
    w=load_workbook(path)
    assert w['Aadhar - DC'].max_column==6
    rules=[r for v in w['Aadhar - DC'].conditional_formatting._cf_rules.values() for r in v]
    assert sum(r.type=='colorScale' and len(r.colorScale.color)==2 for r in rules)==1
    assert w['Aadhar - DC']['F2'].value=='=B1'
    w.close()


def test_unselected_facilities_still_in_base_and_hr(generated):
    w=load_workbook(generated[0],data_only=True)
    assert 'Unselected DC' in [row[7] for row in w['Onboarding Base'].iter_rows(min_row=2,values_only=True)]
    assert 'Synthetic Unselected HR' in [row[0] for row in w['Manual Onb - HR wise details'].iter_rows(min_row=2,values_only=True)]
    w.close()


def test_no_pandas_or_openpyxl_imports_in_runtime():
    for p in (ROOT/'engines/onboarding').glob('*.py'):
        tree=ast.parse(p.read_text())
        for node in ast.walk(tree):
            if isinstance(node,ast.Import):
                assert not any(n.name.split('.')[0] in ['pandas','openpyxl'] for n in node.names),p.name
            if isinstance(node,ast.ImportFrom):
                assert (node.module or '').split('.')[0] not in ['pandas','openpyxl'],p.name


def test_source_text_is_not_executed(tmp_path,profile):
    records=rows_for(profile); records[0][1]='=1+1'
    path=tmp_path/'literal.xlsx'; write_report(path,profile,reference_prepare(profile,records))
    w=load_workbook(path)
    assert w['Onboarding Base']['B2'].value=='=1+1'
    assert w['Onboarding Base']['B2'].data_type=='s'
    w.close()


def test_raw_leading_zeros_and_duplicate_rows(generated):
    w=load_workbook(generated[0])
    base=w['Onboarding Base']
    assert base['A2'].value=='0001'
    assert base['R2'].value=='001234567890'
    assert sum(row[0]=='0008' for row in base.iter_rows(min_row=2,values_only=True))==2
    w.close()


def test_old_template_not_needed(tmp_path,profile):
    # The new runtime needs only the declarative JSON; it does not even check
    # that an XLSX exists at the portal's legacy template_path.
    source=ROOT/'engines/onboarding/onboarding_layout.json'
    shutil.copyfile(source,tmp_path/'onboarding_layout.json')
    assert load_profile(tmp_path/'nonexistent_master.xlsx')['headers']==profile['headers']


def test_legacy_default_persistence_api_writes_json_only(tmp_path):
    shutil.copyfile(ROOT/'engines/onboarding/onboarding_layout.json',tmp_path/'onboarding_layout.json')
    engine.add_new_facilities_to_master({'DC':['Brand New DC','Brand New DC'],'XD':['Brand New XD']},tmp_path/'master_template.xlsx')
    assert not (tmp_path/'master_template.xlsx').exists()
    p=load_profile(tmp_path/'master_template.xlsx')
    assert p['defaults']['DC'].count('Brand New DC')==1
    assert 'Brand New XD' in p['defaults']['XD']


def test_write_failure_preserves_previous_output(tmp_path,profile):
    path=tmp_path/'existing.xlsx'; path.write_bytes(b'PREVIOUS FILE')
    report=reference_prepare(profile,rows_for(profile)); report.row_count+=1
    with pytest.raises(ReportError,match='number of written'):
        write_report(path,profile,report)
    assert path.read_bytes()==b'PREVIOUS FILE'
    assert not list(tmp_path.glob('.onboarding-*.xlsx'))


def test_progress_is_monotonic(generated,tmp_path,profile):
    events=[]
    write_report(tmp_path/'progress.xlsx',profile,reference_prepare(profile,rows_for(profile)),lambda n,t: events.append((n,t)))
    assert [n for n,_ in events]==sorted(n for n,_ in events)
    assert all(isinstance(t,str) and t for _,t in events)


@pytest.mark.parametrize('value,expected',[
    ('03/09/2026',date(2026,9,3)), ('09/13/2026',date(2026,9,13)),
    ('2026-09-01T12:30:00Z',START),('1 September 2026',START),
    ('01-Sep-2026',START), ('2026/09/01',START),
    ('46000',None),(1,date(1900,1,1)),(59,date(1900,2,28)),
    (60,date(1900,2,28)),(61,date(1900,3,1)),(0.5,None),(None,None),('N/A',None),
])
def test_original_date_rules(value,expected): assert parse_date_value(value)==expected


@pytest.mark.parametrize('day,suffix',[(1,'1st'),(2,'2nd'),(3,'3rd'),(11,'11th'),(12,'12th'),(13,'13th'),(21,'21st')])
def test_original_filename(day,suffix): assert output_filename(date(2026,6,day))==f'Onboarding Report - {suffix} June.xlsx'


def test_excel_wildcards_and_nontrimmed_matching():
    assert excel_text_matches('Alpha DC','Alpha*')
    assert excel_text_matches('Alpha* DC','Alpha~* DC')
    assert not excel_text_matches('Alpha DC ','Alpha DC')
    assert excel_text_matches('ALPHA DC','Alpha DC')


def test_portal_loader_uses_unchanged_public_api():
    import inspect
    from portal.engine_loader import onboarding_engine
    loaded=onboarding_engine()
    assert list(inspect.signature(loaded.generate_report).parameters)==[
        'input_path','template_path','selected_start_date','selected_end_date','selected_dc','selected_xd',
        'output_dir','progress_callback','generation_date']

@pytest.mark.parametrize('scenario',['four_days','one_day','calendar_gap','default_sites'])
def test_matches_original_generator_cells_and_formulas(tmp_path,profile,scenario):
    golden=json.loads((ROOT/'tests/fixtures/onboarding_legacy_golden.json').read_text())[scenario]
    cfg=golden['config']
    report=reference_prepare(profile,rows_for(profile),first=date.fromisoformat(cfg['first']),last=date.fromisoformat(cfg['last']),dc=cfg['dc'],xd=cfg['xd'])
    path=tmp_path/'parity.xlsx';write_report(path,profile,report)
    w=load_workbook(path,data_only=False)
    observed={}
    for sheet in w:
        observed[sheet.title]={}
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    observed[sheet.title][cell.coordinate]=({'type':'datetime','value':cell.value.isoformat()} if isinstance(cell.value,datetime)
                                                            else {'type':cell.data_type,'value':cell.value})
    w.close()
    assert observed==golden['cells']


def test_dependency_error_is_actionable(tmp_path,monkeypatch):
    import builtins
    original=builtins.__import__
    def blocked(name,*args,**kwargs):
        if name=='polars': raise ModuleNotFoundError("No module named 'polars'",name='polars')
        return original(name,*args,**kwargs)
    monkeypatch.setattr(builtins,'__import__',blocked)
    # Test only when Polars is not already cached in the data module.
    import sys
    monkeypatch.delitem(sys.modules,'engines.onboarding.onboarding_data',raising=False)
    monkeypatch.delitem(sys.modules,'onboarding_data',raising=False)
    with pytest.raises(ReportError,match='requires Polars'):
        engine._data()
