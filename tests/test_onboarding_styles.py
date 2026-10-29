"""Regressions for Excel's styles.xml recovery prompt and lost formatting."""
from __future__ import annotations

import copy
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import load_workbook

from engines.onboarding.onboarding_common import load_profile
from engines.onboarding.onboarding_styles import format_properties, validate_workbook
from engines.onboarding import onboarding_writer as writer
from engines.onboarding.utils import ReportError
from tests.onboarding_fixtures import rows_for, reference_prepare, END

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

@pytest.fixture
def generated(tmp_path):
    profile = load_profile()
    rows = rows_for(profile)
    # Exercise actual Excel dates, not only text that happens to look like a date.
    for name in ['Registered Date', 'Rejection Date', 'Date of Birth']:
        rows[0][profile['headers'].index(name)] = datetime(2026, 10, 23)
    path = tmp_path / 'styles.xlsx'
    writer.write_report(path, profile, reference_prepare(profile, rows))
    return path


def replace_part(path, part, change):
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    root = ET.fromstring(parts[part])
    change(root)
    parts[part] = ET.tostring(root, encoding='utf-8', xml_declaration=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)


def test_existing_layout_has_integer_font_families():
    assert all(type(props['font_family']) is int for props in load_profile()['formats'].values() if 'font_family' in props)


@pytest.mark.parametrize('value', [0, 1, 2, 2.0, 3.0, 4, 5])
def test_integral_font_family_values_normalized(value):
    result = format_properties({'font_family': value, 'font_size': 11.5})
    assert type(result['font_family']) is int
    assert result['font_family'] == int(value)
    assert result['font_size'] == 11.5


@pytest.mark.parametrize('value', [2.5, True, False, '2.0', None, -1, 6, float('nan'), float('inf')])
def test_invalid_font_family_rejected(value):
    with pytest.raises(ReportError, match='font_family'):
        format_properties({'font_family': value})


def test_old_float_layout_cannot_regress(tmp_path):
    profile = copy.deepcopy(load_profile())
    for properties in profile['formats'].values():
        properties['font_family'] = 2.0
    path = tmp_path / 'old-layout.xlsx'
    writer.write_report(path, profile, reference_prepare(profile, rows_for(profile)))
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read('xl/styles.xml'))
        assert root.findall('.//' + NS + 'family')
        assert all(n.get('val') == '2' for n in root.findall('.//' + NS + 'family'))
    assert validate_workbook(path)['invalid_style_references'] == 0


def test_generated_styles_and_references_valid(generated):
    result = validate_workbook(generated)
    assert result['worksheets'] == 7
    assert result['checked_cells'] > 100
    assert result['conditional_rules'] > 10
    assert result['invalid_style_references'] == 0


def test_regression_rejects_well_formed_but_invalid_family_xml(generated):
    def corrupt(root):
        root.find('.//' + NS + 'family').set('val', '2.0')
    replace_part(generated, 'xl/styles.xml', corrupt)
    # The old "all XML parses" test passes this invalid document. Our guard fails.
    with ZipFile(generated) as archive:
        ET.fromstring(archive.read('xl/styles.xml'))
    with pytest.raises(ReportError, match='family.*integer'):
        validate_workbook(generated)


@pytest.mark.parametrize('part,tag,attribute,value,message', [
    ('xl/styles.xml', 'fonts', 'count', '999', 'incorrect fonts count'),
    ('xl/styles.xml', 'xf', 'fontId', '9999', 'fontId'),
    ('xl/styles.xml', 'xf', 'fillId', '9999', 'fillId'),
    ('xl/styles.xml', 'xf', 'borderId', '9999', 'borderId'),
    ('xl/styles.xml', 'xf', 'numFmtId', '9999', 'undefined number format'),
    ('xl/worksheets/sheet1.xml', 'c', 's', '9999', 'outside its allowed range'),
    ('xl/worksheets/sheet1.xml', 'cfRule', 'dxfId', '9999', 'dxfId'),
    ('xl/worksheets/sheet1.xml', 'cfRule', 'priority', '1.5', 'integer'),
])
def test_corrupt_style_links_rejected(generated, part, tag, attribute, value, message):
    replace_part(generated, part, lambda root: root.find('.//' + NS + tag).set(attribute, value))
    with pytest.raises(ReportError, match=message):
        validate_workbook(generated)


def test_bad_conditional_color_rejected(generated):
    replace_part(generated, 'xl/worksheets/sheet1.xml',
                 lambda root: root.find('.//' + NS + 'colorScale/' + NS + 'color').set('rgb', 'RED'))
    with pytest.raises(ReportError, match='RGB color'):
        validate_workbook(generated)


def test_validation_failure_does_not_replace_existing_output(tmp_path, monkeypatch):
    target = tmp_path / 'previous.xlsx'
    target.write_bytes(b'KEEP EXISTING FILE')
    def fail(path):
        raise ReportError('Injected style validation failure')
    monkeypatch.setattr(writer, 'validate_workbook', fail)
    profile = load_profile()
    with pytest.raises(ReportError, match='style validation'):
        writer.write_report(target, profile, reference_prepare(profile, rows_for(profile)))
    assert target.read_bytes() == b'KEEP EXISTING FILE'
    assert not list(tmp_path.glob('.onboarding-*.xlsx'))


def test_summary_dates_and_hyperverge_link_dates_are_formatted(generated):
    workbook = load_workbook(generated, data_only=True)
    for name in ['Aadhar - DC', 'Aadhar - XD', 'Hyperverge - DC', 'Hyperverge - XD']:
        cell = workbook[name]['B1']
        assert isinstance(cell.value, datetime)
        assert cell.number_format == 'dd-mmm-yyyy'
    for name in ['Aadhar - DC', 'Aadhar - XD']:
        sheet = workbook[name]
        for col in [sheet.max_column - 1, sheet.max_column]:
            cell = sheet.cell(2, col)
            assert isinstance(cell.value, datetime)
            assert cell.number_format == 'dd-mmm-yyyy'
            assert sheet.column_dimensions[cell.column_letter].width >= 12
    workbook.close()


def test_base_dates_have_explicit_date_formats(generated):
    workbook = load_workbook(generated)
    sheet = workbook['Onboarding Base']
    for ref in ['I2', 'J2', 'K2', 'S2']:
        assert isinstance(sheet[ref].value, datetime)
        assert sheet[ref].number_format == 'dd-mmm-yyyy'
    workbook.close()


def test_report_tables_except_plain_base_have_borders(generated):
    workbook = load_workbook(generated)
    for sheet in workbook:
        if sheet.title == 'Onboarding Base':
            continue
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    assert all(getattr(cell.border, side).style == 'thin'
                               for side in ['left', 'right', 'top', 'bottom']), (sheet.title, cell.coordinate)
    workbook.close()


def test_base_headers_are_plain_and_unstyled(generated):
    workbook = load_workbook(generated)
    sheet = workbook['Onboarding Base']
    for cell in sheet[1]:
        assert not cell.has_style
        assert not cell.font.bold
        assert not cell.alignment.wrap_text
        assert cell.fill.patternType is None
        assert all(getattr(cell.border, side).style is None
                   for side in ['left', 'right', 'top', 'bottom'])
    assert not any(dimension.height for dimension in sheet.row_dimensions.values())
    workbook.close()


def test_dates_preserve_values_not_convert_to_text(generated):
    with ZipFile(generated) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet7.xml'))
        for coordinate in ['I2', 'J2', 'K2', 'S2']:
            cell = root.find(f'.//{NS}c[@r="{coordinate}"]')
            assert cell.get('t', 'n') == 'n'
            assert float(cell.find(NS + 'v').text) > 40000
