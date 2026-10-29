from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import hashlib
import csv

import pandas as pd
import pytest
from openpyxl import load_workbook

from engines.facility_scope import is_dcxd_site, filter_pandas_dcxd, filter_polars_dcxd
from tests.dcxd_fixtures import mixed_inputs, FIRST, LAST, DC, XD, EXTRA

CASES = [
    ('Alpha DC', True), ('Alpha XD', True), ('alpha dc', True), ('alpha xD', True),
    ('  Alpha   DC  ', True), ('Alpha\tXD', True), ('Alpha\nDC', True),
    ('Alpha\u00a0XD', True), ('Hyderabad REV DC', True), ('Brand New XD', True),
    ('Alpha-DC', False), ('Alpha_DC', False), ('AlphaDC', False), ('AlphaXD', False),
    ('Alpha DC Annex', False), ('Alpha XD 2', False), ('DC Alpha', False), ('DC', False),
    ('XD', False), (' DC ', False), ('XD ', False), ('', False), ('   ', False),
    (None, False), (float('nan'), False), (123, False), ('Alpha HUB', False),
    ('Alpha DCX', False), ('Alpha XDC', False), ('Alpha dc\n', True),
]

@pytest.mark.parametrize('value,expected', CASES)
def test_reference_suffix_contract(value, expected):
    assert is_dcxd_site(value) is expected


def test_pandas_preserves_order_duplicates_values_types_and_original():
    frame = pd.DataFrame({'site': [x for x, _ in CASES], 'id': range(len(CASES))}, index=range(100,100+len(CASES)))
    before = frame.copy(deep=True)
    result = filter_pandas_dcxd(frame, 'site')
    pd.testing.assert_frame_equal(frame, before)
    expected = before.loc[[100+i for i, (_, valid) in enumerate(CASES) if valid]].reset_index(drop=True)
    pd.testing.assert_frame_equal(result, expected)
    duplicated = pd.DataFrame({'site': ['A DC','A DC'], 'id': ['001','001']})
    pd.testing.assert_frame_equal(filter_pandas_dcxd(duplicated, 'site'), duplicated)

@pytest.mark.parametrize('sites', [[], [None,'','HUB','SiteDC']])
def test_pandas_empty_scope_retains_columns(sites):
    frame = pd.DataFrame({'site': sites, 'id': range(len(sites))})
    result = filter_pandas_dcxd(frame, 'site')
    assert result.empty and list(result.columns) == ['site','id']


def test_missing_column_fails_closed_without_guessing():
    with pytest.raises(ValueError, match='facility column'):
        filter_pandas_dcxd(pd.DataFrame({'Other facility': ['A DC']}), 'Registered Facility')


def test_real_polars_contract_and_immutability():
    pl = pytest.importorskip('polars', reason='Real Polars is required.')
    values = [None if v is None else str(v) for v, _ in CASES]
    frame = pl.DataFrame({'site': values, 'id': range(len(values))})
    before = frame.clone()
    result = filter_polars_dcxd(frame, 'site')
    assert frame.equals(before)
    assert result['id'].to_list() == [i for i, (_, valid) in enumerate(CASES) if valid]
    assert result['site'].to_list() == [values[i] for i, (_, valid) in enumerate(CASES) if valid]
    assert filter_polars_dcxd(frame.head(0), 'site').schema == frame.schema


def bases_are_scoped(book, specs):
    for name, column, first_row in specs:
        ws = book[name]
        values = [ws.cell(r, column).value for r in range(first_row, ws.max_row+1)]
        assert values, name
        assert all(is_dcxd_site(value) for value in values), (name, values)


def assert_valid_package(path):
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith('.xml') or name.endswith('.rels'):
                ET.fromstring(z.read(name))
        root = ET.fromstring(z.read('docProps/core.xml'))
        assert root.find('{http://purl.org/dc/elements/1.1/}title') is None
        styles = ET.fromstring(z.read('xl/styles.xml'))
        for family in styles.findall('.//s:family',ns):
            int(family.get('val'))
    for cached in (True, False):
        wb = load_workbook(path, data_only=cached)
        try:
            for ws in wb:
                for row in ws:
                    for cell in row:
                        assert cell.data_type != 'e', (ws.title, cell.coordinate, cell.value)
                        if cell.data_type == 'f':
                            assert '#REF!' not in cell.value
        finally:
            wb.close()


def test_roster_loaders_scope_before_id_resolution(tmp_path):
    from engines.roster import engine as e
    paths = mixed_inputs(tmp_path)
    roster = e.load_roster(str(paths['roster']))
    attendance = e.load_attendance(str(paths['attendance']), roster)
    assert len(roster.source) == len(attendance.raw) == len(attendance.resolved) == 6
    assert all(is_dcxd_site(v) for v in roster.source[roster.facility_col])
    assert EXTRA in set(roster.source[roster.facility_col])
    assert all(is_dcxd_site(v) for v in attendance.resolved['Facility Name'])
    assert roster.source.index.tolist() == list(range(6))


def test_outside_duplicate_does_not_block_roster_but_inside_duplicate_does(tmp_path):
    from engines.roster import engine as e
    paths = mixed_inputs(tmp_path)
    with paths['roster'].open('a', newline='',encoding='utf-8') as f:
        csv.writer(f).writerow(['v5-001','Outside duplicate','GC','Regional HUB','ACTIVE','Onboarded','01-Oct-2026','','P1','P1'])
    assert len(e.load_roster(str(paths['roster'])).source) == 6
    with paths['roster'].open('a', newline='',encoding='utf-8') as f:
        csv.writer(f).writerow(['v5-001','Inside duplicate','GC',XD,'ACTIVE','Onboarded','01-Oct-2026','','P1','P1'])
    with pytest.raises(e.ReportError, match='unique'):
        e.load_roster(str(paths['roster']))


def test_actual_roster_output_all_three_bases_scoped_and_uploads_untouched(tmp_path):
    from engines.roster import engine as e
    paths = mixed_inputs(tmp_path/'inputs')
    before = {k: hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items()}
    path = e.generate_report(e.ReportInputs(str(paths['roster']),str(paths['attendance']),str(tmp_path/'output'),FIRST,LAST,[DC],[XD]))
    book = load_workbook(path, data_only=True)
    try:
        bases_are_scoped(book, [('Roster Base',4,2),('Master Base',4,3),('Attendance Base',4,2)])
        assert EXTRA in [book['Roster Base'].cell(r,4).value for r in range(2,book['Roster Base'].max_row+1)]
    finally:
        book.close()
    assert before == {k: hashlib.sha256(p.read_bytes()).hexdigest() for k,p in paths.items()}
    assert_valid_package(path)


@pytest.fixture
def attendance_pandas(monkeypatch):
    from engines.attendance import engine
    # Exercise its explicitly supported pandas fallback, even on Polars machines.
    monkeypatch.setattr(engine,'pl',None)
    return engine


def test_attendance_loaders_scope_both_uploads(attendance_pandas,tmp_path):
    e = attendance_pandas
    paths = mixed_inputs(tmp_path)
    users=e.load_users(str(paths['users1']),str(paths['users2']))
    att=e.load_attendance(str(paths['attendance']))
    assert len(users.frame)==len(att.frame)==6
    assert all(is_dcxd_site(v) for v in users.frame[users.required['facility']])
    assert all(is_dcxd_site(v) for v in att.frame[att.identity_map['facility_name']])
    assert EXTRA in set(users.frame[users.required['facility']])
    assert att.dates==[FIRST,LAST]


def test_actual_attendance_main_and_dedicated_exports_scoped(attendance_pandas,tmp_path):
    e=attendance_pandas
    paths=mixed_inputs(tmp_path/'inputs')
    outputs=e.generate_reports(e.ReportInputs(str(paths['users1']),str(paths['users2']),str(paths['attendance']),str(tmp_path/'out'),FIRST,LAST,[DC],[XD],[e.MAIN_REPORT_TOKEN,'Security']))
    assert len(outputs)==2
    for path in outputs:
        wb=load_workbook(path,data_only=True)
        try:
            au=wb['All User data']; ab=wb['Attendance Base']
            au_headers=[au.cell(1,c).value for c in range(1,au.max_column+1)]
            ab_headers=[ab.cell(1,c).value for c in range(1,ab.max_column+1)]
            au_fac=au_headers.index('Registered Facility')+1
            ab_fac=ab_headers.index('facility_name')+1
            bases_are_scoped(wb,[('All User data',au_fac,2),('Attendance Base',ab_fac,3)])
            if 'Security' in path.name:
                role_col=au_headers.index('Role')+1
                assert {au.cell(r,role_col).value for r in range(2,au.max_row+1)}=={'Security'}
        finally:
            wb.close()
        assert_valid_package(path)


def test_orphan_attendance_lookup_has_same_safe_fallback_as_cache(attendance_pandas,tmp_path):
    e=attendance_pandas
    paths=mixed_inputs(tmp_path/'inputs',cross_facility=True)
    output,=e.generate_reports(e.ReportInputs(str(paths['users1']),str(paths['users2']),str(paths['attendance']),str(tmp_path/'out'),FIRST,LAST,[DC],[XD],[e.MAIN_REPORT_TOKEN]))
    formula=load_workbook(output,data_only=False)
    cached=load_workbook(output,data_only=True)
    try:
        ws=formula['Attendance Base']
        row=next(r for r in range(3,ws.max_row+1) if ws.cell(r,3).value=='v5-orphan')
        assert ws.cell(row,2).value.startswith('=IFERROR(')
        assert cached['Attendance Base'].cell(row,2).value is False
        assert 'v5-orphan' not in [r[2] for r in cached['All User data'].iter_rows(min_row=2,values_only=True)]
    finally:
        formula.close();cached.close()


def test_polars_attendance_loaders_and_generation(tmp_path):
    pytest.importorskip('polars',reason='Real Polars is required.')
    from engines.attendance import engine as e
    assert e.pl is not None
    paths=mixed_inputs(tmp_path/'input')
    users=e.load_users(str(paths['users1']),str(paths['users2']))
    att=e.load_attendance(str(paths['attendance']))
    assert users.frame.height==att.frame.height==6
    output,=e.generate_reports(e.ReportInputs(str(paths['users1']),str(paths['users2']),str(paths['attendance']),str(tmp_path/'out'),FIRST,LAST,[DC],[XD],[e.MAIN_REPORT_TOKEN]))
    wb=load_workbook(output,data_only=True)
    try: bases_are_scoped(wb,[('All User data',8,2),('Attendance Base',6,3)])
    finally: wb.close()
    assert_valid_package(output)


def test_polars_flex_loaders_central_and_final_workbook(tmp_path):
    pytest.importorskip('polars',reason='Real Polars is required.')
    pytest.importorskip('xlrd',reason='Flex imports xlrd for legacy .xls support.')
    from engines.flex import engine as e
    paths=mixed_inputs(tmp_path/'input')
    flex=e.load_flex(str(paths['flex']))
    users=e.load_users(str(paths['users1']),str(paths['users2']))
    att=e.load_attendance(str(paths['attendance']))
    assert flex.data.height==12 and users.data.height==att.data.height==6
    assert all(is_dcxd_site(v) for v in flex.data['Location'].to_list())
    central=e._derive_central(flex,users,att)
    assert central.height==12
    output=e.generate_report(e.ReportInputs(str(paths['flex']),str(paths['users1']),str(paths['users2']),str(paths['attendance']),str(tmp_path/'out'),FIRST,LAST,[DC],[XD]))
    wb=load_workbook(output,data_only=True)
    try: bases_are_scoped(wb,[('Central Flex Base',6,2),('Valinor Base',6,2),('Valinor Attendance',4,3)])
    finally: wb.close()
    assert_valid_package(output)


def test_flex_keeps_existing_subcost_and_invalid_date_exclusions(tmp_path):
    pytest.importorskip('polars',reason='Real Polars is required.')
    pytest.importorskip('xlrd',reason='Real Flex imports xlrd.')
    from engines.flex import engine as e
    paths=mixed_inputs(tmp_path)
    with paths['flex'].open('a',encoding='utf-8',newline='') as h:
        w=csv.writer(h)
        w.writerow(['Synthetic','P1','00','00','Vendor',DC,'2026-10-22','FM-RTS','HR','1','FT','Sourcing'])
        w.writerow(['Synthetic','P1','00','00','Vendor',XD,'invalid','DC operations','HR','1','FT','Sourcing'])
    data=e.load_flex(str(paths['flex']))
    assert data.removed_subcost==1 and data.removed_invalid_date==1
    assert data.data.height==12


def test_roster_dedup_prefers_retained_facility_and_ignores_outside_stint(tmp_path):
    from engines.roster import engine as e
    paths=mixed_inputs(tmp_path)
    roster=e.load_roster(str(paths['roster']))
    raw=e._load_attendance_raw(str(paths['attendance']))[0]
    outside=raw.iloc[[0]].copy()
    outside['Facility Name']='Outside HUB'
    outside[FIRST]='P3'
    # Outside duplicate is intentionally first; it must never win ID resolution.
    mixed=pd.concat([outside,raw],ignore_index=True)
    scoped=filter_pandas_dcxd(mixed,'Facility Name')
    resolved=e.resolve_attendance(scoped,roster)
    row=resolved.loc[resolved['Employee Code']=='v5-001'].iloc[0]
    assert row['Facility Name']==DC and row[FIRST]=='P1'


def test_scope_keeps_unknown_but_valid_dcxd_facilities():
    frame=pd.DataFrame({'facility':['Not In Any Master DC','Not In Any Master XD','Not In Any Master HUB']})
    assert filter_pandas_dcxd(frame,'facility')['facility'].tolist()==['Not In Any Master DC','Not In Any Master XD']


def test_onboarding_pipeline_scopes_hr_and_retains_unselected_dc(tmp_path):
    pytest.importorskip('polars',reason='Real Polars is required.')
    from engines.onboarding import onboarding_data as data
    from engines.onboarding import report_generator as generator
    from engines.onboarding.onboarding_common import load_profile
    from tests.onboarding_fixtures import rows_for,START,END,DC as ONB_DC,XD as ONB_XD
    from tests.dcxd_fixtures import write_csv
    profile=load_profile()
    source=write_csv(tmp_path/'onboarding.csv',profile['headers'],rows_for(profile))
    output,*_=generator.generate_report(source,selected_start_date=START,selected_end_date=END,
                                       selected_dc=ONB_DC,selected_xd=ONB_XD,output_dir=tmp_path)
    wb=load_workbook(output,data_only=True)
    try:
        base=wb['Onboarding Base']
        index=profile['headers'].index('Registered Facility')+1
        sites=[base.cell(r,index).value for r in range(2,base.max_row+1)]
        assert sites and all(is_dcxd_site(v) for v in sites)
        assert 'Unselected DC' in sites and 'Other Facility' not in sites
        for name in ('Manual Onb - HR wise details','RED Cases - HR wise details'):
            names=[wb[name].cell(r,1).value for r in range(1,wb[name].max_row+1)]
            assert 'Synthetic Other HR' not in names
            assert 'Synthetic Unselected HR' in names
    finally:wb.close()
    assert_valid_package(output)


def test_onboarding_all_outside_rows_produces_safe_empty_base(tmp_path):
    pytest.importorskip('polars',reason='Real Polars is required.')
    from engines.onboarding import report_generator as generator
    from engines.onboarding.onboarding_common import load_profile
    from tests.onboarding_fixtures import rows_for,START
    from tests.dcxd_fixtures import write_csv
    profile=load_profile()
    row=rows_for(profile)[0]
    row[profile['headers'].index('Registered Facility')]='Outside HUB'
    source=write_csv(tmp_path/'onboarding.csv',profile['headers'],[row])
    output,*_=generator.generate_report(source,selected_start_date=START,selected_end_date=START,
                                       selected_dc=['Synthetic DC'],selected_xd=['Synthetic XD'],output_dir=tmp_path)
    wb=load_workbook(output,data_only=True)
    try:
        assert wb['Onboarding Base'].max_row==1
        assert wb['Aadhar - DC']['B3'].value==0
        assert wb['Aadhar - XD']['B3'].value==0
    finally:wb.close()
    assert_valid_package(output)
