"""Real Polars/Calamine integration tests. They never substitute a fake backend.

Tests are explicitly skipped if dependencies are unavailable. Use the strict
verification runner to reject a partial/skip-only test run before deployment.
"""
import csv
import json
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from engines.onboarding import report_generator as engine
from engines.onboarding.onboarding_common import load_profile
from engines.onboarding.utils import ReportError
from tests.onboarding_fixtures import rows_for, reference_prepare, START, END, DC, XD
from tests.dcxd_fixtures import scoped_reference_rows

ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture
def pipeline():
    pytest.importorskip('polars',reason='Real Polars is required; no mocked dataframe tests count as integration.')
    from engines.onboarding import onboarding_data
    return onboarding_data

@pytest.fixture
def excel_pipeline(pipeline):
    pytest.importorskip('fastexcel',reason='Real Calamine/FastExcel is required for Excel ingestion tests.')
    return pipeline

@pytest.fixture
def profile():return load_profile()


def make_xlsx(tmp_path,profile,rows=None,headers=None,epoch=None):
    p=tmp_path/'onboarding_input.xlsx';w=Workbook();s=w.active;s.title='Export'
    if epoch is not None:w.epoch=epoch
    s.append(headers or profile['headers'])
    for row in rows if rows is not None else rows_for(profile):s.append(row)
    w.save(p);w.close();return p


def make_csv(tmp_path,profile,delimiter=',',encoding='utf-8-sig',rows=None):
    p=tmp_path/'onboarding_input.csv'
    with p.open('w',encoding=encoding,newline='') as handle:
        writer=csv.writer(handle,delimiter=delimiter);writer.writerow(profile['headers']);writer.writerows(rows_for(profile) if rows is None else rows)
    return p


@pytest.mark.parametrize('scenario',['four_days','one_day','calendar_gap','default_sites'])
def test_full_polars_xlsxwriter_matches_dcxd_scope(excel_pipeline,profile,tmp_path,scenario):
    golden=json.loads((ROOT/'tests/fixtures/onboarding_dcxd_golden.json').read_text())[scenario]
    cfg=golden['config'];source=make_xlsx(tmp_path,profile)
    output,first,last=engine.generate_report(source,selected_start_date=date.fromisoformat(cfg['first']),selected_end_date=date.fromisoformat(cfg['last']),selected_dc=cfg['dc'],selected_xd=cfg['xd'],output_dir=tmp_path)
    w=load_workbook(output,data_only=False);observed={}
    for s in w:
        observed[s.title]={}
        for row in s:
            for c in row:
                if c.value is not None:
                    observed[s.title][c.coordinate]={'type':'datetime','value':c.value.isoformat()} if isinstance(c.value,datetime) else {'type':c.data_type,'value':c.value}
    w.close();assert observed==golden['cells']
    assert (first,last)==(date.fromisoformat(cfg['first']),date.fromisoformat(cfg['last']))


@pytest.mark.parametrize('delimiter',[',',';','\t','|'])
def test_csv_delimiters_and_leading_zeros(pipeline,profile,tmp_path,delimiter):
    p=make_csv(tmp_path,profile,delimiter=delimiter)
    frame=pipeline.load_input(p,profile)
    assert frame['ID'][0]=='0001'
    assert frame['Phone Number'][0]=='001234567890'
    report=pipeline.prepare_report(frame,profile,START,END,DC,XD)
    expected=reference_prepare(profile,scoped_reference_rows(rows_for(profile), profile['headers']))
    assert report.counts==expected.counts
    assert report.dc_sites==expected.dc_sites and report.xd_sites==expected.xd_sites
    assert report.manual_counts==expected.manual_counts and report.red_counts==expected.red_counts


def test_cp1252_input(pipeline,profile,tmp_path):
    rows=rows_for(profile);rows[0][1]='Synthetic Caf\u00e9'
    frame=pipeline.load_input(make_csv(tmp_path,profile,encoding='cp1252',rows=rows),profile)
    assert frame['Name'][0]=='Synthetic Caf\u00e9'


def test_numeric_excel_serial_vs_numeric_string(excel_pipeline,profile,tmp_path):
    rows=rows_for(profile)[:2];serial=(START-date(1899,12,30)).days
    rows[0][9]=serial;rows[1][9]=str(serial)
    frame=excel_pipeline.load_input(make_xlsx(tmp_path,profile,rows=rows),profile)
    assert frame['Onboarded Date'].to_list()==[START,None]


def test_excel_1904_date_epoch(excel_pipeline,profile,tmp_path):
    rows=rows_for(profile)[:2]
    frame=excel_pipeline.load_input(make_xlsx(tmp_path,profile,rows=rows,epoch=CALENDAR_MAC_1904),profile)
    assert frame['Onboarded Date'].to_list()==[START,START]


def test_missing_and_duplicate_headers(excel_pipeline,profile,tmp_path):
    headers=profile['headers'].copy();headers[-1]='Name'
    with pytest.raises(ReportError,match='duplicate headers'):
        excel_pipeline.load_input(make_xlsx(tmp_path,profile,headers=headers),profile)
    headers=profile['headers'].copy();headers[0]='Wrong ID'
    with pytest.raises(ReportError,match='Missing required headers: ID'):
        excel_pipeline.load_input(make_xlsx(tmp_path,profile,headers=headers),profile)


def test_reordered_and_case_changed_headers(excel_pipeline,profile,tmp_path):
    headers=[h.upper()+' ' for h in reversed(profile['headers'])]
    rows=[list(reversed(r)) for r in rows_for(profile)]
    frame=excel_pipeline.load_input(make_xlsx(tmp_path,profile,headers=headers,rows=rows),profile)
    assert frame.columns==profile['headers']
    assert frame['ID'][0]=='0001'


def test_date_scan_and_new_facilities(excel_pipeline,profile,tmp_path):
    p=make_xlsx(tmp_path,profile)
    assert engine.inspect_input_date_range(p)==(START-timedelta(days=1),END+timedelta(days=1))
    facilities=engine.inspect_facilities(p,START,END)
    assert facilities['default_DC']==profile['defaults']['DC']
    assert 'Alpha DC' in facilities['DC'] and 'Other Facility' not in facilities['DC']
    assert engine.inspect_new_facilities(p,START,END)['DC']==['Alpha DC','Beta DC','Unselected DC']


@pytest.mark.parametrize('first,last',[(END,START),(START-timedelta(days=2),END),(START,END+timedelta(days=2))])
def test_invalid_date_windows(excel_pipeline,profile,tmp_path,first,last):
    with pytest.raises(ReportError):
        engine.generate_report(make_xlsx(tmp_path,profile),selected_start_date=first,selected_end_date=last,output_dir=tmp_path)


def test_no_valid_dates(pipeline,profile,tmp_path):
    rows=rows_for(profile)[:2]
    for row in rows:row[9]='not a date'
    with pytest.raises(ReportError,match='No valid dates'):
        engine.inspect_input_date_range(make_csv(tmp_path,profile,rows=rows))


@pytest.mark.parametrize('dc,xd',[(DC,[]),([],XD),([],[])])
def test_full_pipeline_empty_facility_groups(excel_pipeline,profile,tmp_path,dc,xd):
    path,*_=engine.generate_report(make_xlsx(tmp_path,profile),selected_start_date=START,selected_end_date=END,selected_dc=dc,selected_xd=xd,output_dir=tmp_path)
    w=load_workbook(path,data_only=True)
    for kind,sites in [('DC',dc),('XD',xd)]:
        if not sites:assert w['Aadhar - '+kind]['B3'].value==0
    w.close()


def test_empty_calendar_day_generates_zero_totals(excel_pipeline,profile,tmp_path):
    day=START+timedelta(days=2)
    p,*_=engine.generate_report(make_xlsx(tmp_path,profile),selected_start_date=day,selected_end_date=day,selected_dc=DC,selected_xd=XD,output_dir=tmp_path)
    w=load_workbook(p,data_only=True)
    assert w['Onboarding Base'].max_row==1
    assert w['Aadhar - DC']['B3'].value==0 and w['Aadhar - DC']['E3'].value==1
    w.close()


def test_csv_formula_looking_text_stays_text(pipeline,profile,tmp_path):
    rows=rows_for(profile);rows[0][1]='=1+1'
    p,*_=engine.generate_report(make_csv(tmp_path,profile,rows=rows),selected_start_date=START,selected_end_date=END,selected_dc=DC,selected_xd=XD,output_dir=tmp_path)
    w=load_workbook(p);assert w['Onboarding Base']['B2'].data_type=='s';w.close()


def test_portal_generate_one_integration(excel_pipeline,profile,tmp_path):
    from portal.generator import ReportRequest,generate_one
    from portal.file_detector import detect_files
    source=make_xlsx(tmp_path,profile)
    filename='onboarding_funnel_all_2026-09-04.xlsx'
    detected=detect_files([filename])
    paths=generate_one(ReportRequest('onboarding',START,END,DC,XD),detected,{filename:source},tmp_path/'out',lambda n,t:None)
    assert len(paths)==1 and paths[0].is_file()
