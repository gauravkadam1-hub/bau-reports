from __future__ import annotations

import csv
import math
import re
import os
import tempfile
import zipfile
from xml.sax.saxutils import escape as xml_escape
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable

import polars as pl
import xlsxwriter
import xlrd
from openpyxl import load_workbook
from xlsxwriter.utility import xl_col_to_name

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".xls"}

FLEX_CANONICAL_COLUMNS = [
    "Name", "Shift", "Mobile Number", "Aadhar Number", "Vendor Name", "Location",
    "Date", "Sub Cost Centre", "HRBP Name", "Flex rate", "Employee Type",
    "Type ( Sourcing / Internal)",
]

EXCLUDED_SUB_COST_CENTRES = {
    "FM-RTS", "LM-FM RTS", "FM/RTS EXPRESS", "FM/RTS",
}

DEFAULT_DC_SITES = [
    "NCR BILASPUR DC", "NCR BAMNOLI DC", "Ahmedabad DC", "Bangalore DC",
    "Hyderabad DC", "Jaipur DC", "Kolkata DC", "LUCKNOW DC", "MUMBAI DC",
    "Nagpur DC", "NCR Tauru DC", "ONE NCR DC", "PATNA DC", "SURAT DC",
    "Noida DC", "Chandigarh DC", "Gorakhpur DC", "Guwahati DC",
    "Ludhiana DC", "CHENNAI DC",
]

DEFAULT_XD_SITES = [
    "Indore XD", "BHUBANESWAR XD", "Trichy XD", "Pune XD", "Ludhiana XD",
    "Coimbatore XD", "Ranchi XD", "Kochi XD", "Raipur XD", "Vanaras XD",
    "Vijayawada XD", "Visakhapatnam XD", "Hubli XD", "Kozhikode XD",
    "Bhopal XD", "Madurai XD", "Kanpur XD", "Siliguri XD", "Jammu XD",
    "Rajahmundry XD", "Purnia XD", "Udaipur XD", "Mysore XD", "Aurangabad XD",
    "SAMBALPUR XD", "SILCHAR XD", "Malda XD", "Chennai XD", "JABALPUR XD",
    "Moradabad XD", "Dibrugarh XD", "GORAKHPUR XD", "VARANASI XD", "Jodhpur XD",
    "MORADABD XD", "CHANDIGARH XD",
]

ATTENDANCE_IDENTITY = [
    "user_id", "user_name", "role", "facility_name", "stint_id", "last_present_day"
]
ATTENDANCE_SUBHEADERS = ["shift", "total_hours", "working_hours", "punch_status"]

ProgressCallback = Callable[[int, str], None]


class ReportError(Exception):
    pass


@dataclass
class FlexData:
    data: pl.DataFrame
    original_rows: int
    removed_subcost: int
    removed_invalid_date: int
    removed_role: int


@dataclass
class UserData:
    data: pl.DataFrame
    id_col: str
    phone_col: str
    aadhaar_col: str
    facility_col: str


@dataclass
class AttendanceData:
    data: pl.DataFrame
    dates: list[date]
    blocks: dict[date, dict[str, str]]
    identity_map: dict[str, str]


@dataclass
class ScanResult:
    dc_sites: list[str]
    xd_sites: list[str]
    report_start: date
    report_end: date
    flex_original_rows: int
    flex_clean_rows: int
    removed_subcost: int
    removed_invalid_date: int
    removed_role: int
    user_rows: int
    attendance_rows: int
    attendance_start: date
    attendance_end: date


@dataclass
class ReportInputs:
    flex_path: str
    users_path_1: str
    users_path_2: str
    attendance_path: str
    output_folder: str
    report_start: date
    report_end: date
    selected_dc: list[str]
    selected_xd: list[str]


def _progress(cb: ProgressCallback | None, pct: int, message: str) -> None:
    if cb:
        cb(pct, message)


def normalize_site(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value).strip()).casefold()


def _clean_space(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value).strip())


def _norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", _cell_text(value).casefold())


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
    return str(value).strip()


def _find_col(columns: Iterable[object], aliases: Iterable[str], label: str, required: bool = True) -> str | None:
    aliases_n = {_norm_header(x) for x in aliases}
    for c in columns:
        if _norm_header(c) in aliases_n:
            return str(c)
    if required:
        raise ReportError(f"Required column not found: {label}")
    return None


def _unique_headers(values: Iterable[object]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for idx, value in enumerate(values, start=1):
        base = _cell_text(value).lstrip("\ufeff") or f"Column {idx}"
        key = base.casefold()
        seen[key] = seen.get(key, 0) + 1
        out.append(base if seen[key] == 1 else f"{base}_{seen[key]}")
    return out


def _stringify_frame(df: pl.DataFrame) -> pl.DataFrame:
    if not df.columns:
        return df
    return df.with_columns([
        pl.col(c).cast(pl.String, strict=False).fill_null("").alias(c)
        for c in df.columns
    ])


def _safe_frame_from_rows(
    headers: list[str],
    rows: Iterable[Iterable[object]],
    *,
    preserve_types: bool = False,
) -> pl.DataFrame:
    """Build a Polars frame without schema inference.

    Excel exports frequently contain mixed physical types in identifier columns
    (for example mostly numeric mobile numbers plus a value like ``98/07901625``).
    Constructing a DataFrame row-wise lets Polars infer a numeric builder first and
    then fail when later text appears.  Building explicit Series avoids that.
    """
    if not headers:
        return pl.DataFrame()

    buckets: list[list[object]] = [[] for _ in headers]
    width = len(headers)
    for row in rows:
        rr = list(row)
        for idx in range(width):
            value = rr[idx] if idx < len(rr) else None
            buckets[idx].append(value if preserve_types else _cell_text(value))

    dtype = pl.Object if preserve_types else pl.String
    series = [
        pl.Series(name, values, dtype=dtype, strict=False)
        for name, values in zip(headers, buckets)
    ]
    return pl.DataFrame(series)


def _read_excel_openpyxl(path: str, *, preserve_types: bool) -> pl.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = _unique_headers(next(rows))
        except StopIteration:
            return pl.DataFrame()
        return _safe_frame_from_rows(headers, rows, preserve_types=preserve_types)
    finally:
        wb.close()


def _read_excel_xls(path: str, *, preserve_types: bool) -> pl.DataFrame:
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        sh = book.sheet_by_index(0)
        if sh.nrows == 0:
            return pl.DataFrame()
        headers = _unique_headers(sh.row_values(0))

        def xls_rows():
            for r in range(1, sh.nrows):
                out: list[object] = []
                for cell in sh.row(r)[:len(headers)]:
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            out.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        except Exception:
                            out.append(cell.value)
                    else:
                        out.append(cell.value)
                yield out

        return _safe_frame_from_rows(headers, xls_rows(), preserve_types=preserve_types)
    finally:
        book.release_resources()


def _read_flat(path: str, *, object_types: bool = False) -> pl.DataFrame:
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ReportError(f"Unsupported file type: {ext}. Use CSV or Excel.")
    try:
        if ext == ".csv":
            df = pl.read_csv(
                path,
                infer_schema_length=0,
                try_parse_dates=False,
                ignore_errors=True,
                encoding="utf8-lossy",
            )
            rename_map = {c: _cell_text(c).lstrip("\ufeff") for c in df.columns}
            df = df.rename(rename_map)
            return _stringify_frame(df)
        if ext in {".xlsx", ".xlsm"}:
            return _read_excel_openpyxl(path, preserve_types=object_types)
        return _read_excel_xls(path, preserve_types=object_types)
    except Exception as exc:
        raise ReportError(f"Could not read {Path(path).name}: {exc}") from exc


def _parse_strict_flex_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    s = _cell_text(value)
    if not s:
        return None
    # Deliberately reject dots/slashes and silent auto-correction. The Flex source is expected
    # to use real Excel dates or hyphen-based date strings.
    if "." in s or "/" in s:
        return None
    formats = (
        "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d-%m-%Y %H:%M:%S",
        "%d-%b-%Y", "%d-%B-%Y", "%d-%b-%Y %H:%M:%S", "%d-%B-%Y %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_header_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = _cell_text(value)
    if not s or s.lower().startswith("unnamed"):
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def _canonical_flex_columns(df: pl.DataFrame) -> dict[str, str]:
    aliases = {
        "Name": ["Name", "Employee Name", "Flex Name"],
        "Shift": ["Shift"],
        "Mobile Number": ["Mobile Number", "Mobile", "Phone Number", "Phone"],
        "Aadhar Number": ["Aadhar Number", "Aadhaar Number", "Aadhar", "Aadhaar"],
        "Vendor Name": ["Vendor Name", "Vendor"],
        "Location": ["Location", "Facility", "Facility Name"],
        "Date": ["Date", "Flex Date", "Usage Date"],
        "Sub Cost Centre": ["Sub Cost Centre", "Sub Cost Center", "Sub CostCentre", "Sub CostCenter"],
        "HRBP Name": ["HRBP Name", "HRBP"],
        "Flex rate": ["Flex rate", "Flex Rate", "Rate"],
        "Employee Type": ["Employee Type", "Emp Type"],
        "Type ( Sourcing / Internal)": [
            "Type ( Sourcing / Internal)", "Type (Sourcing / Internal)",
            "Type ( Sourcing/Internal)", "Type (Sourcing/Internal)", "Sourcing / Internal"
        ],
    }
    return {out: _find_col(df.columns, names, out) for out, names in aliases.items()}


def load_flex(path: str) -> FlexData:
    raw = _read_flat(path, object_types=False)
    if raw.is_empty():
        raise ReportError("Flex Base is empty.")
    mapping = _canonical_flex_columns(raw)

    cols: dict[str, list[object]] = {}
    for out_col in FLEX_CANONICAL_COLUMNS:
        src = mapping[out_col]
        values = raw.get_column(src).to_list()
        # Keep all source fields schema-safe. Excel dates become ISO-like text here and
        # are parsed strictly below; identifiers stay textual so mixed values cannot fail.
        cols[out_col] = [_cell_text(v) for v in values] if out_col == "Date" else [_clean_space(v) for v in values]
    data = pl.DataFrame(cols, strict=False)

    original_rows = data.height
    excluded_norm = {_clean_space(x).upper() for x in EXCLUDED_SUB_COST_CENTRES}
    sub_norm = [_clean_space(x).upper() for x in data.get_column("Sub Cost Centre").to_list()]
    keep_sub = [x not in excluded_norm for x in sub_norm]
    removed_sub = keep_sub.count(False)
    data = data.filter(pl.Series("_keep_sub", keep_sub))

    parsed_dates = [_parse_strict_flex_date(v) for v in data.get_column("Date").to_list()]
    keep_date = [d is not None for d in parsed_dates]
    removed_bad = keep_date.count(False)
    data = data.filter(pl.Series("_keep_date", keep_date))
    good_dates = [d for d in parsed_dates if d is not None]
    data = data.with_columns(pl.Series("Date", good_dates, dtype=pl.Date))

    # Backwards-compatible scan metadata: exclusions are Sub Cost Centre based, not role based.
    removed_role = 0
    return FlexData(data, original_rows, removed_sub, removed_bad, removed_role)


def _sort_user_files(path1: str, path2: str) -> tuple[str, str]:
    date_pat = re.compile(r"(20\d{2})[-_](\d{2})[-_](\d{2})")
    def first_date(p: str):
        m = date_pat.search(Path(p).name)
        return tuple(map(int, m.groups())) if m else (9999, 99, 99)
    pair = [path1, path2]
    pair.sort(key=first_date)
    return pair[0], pair[1]


def _clean_text_expr(source: str, alias: str) -> pl.Expr:
    return (
        pl.col(source)
        .cast(pl.String, strict=False)
        .fill_null("")
        .str.strip_chars()
        .str.replace_all(r"\s+", " ")
        .alias(alias)
    )


def load_users(path1: str, path2: str) -> UserData:
    p1, p2 = _sort_user_files(path1, path2)
    d1 = _read_flat(p1)
    d2 = _read_flat(p2)
    if d1.is_empty() and d2.is_empty():
        raise ReportError("Both All Users Data files are empty.")

    cols = list(d1.columns)
    existing_norm = {_norm_header(c) for c in cols}
    for c in d2.columns:
        n = _norm_header(c)
        if n not in existing_norm:
            cols.append(c)
            existing_norm.add(n)

    def align(df: pl.DataFrame) -> pl.DataFrame:
        source_by_norm = {_norm_header(c): c for c in df.columns}
        exprs: list[pl.Expr] = []
        for c in cols:
            src = source_by_norm.get(_norm_header(c))
            exprs.append(_clean_text_expr(src, c) if src is not None else pl.lit("").alias(c))
        return df.select(exprs)

    data = pl.concat([align(d1), align(d2)], how="vertical")
    id_col = _find_col(data.columns, ["ID", "User ID", "user_id"], "Valinor ID")
    phone_col = _find_col(data.columns, ["Phone Number", "Mobile Number", "Phone", "Mobile"], "Phone Number")
    aadhaar_col = _find_col(data.columns, ["Aadhaar Number", "Aadhar Number", "Aadhaar", "Aadhar"], "Aadhaar Number")
    facility_col = _find_col(data.columns, ["Registered Facility", "Facility Name", "facility_name", "Facility"], "Registered Facility")
    return UserData(data, id_col, phone_col, aadhaar_col, facility_col)


def _attendance_headers_csv(path: str) -> tuple[list[object], list[object]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        try:
            h1 = next(reader)
            h2 = next(reader)
        except StopIteration as exc:
            raise ReportError("Attendance Base must contain two header rows.") from exc
    return h1, h2


def _read_attendance_source(path: str) -> tuple[pl.DataFrame, list[tuple[object, object]]]:
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ReportError(f"Unsupported Attendance file type: {ext}. Use CSV or Excel.")
    try:
        if ext == ".csv":
            h1, h2 = _attendance_headers_csv(path)
            df = pl.read_csv(
                path,
                has_header=False,
                skip_rows=2,
                infer_schema_length=0,
                try_parse_dates=False,
                ignore_errors=True,
                truncate_ragged_lines=True,
                encoding="utf8-lossy",
            )
            df = _stringify_frame(df)
            width = df.width
            h1 = list(h1) + [""] * max(0, width - len(h1))
            h2 = list(h2) + [""] * max(0, width - len(h2))
            return df, list(zip(h1[:width], h2[:width]))

        if ext in {".xlsx", ".xlsm"}:
            wb = load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                h1 = list(next(rows)); h2 = list(next(rows))
                width = max(len(h1), len(h2))
                h1 += [""] * (width - len(h1)); h2 += [""] * (width - len(h2))
                schema = [f"column_{i+1}" for i in range(width)]
                df = _safe_frame_from_rows(schema, rows, preserve_types=False)
                return df, list(zip(h1, h2))
            finally:
                wb.close()

        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sh = book.sheet_by_index(0)
            if sh.nrows < 2:
                raise ReportError("Attendance Base must contain two header rows.")
            h1 = sh.row_values(0); h2 = sh.row_values(1)
            width = max(len(h1), len(h2))
            schema = [f"column_{i+1}" for i in range(width)]
            df = _safe_frame_from_rows(
                schema,
                (sh.row_values(r)[:width] for r in range(2, sh.nrows)),
                preserve_types=False,
            )
            return df, list(zip(h1, h2))
        finally:
            book.release_resources()
    except ReportError:
        raise
    except Exception as exc:
        raise ReportError(f"Could not read Attendance Base: {exc}") from exc


def load_attendance(path: str) -> AttendanceData:
    raw, headers = _read_attendance_source(path)
    if raw.is_empty():
        raise ReportError("Attendance Base is empty.")

    identity_map: dict[str, str | None] = {}
    for desired in ATTENDANCE_IDENTITY:
        aliases = {
            "user_id": ["user_id", "User ID", "Employee Code", "ID"],
            "user_name": ["user_name", "User Name", "Name"],
            "role": ["role", "Role"],
            "facility_name": ["facility_name", "Facility Name", "Facility", "Registered Facility"],
            "stint_id": ["stint_id", "Stint ID"],
            "last_present_day": ["last_present_day", "Last Present Day"],
        }[desired]
        alias_norm = {_norm_header(x) for x in aliases}
        found: str | None = None
        for idx, (top, _sub) in enumerate(headers):
            if idx < raw.width and _norm_header(top) in alias_norm:
                found = raw.columns[idx]
                break
        if found is None and desired in {"user_id", "facility_name"}:
            raise ReportError(f"Required Attendance column not found: {desired}")
        identity_map[desired] = found

    discovered: dict[date, dict[str, str]] = {}
    sub_alias = {
        "shift": {"shift"},
        "total_hours": {"totalhours", "totalhour", "totalhrs"},
        "working_hours": {"workinghours", "workinghour", "workinghrs"},
        "punch_status": {"punchstatus", "status", "punch"},
    }
    current_date: date | None = None
    for idx, (top, sub) in enumerate(headers):
        if idx >= raw.width:
            break
        parsed = _parse_header_date(top)
        if parsed:
            current_date = parsed
        sub_n = _norm_header(sub)
        matched = next((canonical for canonical, aliases in sub_alias.items() if sub_n in aliases), None)
        if matched is not None and current_date is not None:
            discovered.setdefault(current_date, {})[matched] = raw.columns[idx]

    dates = sorted(discovered)
    if not dates:
        raise ReportError("No date blocks were detected in Attendance Base.")

    exprs: list[pl.Expr] = []
    for desired in ATTENDANCE_IDENTITY:
        source = identity_map.get(desired)
        exprs.append(_clean_text_expr(source, desired) if source is not None else pl.lit("").alias(desired))

    blocks: dict[date, dict[str, str]] = {}
    for d in dates:
        blocks[d] = {}
        for sub in ATTENDANCE_SUBHEADERS:
            out_name = f"{d.isoformat()}|{sub}"
            source = discovered[d].get(sub)
            exprs.append(_clean_text_expr(source, out_name) if source is not None else pl.lit("").alias(out_name))
            blocks[d][sub] = out_name

    out = raw.select(exprs)
    return AttendanceData(out, dates, blocks, {k: k for k in ATTENDANCE_IDENTITY})

def _validate_report_window(start: date, end: date) -> tuple[date, date]:
    if start > end:
        raise ReportError("From Date cannot be after To Date.")
    if (start.year, start.month) != (end.year, end.month):
        raise ReportError("From Date and To Date must be within the same calendar month.")
    max_end = date.today() - timedelta(days=2)
    if end > max_end:
        raise ReportError(
            f"Flex reports can run only through D-2. Maximum allowed To Date is {max_end:%d-%b-%Y}."
        )
    return start, end


def _site_display_order(sites: list[str], defaults: list[str]) -> list[str]:
    by_norm = {normalize_site(s): s for s in sites}
    ordered: list[str] = []
    used: set[str] = set()
    for d in defaults:
        n = normalize_site(d)
        if n in by_norm and n not in used:
            # Use the report's canonical casing for known sites.
            ordered.append(d)
            used.add(n)
    for n, label in sorted(by_norm.items(), key=lambda kv: kv[1].casefold()):
        if n not in used:
            ordered.append(label)
            used.add(n)
    return ordered


def _eligible_sites(flex: FlexData, users: UserData, attendance: AttendanceData, start: date, end: date) -> tuple[list[str], list[str]]:
    # Site choices come from the report month in the cleaned Flex Base. Rows after D-2 may still
    # introduce a new site, which is intentional: it will be selectable before its first summary day.
    month_rows = flex.data.filter(
        (pl.col("Date").dt.year() == start.year) & (pl.col("Date").dt.month() == start.month)
    )
    flex_sites_raw = [_clean_space(x) for x in month_rows.get_column("Location").to_list() if _clean_space(x)]

    flex_by_norm: dict[str, str] = {}
    for site in flex_sites_raw:
        n = normalize_site(site)
        if site.upper().endswith(" DC") or site.upper().endswith(" XD"):
            flex_by_norm.setdefault(n, site)

    user_facilities = {
        normalize_site(x) for x in users.data.get_column(users.facility_col).to_list() if _clean_space(x)
    }
    att_facilities = {
        normalize_site(x) for x in attendance.data.get_column("facility_name").to_list() if _clean_space(x)
    }

    eligible: dict[str, str] = {}
    for n, label in flex_by_norm.items():
        if n in user_facilities and n in att_facilities:
            eligible[n] = label

    # Hyderabad is one report row and represents both Hyderabad DC and Hyderabad REV DC.
    hyd = normalize_site("Hyderabad DC")
    hyd_rev = normalize_site("Hyderabad REV DC")
    hyd_available = hyd in eligible or hyd_rev in eligible
    eligible.pop(hyd_rev, None)
    if hyd_available:
        eligible[hyd] = "Hyderabad DC"

    dc = [label for label in eligible.values() if label.upper().endswith(" DC")]
    xd = [label for label in eligible.values() if label.upper().endswith(" XD")]
    return _site_display_order(dc, DEFAULT_DC_SITES), _site_display_order(xd, DEFAULT_XD_SITES)

def scan_inputs(
    flex_path: str,
    users_path_1: str,
    users_path_2: str,
    attendance_path: str,
    report_start: date,
    report_end: date,
) -> ScanResult:
    start, end = _validate_report_window(report_start, report_end)
    flex = load_flex(flex_path)
    users = load_users(users_path_1, users_path_2)
    attendance = load_attendance(attendance_path)

    required_dates = {start + timedelta(days=i) for i in range((end - start).days + 1)}
    missing_dates = sorted(required_dates.difference(attendance.dates))
    if missing_dates:
        preview = ", ".join(d.strftime("%d-%b") for d in missing_dates[:5])
        if len(missing_dates) > 5:
            preview += f" (+{len(missing_dates) - 5} more)"
        raise ReportError(
            f"Attendance does not contain every selected report date. Missing: {preview}. "
            f"Attendance currently runs from {min(attendance.dates):%d-%b-%Y} through {max(attendance.dates):%d-%b-%Y}."
        )

    dc, xd = _eligible_sites(flex, users, attendance, start, end)
    return ScanResult(
        dc, xd, start, end,
        flex.original_rows, len(flex.data), flex.removed_subcost, flex.removed_invalid_date,
        flex.removed_role, len(users.data), len(attendance.data), min(attendance.dates), max(attendance.dates)
    )


def ordinal(n: int) -> str:
    if 10 <= n % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def default_output_filename(as_of_date: date | None = None) -> str:
    d = as_of_date or date.today()
    return f"Flex Adoption Report - {ordinal(d.day)} {d:%b}.xlsx"


def _first_lookup(values: pl.Series, ids: pl.Series) -> dict[str, str]:
    out: dict[str, str] = {}
    for key, val in zip(values.to_list(), ids.to_list()):
        k = _clean_space(key)
        if k and k not in out:
            out[k] = _clean_space(val)
    return out


def _derive_central(flex: FlexData, users: UserData, attendance: AttendanceData) -> pl.DataFrame:
    central = flex.data.clone()

    user_ids = users.data.get_column(users.id_col)
    aad_map = _first_lookup(users.data.get_column(users.aadhaar_col), user_ids)
    phone_map = _first_lookup(users.data.get_column(users.phone_col), user_ids)

    m_vals: list[str] = []
    n_vals: list[str] = []
    o_vals: list[str] = []
    p_vals: list[str] = []

    for aad, phone in zip(
        central.get_column("Aadhar Number").to_list(),
        central.get_column("Mobile Number").to_list(),
    ):
        aad_key = _clean_space(aad)
        phone_key = _clean_space(phone)
        m = aad_map.get(aad_key, "") if aad_key else ""
        n = "" if m else (phone_map.get(phone_key, "") if phone_key else "")
        onboarded = "Onboarded" if (m or n) else "Not Onboarded"
        vid = (m or n) if onboarded == "Onboarded" else ""
        m_vals.append(m); n_vals.append(n); o_vals.append(onboarded); p_vals.append(vid)

    # Excel MATCH takes the first attendance occurrence by Valinor ID. Reproduce that for caches.
    first_att_idx: dict[str, int] = {}
    for idx, uid in enumerate(attendance.data.get_column("user_id").to_list()):
        k = _clean_space(uid)
        if k and k not in first_att_idx:
            first_att_idx[k] = idx

    punch_series = {d: attendance.data.get_column(attendance.blocks[d]["punch_status"]) for d in attendance.dates}
    work_series = {d: attendance.data.get_column(attendance.blocks[d]["working_hours"]) for d in attendance.dates}

    q_vals: list[str] = []
    r_vals: list[str] = []
    s_vals: list[str] = []
    t_vals: list[str] = []
    for vid, d in zip(p_vals, central.get_column("Date").to_list()):
        punch = ""
        work = "0"
        idx = first_att_idx.get(vid) if vid else None
        if idx is not None and d in attendance.blocks:
            punch = _clean_space(punch_series[d][idx])
            work_raw = _clean_space(work_series[d][idx])
            work = work_raw if work_raw not in {"", "-", "#NA"} else "0"
        absent = (not punch) or punch == "-" or punch.upper() == "A" or punch.upper() == "#NA" or punch == "0"
        present = "Absent" if absent else "Present"
        try:
            num = float(work)
        except Exception:
            num = 0.0
        if num <= 0:
            bucket = "-"
        elif num < 5:
            bucket = "Less than 5 hours"
        elif num < 8:
            bucket = "5 to 8 hours"
        else:
            bucket = "greater than 8 hours"
        q_vals.append(punch)
        r_vals.append(present)
        s_vals.append(work)
        t_vals.append(bucket)

    return central.with_columns([
        pl.Series("Aadhar Check", m_vals, dtype=pl.String),
        pl.Series("Mobile Check", n_vals, dtype=pl.String),
        pl.Series("Final Validation", o_vals, dtype=pl.String),
        pl.Series("Valinor ID", p_vals, dtype=pl.String),
        pl.Series("Check", q_vals, dtype=pl.String),
        pl.Series("Flex Check", r_vals, dtype=pl.String),
        pl.Series("Working Hours", s_vals, dtype=pl.String),
        pl.Series("Bucket", t_vals, dtype=pl.String),
    ])


def _summary_cached(central: pl.DataFrame, sites: list[str], dates: list[date], summary_type: int) -> dict[tuple[str, date], tuple]:
    hyd = normalize_site("Hyderabad DC")
    hyd_rev = normalize_site("Hyderabad REV DC")
    site_expr = (
        pl.col("Location")
        .cast(pl.String, strict=False)
        .fill_null("")
        .str.strip_chars()
        .str.replace_all(r"\s+", " ")
        .str.to_lowercase()
    )
    emp_expr = (
        pl.col("Employee Type").cast(pl.String, strict=False).fill_null("")
        .str.strip_chars().str.replace_all(r"\s+", " ").str.to_uppercase()
    )
    check_expr = (
        pl.col("Check").cast(pl.String, strict=False).fill_null("")
        .str.strip_chars().str.to_uppercase()
    )

    prepared = central.with_columns([
        site_expr.alias("_site_norm"),
        emp_expr.alias("_emp_upper"),
        check_expr.alias("_check_upper"),
    ]).with_columns(
        pl.when(pl.col("_site_norm").is_in([hyd, hyd_rev]))
        .then(pl.lit(hyd))
        .otherwise(pl.col("_site_norm"))
        .alias("_report_site_norm")
    )

    agg = prepared.group_by(["_report_site_norm", "Date"]).agg([
        pl.len().alias("total"),
        ((pl.col("Final Validation") == "Onboarded") & (pl.col("Flex Check") == "Present")).sum().alias("punched"),
        ((pl.col("Final Validation") == "Onboarded") & (pl.col("Flex Check") == "Absent")).sum().alias("not_punched"),
        (pl.col("Final Validation") == "Not Onboarded").sum().alias("not_on"),
        ((pl.col("_emp_upper") != "PT") & (pl.col("Final Validation") == "Onboarded") & (pl.col("Flex Check") == "Present")).sum().alias("total_p"),
        ((pl.col("_emp_upper") != "PT") & (pl.col("_check_upper") == "SP")).sum().alias("single"),
        ((pl.col("_emp_upper") != "PT") & (pl.col("Bucket") == "Less than 5 hours")).sum().alias("less5"),
        ((pl.col("_emp_upper") != "PT") & (pl.col("Bucket") == "5 to 8 hours")).sum().alias("five8"),
        ((pl.col("_emp_upper") != "PT") & (pl.col("Bucket") == "greater than 8 hours")).sum().alias("eight"),
    ])
    row_map = {(r["_report_site_norm"], r["Date"]): r for r in agg.iter_rows(named=True)}

    out: dict[tuple[str, date], tuple] = {}
    for site in sites:
        site_n = hyd if normalize_site(site) == hyd else normalize_site(site)
        for d in dates:
            r = row_map.get((site_n, d), {})
            if summary_type == 1:
                total = int(r.get("total", 0) or 0)
                punched = int(r.get("punched", 0) or 0)
                not_punched = int(r.get("not_punched", 0) or 0)
                not_on = int(r.get("not_on", 0) or 0)
                compliance = min(1.0, punched / total) if total else 1.0
                out[(site, d)] = (total, punched, not_punched, not_on, compliance)
            else:
                total_p = int(r.get("total_p", 0) or 0)
                single = int(r.get("single", 0) or 0)
                less5 = int(r.get("less5", 0) or 0)
                five8 = int(r.get("five8", 0) or 0)
                eight = int(r.get("eight", 0) or 0)
                double_ad = min(1.0, (total_p - single) / total_p) if total_p else 1.0
                eight_ad = min(1.0, eight / total_p) if total_p else 1.0
                out[(site, d)] = (total_p, single, less5, five8, eight, double_ad, eight_ad)
    return out

def _formats(workbook: xlsxwriter.Workbook) -> dict[str, object]:
    light_blue = "#DCE6F1"
    grand_fill = "#D9D9D9"
    return {
        "summary_site_header": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": light_blue,
            "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True,
        }),
        "summary_date_header": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": light_blue,
            "border": 1, "align": "center", "valign": "vcenter", "num_format": "dd-mm-yyyy",
        }),
        "summary_metric_header": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": light_blue,
            "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True,
        }),
        "summary_site": workbook.add_format({"font_name": "Calibri", "font_size": 11, "border": 1}),
        "summary_num": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "border": 1, "align": "center", "valign": "vcenter",
        }),
        "summary_pct": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "border": 1, "align": "center", "valign": "vcenter",
            "num_format": "0%",
        }),
        "grand_site": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": grand_fill, "border": 1,
        }),
        "grand_num": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": grand_fill, "border": 1,
            "align": "center", "valign": "vcenter",
        }),
        "grand_pct": workbook.add_format({
            "font_name": "Calibri", "font_size": 11, "bold": True, "bg_color": grand_fill, "border": 1,
            "align": "center", "valign": "vcenter", "num_format": "0%",
        }),
        "central_raw_header": workbook.add_format({
            "font_name": "Calibri", "font_size": 8, "bold": True, "bg_color": "#00FFFF",
            "border": 2, "align": "center", "valign": "vcenter", "text_wrap": True,
        }),
        "central_formula_header": workbook.add_format({
            "font_name": "Aptos Narrow", "font_size": 11, "bold": True, "bg_color": "#FFFF00",
            "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True,
        }),
        "central_raw": workbook.add_format({
            "font_name": "Arial", "font_size": 10, "border": 1, "valign": "vcenter",
        }),
        "central_raw_date": workbook.add_format({
            "font_name": "Arial", "font_size": 10, "border": 1, "valign": "vcenter", "num_format": "dd-mm-yyyy",
        }),
        "central_formula": workbook.add_format({
            "font_name": "Aptos Narrow", "font_size": 11, "border": 1, "valign": "vcenter",
        }),
        "plain_header": workbook.add_format({"font_name": "Aptos Narrow", "font_size": 11}),
        "plain_date": workbook.add_format({"font_name": "Aptos Narrow", "font_size": 11, "num_format": "yyyy-mm-dd"}),
        "att_subheader": workbook.add_format({"font_name": "Aptos Narrow", "font_size": 11}),
    }


def _site_formula_countifs(site: str, criteria_pairs: list[tuple[str, str]]) -> str:
    # criteria_pairs = [(column_letter, criterion_expression)] where criterion is already Excel syntax.
    def one(site_criterion: str) -> str:
        pieces = [f"'Central Flex Base'!$F:$F", site_criterion]
        for col, crit in criteria_pairs:
            pieces.extend([f"'Central Flex Base'!${col}:${col}", crit])
        return "COUNTIFS(" + ",".join(pieces) + ")"
    if normalize_site(site) == normalize_site("Hyderabad DC"):
        return one('"Hyderabad DC"') + "+" + one('"Hyderabad REV DC"')
    return one("$A{row}")


def _write_summary(workbook: xlsxwriter.Workbook, name: str, sites: list[str], dates: list[date], summary_type: int,
                   cached: dict[tuple[str, date], tuple], fmts: dict[str, object]) -> None:
    # Sort only the site rows by the final report date's requested KPI.
    # Grand Total is written separately below and therefore never participates in the sort.
    # Summary 1: Compliance % (tuple index 4).
    # Summary 2: 8 Hrs+ Adherence (tuple index 6).
    if dates:
        last_date = dates[-1]
        sort_idx = 4 if summary_type == 1 else 6
        sites = sorted(
            sites,
            key=lambda site: (float(cached[(site, last_date)][sort_idx]), str(site).casefold()),
        )

    ws = workbook.get_worksheet_by_name(name) or workbook.add_worksheet(name)
    ws.freeze_panes(0, 1)
    ws.hide_gridlines(0)
    # Keep the summary compact for easy sharing. The metric columns are intentionally narrower
    # than the legacy workbook and the wrapped header row gets a fixed tight height so there is
    # no large white cushion around the words.
    ws.set_column(0, 0, 14.4 if name.startswith("XD") else 13.6)
    ws.set_row(0, 14.4)
    ws.set_row(1, 43.2 if summary_type == 2 else 31.2)
    metrics = (
        ["Total Flex Used", "Onboarded & Punched", "Onboarded & Not Punched", "Not Onboarded", "Compliance %"]
        if summary_type == 1 else
        ["Total Punched", "Single Punch", "Less than 5 Hours", "5 to 8 Hours", "8 Hours+", "Double Punch Adherence", "8 Hrs+ Adherence"]
    )
    block = len(metrics)
    ws.merge_range(0, 0, 1, 0, "Site", fmts["summary_site_header"])

    visible_dates = set(dates[-3:])
    for di, d in enumerate(dates):
        c0 = 1 + di * block
        c1 = c0 + block - 1
        ws.merge_range(0, c0, 0, c1, datetime.combine(d, datetime.min.time()), fmts["summary_date_header"])
        for j, metric in enumerate(metrics):
            ws.write(1, c0 + j, metric, fmts["summary_metric_header"])
        # Tight wrapped widths: enough for readable 2-3 line headers without the excessive
        # horizontal whitespace that made the earlier build hard to share.
        widths = [9.0, 9.0, 10.0, 9.0, 7.0, 13.0, 11.0] if summary_type == 2 else [9.0, 12.0, 14.0, 10.0, 9.0]
        for j, width in enumerate(widths):
            opts = {"hidden": True} if d not in visible_dates else None
            ws.set_column(c0 + j, c0 + j, width, None, opts)

    for r_off, site in enumerate(sites):
        row = 2 + r_off
        excel_row = row + 1
        ws.write(row, 0, site, fmts["summary_site"])
        for di, d in enumerate(dates):
            c0 = 1 + di * block
            vals = cached[(site, d)]
            date_ref = f"{xl_col_to_name(c0)}$1"
            if summary_type == 1:
                formulas = [
                    _site_formula_countifs(site, [("G", date_ref)]),
                    _site_formula_countifs(site, [("G", date_ref), ("O", '"Onboarded"'), ("R", '"Present"')]),
                    _site_formula_countifs(site, [("G", date_ref), ("O", '"Onboarded"'), ("R", '"Absent"')]),
                    _site_formula_countifs(site, [("G", date_ref), ("O", '"Not Onboarded"')]),
                ]
                formulas = [f.replace("{row}", str(excel_row)) for f in formulas]
                for j in range(4):
                    ws.write_formula(row, c0 + j, "=" + formulas[j], fmts["summary_num"], vals[j])
                ws.write_formula(row, c0 + 4, f"=IFERROR(MIN(1,{xl_col_to_name(c0+1)}{excel_row}/{xl_col_to_name(c0)}{excel_row}),1)", fmts["summary_pct"], vals[4])
            else:
                criteria_common = [("G", date_ref)]
                formulas = [
                    _site_formula_countifs(site, criteria_common + [("O", '"Onboarded"'), ("R", '"Present"'), ("K", '"<>PT"')]),
                    _site_formula_countifs(site, criteria_common + [("Q", '"SP"'), ("K", '"<>PT"')]),
                    _site_formula_countifs(site, criteria_common + [("T", '"Less than 5 hours"'), ("K", '"<>PT"')]),
                    _site_formula_countifs(site, criteria_common + [("T", '"5 to 8 hours"'), ("K", '"<>PT"')]),
                    _site_formula_countifs(site, criteria_common + [("T", '"greater than 8 hours"'), ("K", '"<>PT"')]),
                ]
                formulas = [f.replace("{row}", str(excel_row)) for f in formulas]
                for j in range(5):
                    ws.write_formula(row, c0 + j, "=" + formulas[j], fmts["summary_num"], vals[j])
                total_l = xl_col_to_name(c0); single_l = xl_col_to_name(c0 + 1); eight_l = xl_col_to_name(c0 + 4)
                ws.write_formula(row, c0 + 5, f"=IFERROR(MIN(1,({total_l}{excel_row}-{single_l}{excel_row})/{total_l}{excel_row}),1)", fmts["summary_pct"], vals[5])
                ws.write_formula(row, c0 + 6, f"=IFERROR(MIN(1,{eight_l}{excel_row}/{total_l}{excel_row}),1)", fmts["summary_pct"], vals[6])

    grand_row = 2 + len(sites)
    grand_excel = grand_row + 1
    ws.write(grand_row, 0, "Grand Total", fmts["grand_site"])
    first_excel = 3
    last_excel = 2 + len(sites)
    for di, d in enumerate(dates):
        c0 = 1 + di * block
        if summary_type == 1:
            totals = [sum(cached[(s, d)][j] for s in sites) for j in range(4)]
            for j in range(4):
                letter = xl_col_to_name(c0 + j)
                ws.write_formula(grand_row, c0 + j, f"=SUM({letter}{first_excel}:{letter}{last_excel})", fmts["grand_num"], totals[j])
            comp = min(1.0, totals[1] / totals[0]) if totals[0] else 1.0
            ws.write_formula(grand_row, c0 + 4, f"=IFERROR(MIN(1,{xl_col_to_name(c0+1)}{grand_excel}/{xl_col_to_name(c0)}{grand_excel}),1)", fmts["grand_pct"], comp)
            pct_cols = [c0 + 4]
        else:
            totals = [sum(cached[(s, d)][j] for s in sites) for j in range(5)]
            for j in range(5):
                letter = xl_col_to_name(c0 + j)
                ws.write_formula(grand_row, c0 + j, f"=SUM({letter}{first_excel}:{letter}{last_excel})", fmts["grand_num"], totals[j])
            double_ad = min(1.0, (totals[0] - totals[1]) / totals[0]) if totals[0] else 1.0
            eight_ad = min(1.0, totals[4] / totals[0]) if totals[0] else 1.0
            ws.write_formula(grand_row, c0 + 5, f"=IFERROR(MIN(1,({xl_col_to_name(c0)}{grand_excel}-{xl_col_to_name(c0+1)}{grand_excel})/{xl_col_to_name(c0)}{grand_excel}),1)", fmts["grand_pct"], double_ad)
            ws.write_formula(grand_row, c0 + 6, f"=IFERROR(MIN(1,{xl_col_to_name(c0+4)}{grand_excel}/{xl_col_to_name(c0)}{grand_excel}),1)", fmts["grand_pct"], eight_ad)
            pct_cols = [c0 + 5, c0 + 6]

        for pc in pct_cols:
            ws.conditional_format(2, pc, grand_row - 1, pc, {
                "type": "3_color_scale", "min_color": "#F8696B", "mid_color": "#FFEB84", "max_color": "#63BE7B",
                "mid_type": "percentile", "mid_value": 50,
            })


def _write_valinor_base(workbook: xlsxwriter.Workbook, users: UserData, fmts: dict[str, object], cb: ProgressCallback | None) -> dict[str, int]:
    """Create the Valinor Base sheet skeleton. Rows are streamed into the XLSX after close."""
    ws = workbook.get_worksheet_by_name("Valinor Base") or workbook.add_worksheet("Valinor Base")
    ws.freeze_panes(1, 0)
    ws.hide_gridlines(0)
    for c, header in enumerate(users.data.columns):
        ws.write(0, c, header, fmts["plain_header"])
        width = min(55, max(9, len(header) + 2))
        ws.set_column(c, c, width)
    return {str(c): i for i, c in enumerate(users.data.columns)}

def _write_valinor_attendance(workbook: xlsxwriter.Workbook, attendance: AttendanceData, fmts: dict[str, object], cb: ProgressCallback | None) -> dict[date, int]:
    """Create Attendance header/column skeleton. Rows are streamed after workbook close."""
    ws = workbook.get_worksheet_by_name("Valinor Attendance") or workbook.add_worksheet("Valinor Attendance")
    ws.freeze_panes(2, 0)
    ws.hide_gridlines(0)
    c = 0
    for ident in ATTENDANCE_IDENTITY:
        ws.write(0, c, ident, fmts["plain_header"])
        ws.write_blank(1, c, None, fmts["att_subheader"])
        ws.set_column(c, c, 17 if ident in {"user_name", "facility_name"} else 12)
        c += 1
    block_starts: dict[date, int] = {}
    for d in attendance.dates:
        block_starts[d] = c
        for j, sub in enumerate(ATTENDANCE_SUBHEADERS):
            if j == 0:
                ws.write_datetime(0, c + j, datetime.combine(d, datetime.min.time()), fmts["plain_date"])
            else:
                ws.write_blank(0, c + j, None, fmts["plain_header"])
            ws.write(1, c + j, sub, fmts["att_subheader"])
            ws.set_column(c + j, c + j, 12)
        c += 4
    return block_starts

def _write_central(workbook: xlsxwriter.Workbook, central: pl.DataFrame, users: UserData,
                   user_col_map: dict[str, int], attendance: AttendanceData, fmts: dict[str, object], cb: ProgressCallback | None) -> None:
    ws = workbook.get_worksheet_by_name("Central Flex Base") or workbook.add_worksheet("Central Flex Base")
    ws.hide_gridlines(0)
    headers = list(central.columns)
    for c, h in enumerate(headers):
        ws.write(0, c, h, fmts["central_raw_header"] if c < 12 else fmts["central_formula_header"])
    widths = [40, 5, 23.9, 26.7, 49, 17.7, 10.4, 23.8, 20.8, 8, 10, 18, 15, 15, 15, 15, 12, 12, 13, 20]
    for c, w in enumerate(widths):
        ws.set_column(c, c, w)

    id_col = user_col_map[users.id_col]
    phone_col = user_col_map[users.phone_col]
    aad_col = user_col_map[users.aadhaar_col]
    id_l = xl_col_to_name(id_col)
    phone_l = xl_col_to_name(phone_col)
    aad_l = xl_col_to_name(aad_col)

    total = len(central)
    for r0, row in enumerate(central.iter_rows(named=False), start=1):
        excel_row = r0 + 1
        # Raw A:L
        for c in range(12):
            val = row[c]
            fmt = fmts["central_raw_date"] if c == 6 else fmts["central_raw"]
            if c == 6 and isinstance(val, date):
                ws.write_datetime(r0, c, datetime.combine(val, datetime.min.time()), fmt)
            elif _cell_text(val) == "":
                ws.write_blank(r0, c, None, fmt)
            else:
                ws.write(r0, c, val, fmt)

        m, n, o, p, q, rr, s, t = row[12:20]
        f_m = f'=IF(TRIM($D{excel_row})="","",IFERROR(INDEX(\'Valinor Base\'!${id_l}:${id_l},MATCH($D{excel_row},\'Valinor Base\'!${aad_l}:${aad_l},0)),""))'
        f_n = f'=IF(LEN(TRIM($M{excel_row}))>0,"",IF(TRIM($C{excel_row})="","",IFERROR(INDEX(\'Valinor Base\'!${id_l}:${id_l},MATCH($C{excel_row},\'Valinor Base\'!${phone_l}:${phone_l},0)),"")))'
        f_o = f'=IF(AND(TRIM(M{excel_row})="",TRIM(N{excel_row})=""),"Not Onboarded","Onboarded")'
        f_p = f'=IF(O{excel_row}="Onboarded",IF(TRIM(M{excel_row})<>"",M{excel_row},N{excel_row}),"")'
        f_q = f'=IFERROR(INDEX(\'Valinor Attendance\'!$1:$1048576,MATCH($P{excel_row},\'Valinor Attendance\'!$A:$A,0),MATCH($G{excel_row},\'Valinor Attendance\'!$1:$1,0)+3),"")'
        f_r = f'=IFERROR(IF(OR(TRIM(Q{excel_row})="",TRIM(Q{excel_row})="-",UPPER(TRIM(Q{excel_row}))="A",TRIM(Q{excel_row})="#NA",Q{excel_row}=0),"Absent","Present"),"Absent")'
        f_s = f'=IFERROR(INDEX(\'Valinor Attendance\'!$1:$1048576,MATCH($P{excel_row},\'Valinor Attendance\'!$A:$A,0),MATCH($G{excel_row},\'Valinor Attendance\'!$1:$1,0)+2),"0")'
        f_t = f'=IFERROR(IF(OR(TRIM(S{excel_row})="",TRIM(S{excel_row})="-",TRIM(S{excel_row})="#NA",VALUE(S{excel_row})=0),"-",IF(VALUE(S{excel_row})<5,"Less than 5 hours",IF(VALUE(S{excel_row})<8,"5 to 8 hours","greater than 8 hours"))),"-")'
        formulas = [f_m, f_n, f_o, f_p, f_q, f_r, f_s, f_t]
        caches = [m, n, o, p, q, rr, s, t]
        for j, (formula, cache) in enumerate(zip(formulas, caches)):
            ws.write_formula(r0, 12 + j, formula, fmts["central_formula"], cache if cache != "" else "")
        if cb and r0 % 2000 == 0:
            _progress(cb, 68 + int(13 * r0 / max(total, 1)), f"Writing Central Flex Base formulas... {r0:,}/{total:,}")


def _xml_text(value: object) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    # Excel XML cannot contain most C0 controls. Keep tab/newline/carriage-return.
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)
    if len(text) > 32767:
        text = text[:32767]
    return xml_escape(text)


def _inline_cell(ref: str, value: object, style: str | None = None) -> str:
    text = _xml_text(value)
    if not text:
        return ""
    sattr = f' s="{style}"' if style is not None else ""
    preserve = ' xml:space="preserve"' if text[:1].isspace() or text[-1:].isspace() else ""
    return f'<c r="{ref}"{sattr} t="inlineStr"><is><t{preserve}>{text}</t></is></c>'


def _number_cell(ref: str, value: int | float, style: str | None = None) -> str:
    sattr = f' s="{style}"' if style is not None else ""
    return f'<c r="{ref}"{sattr}><v>{value}</v></c>'


def _escape_fast(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        text = value
    else:
        text = str(value)
    if not text:
        return ""
    if len(text) > 32767:
        text = text[:32767]
    # Inputs have already been cleaned while loading. Escape only XML metacharacters here.
    if "&" in text: text = text.replace("&", "&amp;")
    if "<" in text: text = text.replace("<", "&lt;")
    if ">" in text: text = text.replace(">", "&gt;")
    return text


def _write_streamed_valinor_xml(path: Path, prefix: bytes, suffix: bytes, users: UserData, cb: ProgressCallback | None) -> None:
    col_count = len(users.data.columns)
    letters = [xl_col_to_name(c) for c in range(col_count)]
    max_col = letters[-1]
    prefix = re.sub(rb'<dimension ref="[^"]*"\s*/>', f'<dimension ref="A1:{max_col}{len(users.data)+1}"/>'.encode(), prefix, count=1)
    with open(path, "wb", buffering=1024*1024) as f:
        f.write(prefix); f.write(b"<sheetData>")
        cells=[]
        for c,h in enumerate(users.data.columns):
            text=_escape_fast(h)
            if text: cells.append(f'<c r="{letters[c]}1" t="inlineStr"><is><t>{text}</t></is></c>')
        f.write((f'<row r="1">{"".join(cells)}</row>').encode())
        total=len(users.data)
        for i,row in enumerate(users.data.iter_rows(named=False), start=2):
            parts=[]
            for c,value in enumerate(row):
                if value == "" or value is None: continue
                text=_escape_fast(value)
                if text: parts.append(f'<c r="{letters[c]}{i}" t="inlineStr"><is><t>{text}</t></is></c>')
            f.write((f'<row r="{i}">{"".join(parts)}</row>').encode())
            if cb and i % 15000 == 0:
                _progress(cb, 84 + int(5 * (i-1) / max(total, 1)), f"Streaming Valinor Base... {i-1:,}/{total:,}")
        f.write(b"</sheetData>"); f.write(suffix)


def _write_streamed_attendance_xml(path: Path, prefix: bytes, suffix: bytes, attendance: AttendanceData,
                                   date_style: str | None, cb: ProgressCallback | None) -> None:
    total_cols = 6 + 4 * len(attendance.dates)
    letters=[xl_col_to_name(c) for c in range(total_cols)]
    max_col=letters[-1]
    prefix = re.sub(rb'<dimension ref="[^"]*"\s*/>', f'<dimension ref="A1:{max_col}{len(attendance.data)+2}"/>'.encode(), prefix, count=1)
    sattr=f' s="{date_style}"' if date_style is not None else ""
    with open(path, "wb", buffering=1024*1024) as f:
        f.write(prefix); f.write(b"<sheetData>")
        row1=[]
        for c,h in enumerate(ATTENDANCE_IDENTITY):
            row1.append(f'<c r="{letters[c]}1" t="inlineStr"><is><t>{_escape_fast(h)}</t></is></c>')
        c=6
        for d in attendance.dates:
            serial=(d-date(1899,12,30)).days
            row1.append(f'<c r="{letters[c]}1"{sattr}><v>{serial}</v></c>'); c += 4
        f.write((f'<row r="1">{"".join(row1)}</row>').encode())
        row2=[]; c=6
        for _d in attendance.dates:
            for sub in ATTENDANCE_SUBHEADERS:
                row2.append(f'<c r="{letters[c]}2" t="inlineStr"><is><t>{sub}</t></is></c>'); c += 1
        f.write((f'<row r="2">{"".join(row2)}</row>').encode())
        total=len(attendance.data)
        for ridx,row in enumerate(attendance.data.iter_rows(named=False), start=3):
            parts=[]
            for c,value in enumerate(row):
                if value == "" or value is None: continue
                text=_escape_fast(value)
                if text: parts.append(f'<c r="{letters[c]}{ridx}" t="inlineStr"><is><t>{text}</t></is></c>')
            f.write((f'<row r="{ridx}">{"".join(parts)}</row>').encode())
            if cb and ridx % 8000 == 0:
                _progress(cb, 89 + int(7 * (ridx-2) / max(total, 1)), f"Streaming Valinor Attendance... {ridx-2:,}/{total:,}")
        f.write(b"</sheetData>"); f.write(suffix)

def _inject_large_base_sheets(workbook_path: Path, users: UserData, attendance: AttendanceData,
                              cb: ProgressCallback | None) -> None:
    replacement = workbook_path.with_name(workbook_path.stem + ".streamed.xlsx")
    with zipfile.ZipFile(workbook_path, "r") as zin:
        vb = zin.read("xl/worksheets/sheet6.xml")
        att = zin.read("xl/worksheets/sheet7.xml")
        vb_open=vb.index(b"<sheetData>"); vb_close=vb.index(b"</sheetData>")+len(b"</sheetData>")
        att_open=att.index(b"<sheetData>"); att_close=att.index(b"</sheetData>")+len(b"</sheetData>")
        vb_prefix=vb[:vb_open]; vb_suffix=vb[vb_close:]
        att_prefix=att[:att_open]; att_suffix=att[att_close:]
        m=re.search(rb'<c r="G1"(?: s="(\d+)")?', att)
        date_style=(m.group(1).decode() if m and m.group(1) else None)
        with tempfile.TemporaryDirectory(prefix="flex_stream_") as td:
            vb_file=Path(td)/"sheet6.xml"; att_file=Path(td)/"sheet7.xml"
            _write_streamed_valinor_xml(vb_file, vb_prefix, vb_suffix, users, cb)
            _write_streamed_attendance_xml(att_file, att_prefix, att_suffix, attendance, date_style, cb)
            with zipfile.ZipFile(replacement, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=1) as zout:
                for info in zin.infolist():
                    if info.filename == "xl/worksheets/sheet6.xml":
                        zout.write(vb_file, info.filename)
                    elif info.filename == "xl/worksheets/sheet7.xml":
                        zout.write(att_file, info.filename)
                    else:
                        zout.writestr(info, zin.read(info.filename))
    os.replace(replacement, workbook_path)


def _validate_selected(selected: list[str], eligible: list[str], label: str) -> None:
    eligible_n = {normalize_site(s) for s in eligible}
    bad = [s for s in selected if normalize_site(s) not in eligible_n]
    if bad:
        raise ReportError(f"Selected {label} site(s) are no longer eligible: {', '.join(bad)}. Refresh inputs.")


def generate_report(inputs: ReportInputs, progress: ProgressCallback | None = None) -> str:
    start, end = _validate_report_window(inputs.report_start, inputs.report_end)
    if not inputs.selected_dc and not inputs.selected_xd:
        raise ReportError("Select at least one DC or XD site.")

    # Load each large source once with Polars. CSV parsing, filtering, joins/caches and summary
    # aggregation stay in the Polars pipeline; XlsxWriter is used only for the final Excel output.
    _progress(progress, 2, "Loading and validating four source files with Polars...")
    flex = load_flex(inputs.flex_path)
    _progress(progress, 6, "Combining All Users Data with Polars...")
    users = load_users(inputs.users_path_1, inputs.users_path_2)
    _progress(progress, 10, "Normalizing Valinor Attendance with Polars...")
    attendance = load_attendance(inputs.attendance_path)

    required_dates = {start + timedelta(days=i) for i in range((end - start).days + 1)}
    missing_dates = sorted(required_dates.difference(attendance.dates))
    if missing_dates:
        preview = ", ".join(d.strftime("%d-%b") for d in missing_dates[:5])
        if len(missing_dates) > 5:
            preview += f" (+{len(missing_dates) - 5} more)"
        raise ReportError(
            f"Attendance does not contain every selected report date. Missing: {preview}. "
            f"Attendance currently runs from {min(attendance.dates):%d-%b-%Y} through {max(attendance.dates):%d-%b-%Y}."
        )

    eligible_dc, eligible_xd = _eligible_sites(flex, users, attendance, start, end)
    _validate_selected(inputs.selected_dc, eligible_dc, "DC")
    _validate_selected(inputs.selected_xd, eligible_xd, "XD")
    report_dates = [start + timedelta(days=i) for i in range((end - start).days + 1)]

    _progress(progress, 14, "Building Central Flex Base caches with Polars...")
    central = _derive_central(flex, users, attendance)
    all_sites = inputs.selected_dc + inputs.selected_xd
    cache1 = _summary_cached(central, all_sites, report_dates, 1)
    cache2 = _summary_cached(central, all_sites, report_dates, 2)

    out_dir = Path(inputs.output_folder)
    out_dir.mkdir(parents=True, exist_ok=True)
    output = out_dir / default_output_filename()
    temp = output.with_suffix(".building.xlsx")
    if temp.exists():
        temp.unlink()

    _progress(progress, 20, "Creating compact summary sheets...")
    workbook = xlsxwriter.Workbook(temp)
    workbook.set_calc_mode("auto")
    workbook.set_properties({
        "title": "Flex Adoption Report", "subject": "Flex adoption and attendance compliance",
        "comments": "Generated by Flex Automation Project",
    })
    fmts = _formats(workbook)

    for sheet_name in [
        "DC Summary - 1", "DC Summary - 2", "XD Summary - 1", "XD Summary - 2",
        "Central Flex Base", "Valinor Base", "Valinor Attendance",
    ]:
        workbook.add_worksheet(sheet_name)

    _write_summary(workbook, "DC Summary - 1", inputs.selected_dc, report_dates, 1, cache1, fmts)
    _write_summary(workbook, "DC Summary - 2", inputs.selected_dc, report_dates, 2, cache2, fmts)
    _write_summary(workbook, "XD Summary - 1", inputs.selected_xd, report_dates, 1, cache1, fmts)
    _write_summary(workbook, "XD Summary - 2", inputs.selected_xd, report_dates, 2, cache2, fmts)

    _progress(progress, 38, "Writing combined Valinor Base...")
    user_col_map = _write_valinor_base(workbook, users, fmts, progress)
    _progress(progress, 52, "Writing normalized Valinor Attendance...")
    _write_valinor_attendance(workbook, attendance, fmts, progress)
    _progress(progress, 67, "Writing Central Flex Base with formulas...")
    _write_central(workbook, central, users, user_col_map, attendance, fmts, progress)

    _progress(progress, 82, "Finalizing formula workbook...")
    workbook.close()
    _progress(progress, 84, "Streaming large source sheets...")
    _inject_large_base_sheets(temp, users, attendance, progress)
    if output.exists():
        output.unlink()
    temp.replace(output)
    _progress(progress, 100, f"Report generated: {output}")
    return str(output)

