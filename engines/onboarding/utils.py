"""Utility helpers for the Onboarding Report generator."""

from copy import copy
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


class ReportError(Exception):
    """User-facing error raised for validation or workbook processing failures."""


def normalize_header(value):
    """Return a stable representation for Excel header comparison."""
    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def parse_date_value(value):
    """Convert a common Excel/date value into a date, or return None."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    # Excel serials normally arrive as datetime when the cell has a date format,
    # but exported files occasionally contain a raw number.
    if isinstance(value, (int, float)) and value > 0:
        try:
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except (TypeError, ValueError, OverflowError):
            return None

    if not isinstance(value, str):
        return None

    text = value.strip()
    if not text or text in {"-", "--", "NA", "N/A"}:
        return None

    # ISO strings are common in website exports and may include a time portion.
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass

    formats = (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d-%b-%Y",
        "%d %b %Y",
        "%d-%B-%Y",
        "%d %B %Y",
        "%m/%d/%Y",
        "%d-%m-%Y %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def date_to_excel_datetime(value):
    """Return midnight datetime for an Excel date header/cell."""
    return datetime.combine(value, time.min)


def inclusive_date_range(first_date, latest_date):
    """Build every calendar date from first_date through latest_date."""
    if first_date > latest_date:
        raise ReportError("The selected From date cannot be after the To date.")
    days = (latest_date - first_date).days
    return [first_date + timedelta(days=offset) for offset in range(days + 1)]


def ordinal_day(day):
    """Return 1st, 2nd, 3rd, 4th ... with correct teen suffixes."""
    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return f"{day}{suffix}"


def output_filename(report_date=None):
    """Return the required daily report filename using the generation date."""
    report_date = report_date or date.today()
    month_labels = (
        "Jan", "Feb", "Mar", "Apr", "May", "June",
        "July", "Aug", "Sep", "Oct", "Nov", "Dec",
    )
    month_text = month_labels[report_date.month - 1]
    return f"Onboarding Report - {ordinal_day(report_date.day)} {month_text}.xlsx"


def snapshot_cell(cell):
    """Capture a cell's value and presentation properties for later cloning."""
    return {
        "value": cell.value,
        "style": copy(cell._style),
        "number_format": cell.number_format,
        "comment": copy(cell.comment),
        "hyperlink": copy(cell.hyperlink),
    }


def restore_cell(target, snapshot, formula_origin=None):
    """Restore a captured cell, translating formulas when an origin is supplied."""
    value = snapshot["value"]
    if formula_origin and isinstance(value, str) and value.startswith("="):
        value = Translator(value, origin=formula_origin).translate_formula(target.coordinate)

    target.value = value
    target._style = copy(snapshot["style"])
    target.number_format = snapshot["number_format"]
    target.comment = copy(snapshot["comment"])
    target._hyperlink = copy(snapshot["hyperlink"])


def snapshot_column_dimension(ws, column_index):
    """Capture column-dimension properties that affect appearance/layout."""
    dim = ws.column_dimensions[get_column_letter(column_index)]
    return {
        "width": dim.width,
        "hidden": bool(dim.hidden),
        "bestFit": bool(dim.bestFit),
        "outlineLevel": dim.outlineLevel,
        "collapsed": bool(dim.collapsed),
    }


def restore_column_dimension(ws, column_index, snapshot, hidden_override=None):
    """Apply captured column width/layout while always keeping the column visible."""
    dim = ws.column_dimensions[get_column_letter(column_index)]
    dim.width = snapshot["width"]
    dim.hidden = False
    dim.bestFit = snapshot["bestFit"]
    dim.outlineLevel = snapshot["outlineLevel"]
    dim.collapsed = False


def clean_path(value):
    """Normalize a GUI-provided path."""
    return Path(value).expanduser().resolve()
