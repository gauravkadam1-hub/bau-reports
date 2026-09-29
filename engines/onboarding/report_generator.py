"""Core workbook automation for the daily Onboarding Report."""

from copy import copy
import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import (
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)
from openpyxl.utils.cell import coordinate_from_string

from utils import (
    ReportError,
    date_to_excel_datetime,
    inclusive_date_range,
    normalize_header,
    output_filename,
    parse_date_value,
    restore_cell,
    restore_column_dimension,
    snapshot_cell,
    snapshot_column_dimension,
)


BASE_SHEET = "Onboarding Base"
AADHAR_DC = "Aadhar - DC"
AADHAR_XD = "Aadhar - XD"
HYPERVERGE_DC = "Hyperverge - DC"
HYPERVERGE_XD = "Hyperverge - XD"
MANUAL_HR = "Manual Onb - HR wise details"
RED_HR = "RED Cases - HR wise details"

REQUIRED_TEMPLATE_SHEETS = (
    AADHAR_DC,
    AADHAR_XD,
    HYPERVERGE_DC,
    HYPERVERGE_XD,
    MANUAL_HR,
    RED_HR,
    BASE_SHEET,
)


def _facility_type(value):
    """Return DC/XD only for facilities with a separate trailing suffix."""
    if value is None:
        return None
    text = " ".join(str(value).strip().split())
    folded = text.casefold()
    if folded.endswith(" dc"):
        return "DC"
    if folded.endswith(" xd"):
        return "XD"
    return None


def _clean_facility_name(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def _notify(callback, percent, message):
    if callback:
        callback(int(percent), message)


def _load_workbook(path, label):
    try:
        return load_workbook(path, data_only=False, keep_links=True)
    except FileNotFoundError as exc:
        raise ReportError(f"{label} file was not found: {path}") from exc
    except PermissionError as exc:
        raise ReportError(
            f"{label} file is open or not readable. Close it in Excel and try again."
        ) from exc
    except Exception as exc:
        raise ReportError(
            f"Could not open the {label.lower()} workbook. It may be corrupted or not a valid Excel file.\n\n{exc}"
        ) from exc


def _load_csv_workbook(path):
    """Load a CSV export into an in-memory openpyxl workbook."""
    encodings = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
    last_error = None

    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(65536)
                handle.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel

                wb = Workbook()
                ws = wb.active
                ws.title = BASE_SHEET
                for row in csv.reader(handle, dialect):
                    ws.append(row)
                return wb
        except UnicodeDecodeError as exc:
            last_error = exc
        except PermissionError as exc:
            raise ReportError(
                "The selected CSV file is open or cannot be read. Close it and try again."
            ) from exc
        except Exception as exc:
            last_error = exc
            break

    raise ReportError(f"Could not read the selected CSV file.\n\n{last_error}")


def _load_input_workbook(path):
    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm"}:
        return _load_workbook(path, "Uploaded")
    if suffix == ".csv":
        return _load_csv_workbook(path)
    if suffix == ".xls":
        raise ReportError(
            "Old .xls files are not supported by openpyxl. Please save/export the file as .xlsx or .csv and try again."
        )
    raise ReportError("Please select an .xlsx, .xlsm, or .csv onboarding export.")

def _template_headers(template_ws):
    headers = [template_ws.cell(1, col).value for col in range(1, template_ws.max_column + 1)]
    if not headers or any(value is None for value in headers):
        raise ReportError("The master template has missing Onboarding Base headers.")
    return headers


def _header_map(ws):
    result = {}
    duplicates = []
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(1, col).value
        key = normalize_header(raw)
        if not key:
            continue
        if key in result:
            duplicates.append(str(raw))
        else:
            result[key] = col
    if duplicates:
        raise ReportError(
            "The uploaded base contains duplicate headers: " + ", ".join(sorted(set(duplicates)))
        )
    return result


def _find_input_base_sheet(input_wb, expected_headers):
    """Find the uploaded data sheet by name or by the exact required header set."""
    expected_keys = [normalize_header(value) for value in expected_headers]

    candidates = []
    if BASE_SHEET in input_wb.sheetnames:
        candidates.append(input_wb[BASE_SHEET])
    candidates.extend(ws for ws in input_wb.worksheets if ws.title != BASE_SHEET)

    best_missing = None
    for ws in candidates:
        mapping = _header_map(ws)
        missing = [
            expected_headers[index]
            for index, key in enumerate(expected_keys)
            if key not in mapping
        ]
        if not missing:
            return ws, mapping
        if best_missing is None or len(missing) < len(best_missing):
            best_missing = missing

    missing_text = ", ".join(str(value) for value in (best_missing or expected_headers))
    raise ReportError(
        "The uploaded Excel does not contain a valid Onboarding Base header row.\n\n"
        f"Missing required headers: {missing_text}"
    )


def _validate_template(template_wb):
    missing = [name for name in REQUIRED_TEMPLATE_SHEETS if name not in template_wb.sheetnames]
    if missing:
        raise ReportError(
            "The master template is missing required sheet(s): " + ", ".join(missing)
        )

    base_ws = template_wb[BASE_SHEET]
    required = (
        "Registered Facility",
        "Onboarded Date",
        "Approved By",
        "Onboarded by Aadhaar (Y/N)",
        "Hyperverge Status",
    )
    mapping = _header_map(base_ws)
    missing_headers = [header for header in required if normalize_header(header) not in mapping]
    if missing_headers:
        raise ReportError(
            "The master template has invalid Onboarding Base headers: "
            + ", ".join(missing_headers)
        )


def _copy_base_row_style(ws, source_row, target_row, max_column):
    """Use the template's first data-row styling for rows beyond its current used range."""
    for col in range(1, max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        target._style = copy(source._style)
        target.number_format = source.number_format
    if ws.row_dimensions[source_row].height is not None:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    # User requested no automatically hidden rows.
    ws.row_dimensions[target_row].hidden = False

def _replace_onboarding_base(
    target_ws,
    source_ws,
    source_header_map,
    expected_headers,
    selected_start_date,
    selected_end_date,
    progress_callback=None,
):
    """Replace template data with only rows inside the selected onboarding date range."""
    target_header_map = {
        normalize_header(value): index + 1 for index, value in enumerate(expected_headers)
    }
    source_rows = max(0, source_ws.max_row - 1)
    original_max_row = target_ws.max_row
    target_max_col = len(expected_headers)
    style_row = 2 if original_max_row >= 2 else 1

    onboarded_key = normalize_header("Onboarded Date")
    approved_key = normalize_header("Approved By")
    aadhar_key = normalize_header("Onboarded by Aadhaar (Y/N)")
    hv_status_key = normalize_header("Hyperverge Status")

    manual_names = set()
    red_names = set()

    # Clear only rows that actually exist in the lightweight template. Do not
    # pre-create thousands of blank cells based on the uploaded source size.
    for row in range(2, original_max_row + 1):
        for col in range(1, target_max_col + 1):
            target_ws.cell(row, col).value = None
        target_ws.row_dimensions[row].hidden = False

    target_row = 2
    copied_rows = 0
    date_source_col = source_header_map[onboarded_key]

    for source_row in range(2, source_ws.max_row + 1):
        parsed_date = parse_date_value(source_ws.cell(source_row, date_source_col).value)
        if parsed_date is None or not (selected_start_date <= parsed_date <= selected_end_date):
            if source_rows and source_row % 500 == 0:
                fraction = (source_row - 1) / source_rows
                _notify(progress_callback, 12 + fraction * 23, "Filtering Onboarding Base by selected dates...")
            continue

        if target_row > original_max_row:
            _copy_base_row_style(target_ws, style_row, target_row, target_max_col)

        row_values = {}
        for header in expected_headers:
            key = normalize_header(header)
            source_col = source_header_map[key]
            target_col = target_header_map[key]
            value = source_ws.cell(source_row, source_col).value

            if key == onboarded_key:
                value = date_to_excel_datetime(parsed_date)

            target_ws.cell(target_row, target_col).value = value
            row_values[key] = value

        approved = row_values.get(approved_key)
        approved_text = "" if approved is None else str(approved).strip()
        aadhar_value = row_values.get(aadhar_key)
        hv_value = row_values.get(hv_status_key)

        if approved_text:
            if str(aadhar_value).strip().upper() == "N":
                manual_names.add(approved_text)
            if str(hv_value).strip().casefold() == "red":
                red_names.add(approved_text)

        copied_rows += 1
        target_row += 1

        if source_rows and source_row % 500 == 0:
            fraction = (source_row - 1) / source_rows
            _notify(progress_callback, 12 + fraction * 23, "Filtering Onboarding Base by selected dates...")

    return manual_names, red_names, copied_rows

def _collect_onboarding_dates(source_ws, source_header_map):
    """Return all valid onboarding dates found in the uploaded base."""
    key = normalize_header("Onboarded Date")
    if key not in source_header_map:
        raise ReportError("The uploaded base is missing the 'Onboarded Date' column.")

    source_col = source_header_map[key]
    dates = []
    for row in range(2, source_ws.max_row + 1):
        parsed = parse_date_value(source_ws.cell(row, source_col).value)
        if parsed is not None:
            dates.append(parsed)
    if not dates:
        raise ReportError("No valid dates were found under 'Onboarded Date' in the uploaded base.")
    return dates


def inspect_input_date_range(input_path, template_path=None):
    """Return the first/latest valid onboarding dates for GUI date selection."""
    input_path = Path(input_path).expanduser().resolve()
    if not input_path.exists():
        raise ReportError("Please select an existing onboarding export.")

    if template_path is None:
        template_path = Path(__file__).with_name("master_template.xlsx")
    template_path = Path(template_path).expanduser().resolve()

    template_wb = _load_workbook(template_path, "Master template")
    try:
        _validate_template(template_wb)
        expected_headers = _template_headers(template_wb[BASE_SHEET])
    finally:
        template_wb.close()

    input_wb = _load_input_workbook(input_path)
    try:
        source_ws, source_header_map = _find_input_base_sheet(input_wb, expected_headers)
        dates = _collect_onboarding_dates(source_ws, source_header_map)
        return min(dates), max(dates)
    finally:
        input_wb.close()


def _template_facility_order(wb, facility_kind):
    """Return stable unique facility order, correcting any old DC/XD misplacement."""
    if facility_kind == "DC":
        sheet_order = (AADHAR_DC, HYPERVERGE_DC, AADHAR_XD, HYPERVERGE_XD)
    else:
        sheet_order = (AADHAR_XD, HYPERVERGE_XD, AADHAR_DC, HYPERVERGE_DC)

    result = []
    seen = set()
    for sheet_name in sheet_order:
        ws = wb[sheet_name]
        total_row = _find_grand_total_row(ws)
        for row in range(3, total_row):
            facility = _clean_facility_name(ws.cell(row, 1).value)
            if not facility or _facility_type(facility) != facility_kind:
                continue
            key = facility.casefold()
            if key not in seen:
                result.append(facility)
                seen.add(key)
    return result


def _collect_source_facilities(source_ws, source_header_map, start_date, end_date):
    facility_key = normalize_header("Registered Facility")
    date_key = normalize_header("Onboarded Date")
    if facility_key not in source_header_map:
        raise ReportError("The uploaded base is missing the 'Registered Facility' column.")
    if date_key not in source_header_map:
        raise ReportError("The uploaded base is missing the 'Onboarded Date' column.")

    facility_col = source_header_map[facility_key]
    date_col = source_header_map[date_key]
    found = {"DC": [], "XD": []}
    seen = {"DC": set(), "XD": set()}

    for row in range(2, source_ws.max_row + 1):
        onboarded = parse_date_value(source_ws.cell(row, date_col).value)
        if onboarded is None or not (start_date <= onboarded <= end_date):
            continue

        facility = _clean_facility_name(source_ws.cell(row, facility_col).value)
        kind = _facility_type(facility)
        if kind is None:
            continue
        key = facility.casefold()
        if key not in seen[kind]:
            found[kind].append(facility)
            seen[kind].add(key)
    return found


def inspect_facilities(
    input_path,
    selected_start_date=None,
    selected_end_date=None,
    template_path=None,
):
    """Return all eligible DC/XD facilities plus the template defaults.

    Existing master-template facilities are the default selected set. Facilities
    found only in the uploaded source are returned as available but not as
    defaults, allowing the Streamlit layer to mark them NEW and keep them
    unchecked until the user explicitly selects them.
    """
    input_path = Path(input_path).expanduser().resolve()
    if not input_path.exists():
        raise ReportError("Please select an existing onboarding export.")

    if template_path is None:
        template_path = Path(__file__).with_name("master_template.xlsx")
    template_path = Path(template_path).expanduser().resolve()

    template_wb = _load_workbook(template_path, "Master template")
    input_wb = None
    try:
        _validate_template(template_wb)
        expected_headers = _template_headers(template_wb[BASE_SHEET])
        defaults = {
            "DC": _template_facility_order(template_wb, "DC"),
            "XD": _template_facility_order(template_wb, "XD"),
        }

        input_wb = _load_input_workbook(input_path)
        source_ws, source_header_map = _find_input_base_sheet(input_wb, expected_headers)
        dates = _collect_onboarding_dates(source_ws, source_header_map)
        start_date = selected_start_date or min(dates)
        end_date = selected_end_date or max(dates)
        if isinstance(start_date, datetime):
            start_date = start_date.date()
        if isinstance(end_date, datetime):
            end_date = end_date.date()
        if start_date > end_date:
            raise ReportError("The selected From date cannot be after the To date.")

        source = _collect_source_facilities(
            source_ws, source_header_map, start_date, end_date
        )

        available = {}
        for kind in ("DC", "XD"):
            ordered = list(defaults[kind])
            seen = {item.casefold() for item in ordered}
            for facility in source[kind]:
                key = facility.casefold()
                if key not in seen:
                    ordered.append(facility)
                    seen.add(key)
            available[kind] = ordered

        return {
            "DC": available["DC"],
            "XD": available["XD"],
            "default_DC": defaults["DC"],
            "default_XD": defaults["XD"],
        }
    finally:
        if input_wb is not None:
            input_wb.close()
        template_wb.close()


def inspect_new_facilities(
    input_path,
    selected_start_date,
    selected_end_date,
    template_path=None,
):
    """Return new spaced-suffix DC/XD facilities found in the selected report range."""
    input_path = Path(input_path).expanduser().resolve()
    if not input_path.exists():
        raise ReportError("Please select an existing onboarding export.")

    if template_path is None:
        template_path = Path(__file__).with_name("master_template.xlsx")
    template_path = Path(template_path).expanduser().resolve()

    if selected_start_date > selected_end_date:
        raise ReportError("The selected From date cannot be after the To date.")

    template_wb = _load_workbook(template_path, "Master template")
    input_wb = None
    try:
        _validate_template(template_wb)
        expected_headers = _template_headers(template_wb[BASE_SHEET])
        known = {
            "DC": {item.casefold() for item in _template_facility_order(template_wb, "DC")},
            "XD": {item.casefold() for item in _template_facility_order(template_wb, "XD")},
        }

        input_wb = _load_input_workbook(input_path)
        source_ws, source_header_map = _find_input_base_sheet(input_wb, expected_headers)
        source_facilities = _collect_source_facilities(
            source_ws,
            source_header_map,
            selected_start_date,
            selected_end_date,
        )

        return {
            kind: sorted(
                [item for item in source_facilities[kind] if item.casefold() not in known[kind]],
                key=str.casefold,
            )
            for kind in ("DC", "XD")
        }
    finally:
        if input_wb is not None:
            input_wb.close()
        template_wb.close()


def _restore_total_row(ws, total_snapshot, old_total_row, new_total_row):
    """Restore Grand Total styling and formulas after facility row count changes."""
    data_end = new_total_row - 1
    for col, item in enumerate(total_snapshot, start=1):
        target = ws.cell(new_total_row, col)
        restore_cell(target, item)
        value = item["value"]
        if isinstance(value, str) and value.startswith("="):
            if value.upper().startswith("=SUM("):
                letter = get_column_letter(col)
                target.value = f"=SUM({letter}3:{letter}{data_end})"
            else:
                origin = f"{get_column_letter(col)}{old_total_row}"
                target.value = Translator(value, origin=origin).translate_formula(target.coordinate)


def _rebuild_facility_rows(ws, facilities):
    """Make one summary sheet use exactly the requested facility list."""
    old_total_row = _find_grand_total_row(ws)
    if old_total_row <= 3:
        raise ReportError(f"'{ws.title}' does not contain a facility template row.")

    max_col = ws.max_column
    prototype_row = 3
    prototype = [snapshot_cell(ws.cell(prototype_row, col)) for col in range(1, max_col + 1)]
    total_snapshot = [snapshot_cell(ws.cell(old_total_row, col)) for col in range(1, max_col + 1)]
    prototype_height = ws.row_dimensions[prototype_row].height
    total_height = ws.row_dimensions[old_total_row].height

    current_count = old_total_row - 3
    target_count = len(facilities)
    if target_count > current_count:
        ws.insert_rows(old_total_row, amount=target_count - current_count)
    elif target_count < current_count:
        ws.delete_rows(3 + target_count, amount=current_count - target_count)

    new_total_row = 3 + target_count
    for index, facility in enumerate(facilities, start=3):
        for col, item in enumerate(prototype, start=1):
            origin = f"{get_column_letter(col)}{prototype_row}"
            restore_cell(ws.cell(index, col), item, formula_origin=origin)
        ws.cell(index, 1).value = facility
        if prototype_height is not None:
            ws.row_dimensions[index].height = prototype_height
        ws.row_dimensions[index].hidden = False

    _restore_total_row(ws, total_snapshot, old_total_row, new_total_row)
    if total_height is not None:
        ws.row_dimensions[new_total_row].height = total_height
    ws.row_dimensions[new_total_row].hidden = False


def _sync_master_facility_rows(wb, dc_facilities, xd_facilities):
    """Keep Aadhar and Hyperverge site rows aligned by DC/XD suffix."""
    for sheet_name in (AADHAR_DC, HYPERVERGE_DC):
        _rebuild_facility_rows(wb[sheet_name], dc_facilities)
    for sheet_name in (AADHAR_XD, HYPERVERGE_XD):
        _rebuild_facility_rows(wb[sheet_name], xd_facilities)


def _latest_date_sort_counts(base_ws, report_date):
    """Return D-1 Hyperverge RED and Manual counts keyed by facility.

    The user wants one shared row order for each Aadhar/Hyperverge DC/XD pair:
    Hyperverge RED count descending first, then Manual count descending.  The
    counts are calculated from the filtered Onboarding Base so sorting does not
    depend on Excel recalculating formulas after the workbook is opened.
    """
    mapping = _header_map(base_ws)
    required = {
        "facility": normalize_header("Registered Facility"),
        "date": normalize_header("Onboarded Date"),
        "aadhar": normalize_header("Onboarded by Aadhaar (Y/N)"),
        "hyperverge": normalize_header("Hyperverge Status"),
    }
    missing = [label for label, key in required.items() if key not in mapping]
    if missing:
        raise ReportError(
            "The Onboarding Base is missing column(s) needed for D-1 sorting: "
            + ", ".join(missing)
        )

    facility_col = mapping[required["facility"]]
    date_col = mapping[required["date"]]
    aadhar_col = mapping[required["aadhar"]]
    hyperverge_col = mapping[required["hyperverge"]]

    counts = {}
    for row in range(2, base_ws.max_row + 1):
        onboarded = parse_date_value(base_ws.cell(row, date_col).value)
        if onboarded != report_date:
            continue

        facility = _clean_facility_name(base_ws.cell(row, facility_col).value)
        if not facility:
            continue
        key = facility.casefold()
        item = counts.setdefault(key, {"red": 0, "manual": 0})

        hv_status = str(base_ws.cell(row, hyperverge_col).value or "").strip().casefold()
        if hv_status == "red":
            item["red"] += 1

        aadhar_flag = str(base_ws.cell(row, aadhar_col).value or "").strip().casefold()
        if aadhar_flag == "n":
            item["manual"] += 1

    return counts


def _sort_facilities_for_latest_date(
    wb, base_ws, facility_kind, report_date, facilities=None
):
    """Stable-sort sites by D-1 RED desc, then D-1 Manual desc.

    When ``facilities`` is supplied, only that explicit user-selected set is
    included. Otherwise the current master-template order is used.
    """
    if facilities is None:
        facilities = _template_facility_order(wb, facility_kind)
    else:
        cleaned = []
        seen = set()
        for raw in facilities:
            facility = _clean_facility_name(raw)
            if _facility_type(facility) != facility_kind:
                continue
            key = facility.casefold()
            if key not in seen:
                cleaned.append(facility)
                seen.add(key)
        facilities = cleaned

    counts = _latest_date_sort_counts(base_ws, report_date)

    # Python's sort is stable, so sites tied on both metrics keep the user's
    # selected/master order. Grand Total is rebuilt separately and never sorted.
    return sorted(
        facilities,
        key=lambda facility: (
            -counts.get(facility.casefold(), {}).get("red", 0),
            -counts.get(facility.casefold(), {}).get("manual", 0),
        ),
    )


def add_new_facilities_to_master(facilities, template_path=None):
    """Persist user-approved new facilities into the lightweight master template."""
    if template_path is None:
        template_path = Path(__file__).with_name("master_template.xlsx")
    template_path = Path(template_path).expanduser().resolve()

    wb = _load_workbook(template_path, "Master template")
    temp_path = template_path.with_name(template_path.stem + ".updating.xlsx")
    try:
        _validate_template(wb)
        dc = _template_facility_order(wb, "DC")
        xd = _template_facility_order(wb, "XD")

        for kind, target in (("DC", dc), ("XD", xd)):
            seen = {item.casefold() for item in target}
            for raw in facilities.get(kind, []):
                facility = _clean_facility_name(raw)
                if _facility_type(facility) != kind:
                    continue
                key = facility.casefold()
                if key not in seen:
                    target.append(facility)
                    seen.add(key)

        _sync_master_facility_rows(wb, dc, xd)
        _remove_freeze_and_hidden_state(wb)
        formula_errors = _scan_formula_errors(wb)
        if formula_errors:
            raise ReportError(
                "The master-template facility update created broken formula references: "
                + ", ".join(formula_errors)
            )
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcOnSave = True

        try:
            wb.save(temp_path)
        except PermissionError as exc:
            raise ReportError(
                "The master template cannot be updated because it is open in Excel or the app folder is not writable."
            ) from exc
    finally:
        wb.close()

    try:
        temp_path.replace(template_path)
    except PermissionError as exc:
        if temp_path.exists():
            temp_path.unlink()
        raise ReportError(
            "The master template could not be replaced. Close it in Excel and make sure the app folder is writable."
        ) from exc
    except Exception:
        if temp_path.exists():
            temp_path.unlink()
        raise


def _remove_freeze_and_hidden_state(wb):
    """Leave every row/column visible and remove freeze panes, per user preference."""
    for ws in wb.worksheets:
        ws.freeze_panes = None
        for dimension in ws.column_dimensions.values():
            dimension.hidden = False
            dimension.collapsed = False
        for dimension in ws.row_dimensions.values():
            dimension.hidden = False
            dimension.collapsed = False


def _hide_older_report_date_blocks(ws, report_dates, block_width):
    """Hide all generated date blocks except the latest two report dates."""
    keep_from = max(0, len(report_dates) - 2)
    for index, _group_date in enumerate(report_dates):
        hidden = index < keep_from
        start_col = 2 + index * block_width
        for col in range(start_col, start_col + block_width):
            dimension = ws.column_dimensions[get_column_letter(col)]
            dimension.hidden = hidden
            dimension.collapsed = False


def _excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _date_groups(ws):
    """Read merged date headers as (date, start_col, width) tuples."""
    groups = []
    row_one_merges = [
        merged
        for merged in ws.merged_cells.ranges
        if merged.min_row == 1 and merged.max_row == 1
    ]
    for merged in sorted(row_one_merges, key=lambda item: item.min_col):
        value = ws.cell(1, merged.min_col).value
        parsed = _excel_date(value)
        if parsed is not None:
            groups.append((parsed, merged.min_col, merged.max_col - merged.min_col + 1))
    return groups


def _find_formula_template_block(ws, width):
    """Find a merged report block by its metric headers when master dates are blank."""
    candidates = []
    for merged in ws.merged_cells.ranges:
        if merged.min_row != 1 or merged.max_row != 1:
            continue
        if merged.max_col - merged.min_col + 1 != width:
            continue
        if merged.min_col < 2:
            continue

        first_metric = normalize_header(ws.cell(2, merged.min_col).value)
        if first_metric != normalize_header("Total Onboarding"):
            continue
        candidates.append(merged.min_col)

    if not candidates:
        return None
    return max(candidates)


def _find_grand_total_row(ws):
    for row in range(3, ws.max_row + 1):
        if str(ws.cell(row, 1).value).strip().casefold() == "grand total":
            return row
    raise ReportError(f"Could not find the 'Grand Total' row in '{ws.title}'.")


def _snapshot_block(ws, start_col, width):
    cells = []
    for row in range(1, ws.max_row + 1):
        row_cells = []
        for offset in range(width):
            col = start_col + offset
            row_cells.append(
                {
                    "origin": f"{get_column_letter(col)}{row}",
                    "cell": snapshot_cell(ws.cell(row, col)),
                }
            )
        cells.append(row_cells)
    return {
        "start_col": start_col,
        "width": width,
        "cells": cells,
        "dimensions": [snapshot_column_dimension(ws, start_col + offset) for offset in range(width)],
    }


def _find_cf_rules(ws, exact_range):
    """Capture rules for the template metric column, even if site rows changed."""
    result = []
    min_col, min_row, max_col, max_row = range_boundaries(exact_range)
    for conditional, rules in ws.conditional_formatting._cf_rules.items():
        if str(conditional.sqref) == exact_range:
            result.extend(copy(rule) for rule in rules)
            continue

        # Facility rows can be appended to the persistent master. openpyxl does
        # not always expand conditional-formatting sqref ranges on row inserts,
        # so match the same metric column and first data row as a safe fallback.
        try:
            ranges = list(conditional.sqref.ranges)
        except AttributeError:
            ranges = []
        for cell_range in ranges:
            if (
                cell_range.min_col == min_col
                and cell_range.max_col == max_col
                and cell_range.min_row <= min_row <= cell_range.max_row
            ):
                result.extend(copy(rule) for rule in rules)
                break
    return result


def _clear_dynamic_columns(ws):
    """Remove all generated columns from B onward while preserving column A and sheet settings."""
    for merged in list(ws.merged_cells.ranges):
        if merged.max_col >= 2:
            ws.unmerge_cells(str(merged))

    old_max_col = ws.max_column
    if old_max_col >= 2:
        ws.delete_cols(2, old_max_col - 1)

    # delete_cols does not reliably remove stale column-dimension records.
    for key in list(ws.column_dimensions.keys()):
        try:
            index = column_index_from_string(key)
        except (TypeError, ValueError):
            continue
        if index >= 2:
            del ws.column_dimensions[key]

    ws.conditional_formatting._cf_rules.clear()


def _paste_block(ws, block, target_start_col):
    for row_index, row_cells in enumerate(block["cells"], start=1):
        for offset, item in enumerate(row_cells):
            target = ws.cell(row_index, target_start_col + offset)
            restore_cell(target, item["cell"], formula_origin=item["origin"])


def _apply_block_dimensions(ws, block, target_start_col, hidden_flags=None):
    """Restore widths/layout but always keep generated columns visible."""
    for offset, dimension in enumerate(block["dimensions"]):
        restore_column_dimension(
            ws,
            target_start_col + offset,
            dimension,
            hidden_override=False,
        )

def _rebuild_hyperverge_sheet(ws, report_dates):
    groups = _date_groups(ws)
    if groups:
        # The most recent block is the canonical two-column template.
        _, source_start, source_width = groups[-1]
    else:
        source_start = _find_formula_template_block(ws, 2)
        source_width = 2 if source_start is not None else None
        if source_start is None:
            raise ReportError(f"No formula template block was found in '{ws.title}'.")
    if source_width != 2:
        raise ReportError(f"Unexpected date-block width in '{ws.title}'.")

    total_row = _find_grand_total_row(ws)
    data_end = total_row - 1
    source_block = _snapshot_block(ws, source_start, source_width)
    source_red_col = source_start + 1
    if data_end >= 3:
        red_range = f"{get_column_letter(source_red_col)}3:{get_column_letter(source_red_col)}{data_end}"
        red_cf_rules = _find_cf_rules(ws, red_range)
    else:
        red_cf_rules = []

    _clear_dynamic_columns(ws)

    red_columns = {}
    target_start = 2
    for group_date in report_dates:
        _paste_block(ws, source_block, target_start)
        ws.cell(1, target_start).value = date_to_excel_datetime(group_date)
        ws.merge_cells(start_row=1, start_column=target_start, end_row=1, end_column=target_start + 1)

        _apply_block_dimensions(ws, source_block, target_start)

        red_col = target_start + 1
        red_columns[group_date] = red_col
        if data_end >= 3:
            red_target_range = f"{get_column_letter(red_col)}3:{get_column_letter(red_col)}{data_end}"
            for rule in red_cf_rules:
                ws.conditional_formatting.add(red_target_range, copy(rule))

        target_start += 2

    return red_columns

def _find_hv_summary_start(ws):
    for col in range(2, ws.max_column + 1):
        value = ws.cell(1, col).value
        if isinstance(value, str) and value.strip().casefold() == "hyperverge red":
            return col
    raise ReportError(f"Could not find the Hyperverge Red summary in '{ws.title}'.")


def _external_row_from_formula(formula, expected_sheet):
    """Read only the row mapping from a template cross-sheet formula."""
    if not isinstance(formula, str) or not formula.startswith("=") or "!" not in formula:
        return None
    left, right = formula[1:].split("!", 1)
    sheet_name = left.strip().strip("'")
    if sheet_name.casefold() != expected_sheet.casefold():
        return None
    reference = right.strip().replace("$", "")
    try:
        _, row = coordinate_from_string(reference)
    except ValueError:
        return None
    return row


def _rebuild_aadhar_sheet(ws, report_dates, hyper_sheet_name, red_columns):
    groups = _date_groups(ws)
    # Latest Aadhar date blocks in the supplied template contain the complete
    # four-metric layout: Total Onboarding, Aadhar, Manual, Compliance %.
    canonical = next((item for item in reversed(groups) if item[2] == 4), None)
    if canonical is not None:
        _, source_start, source_width = canonical
    else:
        source_start = _find_formula_template_block(ws, 4)
        source_width = 4 if source_start is not None else None
        if source_start is None:
            raise ReportError(f"Could not find a four-column formula template in '{ws.title}'.")

    total_row = _find_grand_total_row(ws)
    data_end = total_row - 1
    source_block = _snapshot_block(ws, source_start, source_width)
    manual_col = source_start + 2
    compliance_col = source_start + 3
    if data_end >= 3:
        manual_range = f"{get_column_letter(manual_col)}3:{get_column_letter(manual_col)}{data_end}"
        compliance_range = f"{get_column_letter(compliance_col)}3:{get_column_letter(compliance_col)}{data_end}"
        manual_cf_rules = _find_cf_rules(ws, manual_range)
        compliance_cf_rules = _find_cf_rules(ws, compliance_range)
    else:
        manual_cf_rules = []
        compliance_cf_rules = []

    hv_source_start = _find_hv_summary_start(ws)
    hv_source_block = _snapshot_block(ws, hv_source_start, 2)

    _clear_dynamic_columns(ws)

    date_start_columns = {}
    target_start = 2
    for group_date in report_dates:
        _paste_block(ws, source_block, target_start)
        ws.cell(1, target_start).value = date_to_excel_datetime(group_date)
        ws.merge_cells(
            start_row=1,
            start_column=target_start,
            end_row=1,
            end_column=target_start + 3,
        )

        _apply_block_dimensions(ws, source_block, target_start)
        date_start_columns[group_date] = target_start

        target_manual_col = target_start + 2
        target_compliance_col = target_start + 3
        if data_end >= 3:
            target_manual_range = (
                f"{get_column_letter(target_manual_col)}3:"
                f"{get_column_letter(target_manual_col)}{data_end}"
            )
            target_compliance_range = (
                f"{get_column_letter(target_compliance_col)}3:"
                f"{get_column_letter(target_compliance_col)}{data_end}"
            )
            for rule in manual_cf_rules:
                ws.conditional_formatting.add(target_manual_range, copy(rule))
            for rule in compliance_cf_rules:
                ws.conditional_formatting.add(target_compliance_range, copy(rule))

        target_start += 4

    # Hyperverge RED summary: only the latest two available report dates.
    hv_dates = report_dates[-2:]
    hv_width = len(hv_dates)
    hv_target_start = target_start

    for offset in range(hv_width):
        # Clone styles/formulas first, then replace only the date/link formulas.
        for row in range(1, ws.max_row + 1):
            item = hv_source_block["cells"][row - 1][offset]
            # Clone presentation first. Cross-sheet formulas in this block are
            # rebuilt below because translating them by the Aadhar column offset
            # would shift Hyperverge references by the wrong number of columns.
            restore_cell(
                ws.cell(row, hv_target_start + offset),
                item["cell"],
            )
        restore_column_dimension(
            ws,
            hv_target_start + offset,
            hv_source_block["dimensions"][offset],
            hidden_override=False,
        )

    if hv_width:
        ws.cell(1, hv_target_start).value = hv_source_block["cells"][0][0]["cell"]["value"]
        if hv_width == 2:
            ws.merge_cells(
                start_row=1,
                start_column=hv_target_start,
                end_row=1,
                end_column=hv_target_start + 1,
            )

    for offset, group_date in enumerate(hv_dates):
        date_header_col = date_start_columns[group_date]
        ws.cell(2, hv_target_start + offset).value = f"={get_column_letter(date_header_col)}1"

        hyper_red_col = red_columns[group_date]
        hyper_red_letter = get_column_letter(hyper_red_col)
        for row in range(3, data_end + 1):
            # Aadhar and Hyperverge facility rows are deliberately kept in the
            # same order by the master-facility synchronizer. Link by the same
            # row number instead of trusting historical cross-sheet mappings,
            # which can become stale when sites are moved between DC and XD.
            mapped_row = row
            ws.cell(row, hv_target_start + offset).value = (
                f"='{hyper_sheet_name}'!{hyper_red_letter}{mapped_row}"
            )

        # Grand Total formula remains a translated copy of the master template.
        source_total = hv_source_block["cells"][total_row - 1][offset]
        ws.cell(total_row, hv_target_start + offset).value = Translator(
            source_total["cell"]["value"], origin=source_total["origin"]
        ).translate_formula(ws.cell(total_row, hv_target_start + offset).coordinate)


def _rebuild_hr_detail_sheet(ws, relevant_names):
    """Preserve existing HR order, append new HRs, and copy the template COUNTIFS formula."""
    if ws.max_row < 2:
        raise ReportError(f"'{ws.title}' does not contain a template detail row.")

    existing_order = []
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 1).value
        if value is not None and str(value).strip():
            existing_order.append(str(value).strip())

    prototype_a = snapshot_cell(ws.cell(2, 1))
    prototype_b = snapshot_cell(ws.cell(2, 2))
    prototype_formula = ws.cell(2, 2).value
    if not isinstance(prototype_formula, str) or not prototype_formula.startswith("="):
        raise ReportError(f"'{ws.title}' is missing its template formula in B2.")

    ordered = [name for name in existing_order if name in relevant_names]
    known = set(ordered)
    ordered.extend(sorted((name for name in relevant_names if name not in known), key=str.casefold))

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row, name in enumerate(ordered, start=2):
        restore_cell(ws.cell(row, 1), prototype_a)
        restore_cell(ws.cell(row, 2), prototype_b)
        ws.cell(row, 1).value = name
        ws.cell(row, 2).value = Translator(
            prototype_formula, origin="B2"
        ).translate_formula(f"B{row}")


def _scan_formula_errors(wb):
    """Catch formulas that were broken while columns were regenerated."""
    problems = []
    for sheet_name in (AADHAR_DC, AADHAR_XD, HYPERVERGE_DC, HYPERVERGE_XD, MANUAL_HR, RED_HR):
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("=") and "#REF!" in value.upper():
                    problems.append(f"{sheet_name}!{cell.coordinate}")
                    if len(problems) >= 10:
                        return problems
    return problems


def generate_report(
    input_path,
    template_path=None,
    selected_start_date=None,
    selected_end_date=None,
    selected_dc=None,
    selected_xd=None,
    output_dir=None,
    progress_callback=None,
    generation_date=None,
):
    """
    Generate a completed report from a fresh Onboarding Base export.

    The uploaded base is filtered to the user-selected onboarding date range so
    all summaries and HR-wise counts represent exactly that same range.

    Returns:
        tuple(Path, date, date): output path, selected start date, selected end date
    """
    input_path = Path(input_path).expanduser().resolve()
    if not input_path.exists():
        raise ReportError("Please select an existing onboarding export.")

    if template_path is None:
        template_path = Path(__file__).with_name("master_template.xlsx")
    template_path = Path(template_path).expanduser().resolve()

    _notify(progress_callback, 3, "Opening master template...")
    template_wb = _load_workbook(template_path, "Master template")
    _validate_template(template_wb)

    input_wb = None
    try:
        _notify(progress_callback, 7, "Opening uploaded onboarding data...")
        input_wb = _load_input_workbook(input_path)
        expected_headers = _template_headers(template_wb[BASE_SHEET])
        source_ws, source_header_map = _find_input_base_sheet(input_wb, expected_headers)

        _notify(progress_callback, 10, "Validating date range...")
        available_dates = _collect_onboarding_dates(source_ws, source_header_map)
        available_first = min(available_dates)
        available_latest = max(available_dates)

        start_date = selected_start_date or available_first
        end_date = selected_end_date or available_latest
        if isinstance(start_date, datetime):
            start_date = start_date.date()
        if isinstance(end_date, datetime):
            end_date = end_date.date()
        if not isinstance(start_date, date) or not isinstance(end_date, date):
            raise ReportError("Please select a valid From and To date.")
        if start_date > end_date:
            raise ReportError("The selected From date cannot be after the To date.")
        if start_date < available_first or end_date > available_latest:
            raise ReportError(
                "The selected date range must stay inside the dates available in the uploaded file.\n\n"
                f"Available: {available_first.strftime('%d %b %Y')} to {available_latest.strftime('%d %b %Y')}"
            )

        report_dates = inclusive_date_range(start_date, end_date)

        _notify(progress_callback, 12, "Filtering Onboarding Base by selected dates...")
        manual_names, red_names, copied_rows = _replace_onboarding_base(
            template_wb[BASE_SHEET],
            source_ws,
            source_header_map,
            expected_headers,
            start_date,
            end_date,
            progress_callback,
        )

        # Sort each DC/XD pair once using the latest report date so Aadhar and
        # Hyperverge sheets remain row-aligned.  Priority is Hyperverge RED
        # descending, then Manual descending.  Grand Total is rebuilt separately
        # by _rebuild_facility_rows and therefore never participates in sorting.
        latest_report_date = report_dates[-1]
        sorted_dc = _sort_facilities_for_latest_date(
            template_wb,
            template_wb[BASE_SHEET],
            "DC",
            latest_report_date,
            facilities=selected_dc,
        )
        sorted_xd = _sort_facilities_for_latest_date(
            template_wb,
            template_wb[BASE_SHEET],
            "XD",
            latest_report_date,
            facilities=selected_xd,
        )
        _sync_master_facility_rows(template_wb, sorted_dc, sorted_xd)

        _notify(progress_callback, 40, "Updating Hyperverge - DC...")
        dc_red_columns = _rebuild_hyperverge_sheet(template_wb[HYPERVERGE_DC], report_dates)

        _notify(progress_callback, 52, "Updating Hyperverge - XD...")
        xd_red_columns = _rebuild_hyperverge_sheet(template_wb[HYPERVERGE_XD], report_dates)

        _notify(progress_callback, 64, "Updating Aadhar - DC summary...")
        _rebuild_aadhar_sheet(
            template_wb[AADHAR_DC], report_dates, HYPERVERGE_DC, dc_red_columns
        )

        _notify(progress_callback, 76, "Updating Aadhar - XD summary...")
        _rebuild_aadhar_sheet(
            template_wb[AADHAR_XD], report_dates, HYPERVERGE_XD, xd_red_columns
        )

        _notify(progress_callback, 85, "Updating HR-wise detail sheets...")
        _rebuild_hr_detail_sheet(template_wb[MANUAL_HR], manual_names)
        _rebuild_hr_detail_sheet(template_wb[RED_HR], red_names)

        # Keep the existing no-freeze behavior, then hide only historical date
        # blocks in the four summary sheets.  Full date-range formulas remain in
        # the workbook; only D-2 and D-1 stay visible.
        _remove_freeze_and_hidden_state(template_wb)
        for sheet_name in (HYPERVERGE_DC, HYPERVERGE_XD):
            _hide_older_report_date_blocks(
                template_wb[sheet_name], report_dates, block_width=2
            )
        for sheet_name in (AADHAR_DC, AADHAR_XD):
            _hide_older_report_date_blocks(
                template_wb[sheet_name], report_dates, block_width=4
            )

        errors = _scan_formula_errors(template_wb)
        if errors:
            raise ReportError(
                "Formula validation found broken references: " + ", ".join(errors)
            )

        # openpyxl does not evaluate formulas. Force Excel to perform a full
        # calculation when the output file is opened.
        template_wb.calculation.calcMode = "auto"
        template_wb.calculation.fullCalcOnLoad = True
        template_wb.calculation.forceFullCalc = True
        template_wb.calculation.calcOnSave = True

        if output_dir is None:
            output_dir = input_path.parent
        output_dir = Path(output_dir).expanduser().resolve()
        if not output_dir.exists() or not output_dir.is_dir():
            raise ReportError("The selected destination folder does not exist.")

        generation_date = generation_date or date.today()
        output_path = output_dir / output_filename(generation_date)
        if output_path == input_path:
            raise ReportError(
                "The output path is the same as the uploaded source file. Please choose a different destination folder."
            )

        _notify(
            progress_callback,
            94,
            f"Saving final report ({copied_rows:,} matching rows)...",
        )
        try:
            template_wb.save(output_path)
        except PermissionError as exc:
            raise ReportError(
                f"Cannot save '{output_path.name}' because it is open in Excel or the folder is not writable."
            ) from exc
        except Exception as exc:
            raise ReportError(f"Could not save the final report.\n\n{exc}") from exc

        _notify(progress_callback, 100, "Report generated successfully.")
        return output_path, start_date, end_date
    finally:
        if input_wb is not None:
            input_wb.close()
        template_wb.close()
