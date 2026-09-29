from __future__ import annotations

import csv
import json
import math
import re
import zipfile
import tempfile
import os
from bisect import bisect_right
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Iterable, Iterator

try:
    import polars as pl  # primary/production backend
except Exception:  # development fallback only; RUN_ME installs Polars
    pl = None

try:
    import pandas as pd  # optional fallback for environments where Polars is unavailable
except Exception:
    pd = None

import xlsxwriter
from openpyxl import load_workbook
from xlsxwriter.utility import xl_col_to_name

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}
ProgressCallback = Callable[[int, str], None]

ASSET_DIR = Path(__file__).resolve().parent / "assets"
CONFIG = json.loads((ASSET_DIR / "config.json").read_text(encoding="utf-8"))
DEFAULT_DC_SITES = CONFIG["default_dc_sites"]
DEFAULT_XD_SITES = CONFIG["default_xd_sites"]
MAIN_REPORT_TOKEN = "__main__"
MAIN_REPORT_LABEL = "Main Attendance"
MASTER_HC_PATH = Path(__file__).resolve().parent / "Deployed HC Master.xlsx"
REF_MASTER_PATH = Path(__file__).resolve().parent / "REF Master.xlsx"
REF_MASTER_SHEET = "Ref"
DEFAULT_ROLE_REPORTS = ["Housekeeping", "Security"]  # legacy labels only; deployed master drives dedicated roles

ATTENDANCE_SUBHEADERS = ["shift", "total_hours", "working_hours", "punch_status"]
METRIC_HEADERS = [
    "HC", "Total \nPunched", "Punched %", "Single \nPunch", "Less than \n5 Hours",
    "5 to 8 \nHours", "8 Hours +", "Double Punch\nAdherence", "<5 Hrs\nAdherence",
    "5 to 8 Hrs \nAdherence", "8 Hrs + \nAdherence",
]
METRIC_WIDTHS = [6.55, 7.78, 7.78, 5.78, 7.11, 5.66, 7.55, 9.55, 7.78, 9.11, 7.78]


class ReportError(Exception):
    pass


@dataclass
class UserData:
    frame: object
    columns: list[str]
    required: dict[str, str]


@dataclass
class AttendanceData:
    frame: object
    columns: list[str]
    identity_columns: list[str]
    identity_map: dict[str, str]
    blocks: dict[date, dict[str, str]]
    dates: list[date]


@dataclass
class ScanResult:
    dc_sites: list[str]
    xd_sites: list[str]
    roles: list[str]
    dedicated_roles: list[str]
    master_only_roles: list[str]
    user_rows: int
    attendance_rows: int
    attendance_start: date
    attendance_end: date


@dataclass(frozen=True)
class ReportSpec:
    key: str
    display_name: str
    roles: tuple[str, ...]
    deployed_sheet_name: str | None
    hc_label: str
    is_main: bool = False


@dataclass
class ReportInputs:
    users_path_1: str
    users_path_2: str
    attendance_path: str
    output_folder: str
    report_start: date
    report_end: date
    selected_dc: list[str]
    selected_xd: list[str]
    selected_reports: list[str] | None = None


def _progress(cb: ProgressCallback | None, pct: int, message: str) -> None:
    if cb:
        cb(max(0, min(100, pct)), message)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
    return str(value).strip()


def _clean_space(value: object) -> str:
    return re.sub(r"\s+", " ", _cell_text(value).strip())


def _norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", _cell_text(value).casefold())


def normalize_site(value: object) -> str:
    return _clean_space(value).casefold()


def normalize_role(value: object) -> str:
    return _clean_space(value).casefold()


def _ref_sheet_name(wb) -> str | None:
    target = _clean_space(REF_MASTER_SHEET).casefold()
    for name in wb.sheetnames:
        if _clean_space(name).casefold() == target:
            return name
    return None


def load_ref_master_roles() -> list[str]:
    """Read the editable Main Attendance REF master in workbook order.

    The master intentionally has a single role column with no mandatory header. A
    simple Role/Ref header is tolerated and ignored if the user adds one later.
    """
    if not REF_MASTER_PATH.exists():
        raise ReportError(
            f"{REF_MASTER_PATH.name} was not found in the project folder. "
            f"Create a sheet named '{REF_MASTER_SHEET}' containing the Main Attendance roles."
        )
    wb = load_workbook(REF_MASTER_PATH, read_only=True, data_only=True)
    try:
        sheet_name = _ref_sheet_name(wb)
        if not sheet_name:
            raise ReportError(
                f"{REF_MASTER_PATH.name} must contain a sheet named '{REF_MASTER_SHEET}'."
            )
        ws = wb[sheet_name]
        roles: list[str] = []
        seen: set[str] = set()
        first_nonblank = True
        header_tokens = {"role", "roles", "ref", "reference"}
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            role = _clean_space(row[0] if row else "")
            if not role:
                continue
            key = normalize_role(role)
            if first_nonblank and key in header_tokens:
                first_nonblank = False
                continue
            first_nonblank = False
            if key in seen:
                continue
            seen.add(key)
            roles.append(role)
        if not roles:
            raise ReportError(
                f"No roles were found in {REF_MASTER_PATH.name} / {REF_MASTER_SHEET}."
            )
        return roles
    finally:
        wb.close()


def dedicated_master_roles() -> list[str]:
    """Return role names represented by `<Role> - Deployed HC` sheets."""
    if not MASTER_HC_PATH.exists():
        return []
    wb = load_workbook(MASTER_HC_PATH, read_only=True, data_only=False)
    try:
        out: list[str] = []
        seen: set[str] = set()
        pat = re.compile(r"^(.*?)\s*-\s*deployed\s+hc\s*$", re.IGNORECASE)
        for sheet in wb.sheetnames:
            m = pat.match(_clean_space(sheet))
            if not m:
                continue
            role = _clean_space(m.group(1))
            key = normalize_role(role)
            if role and key not in seen:
                seen.add(key)
                out.append(role)
        return out
    finally:
        wb.close()


def effective_main_roles() -> list[str]:
    """Main REF after automatically removing all dedicated deployed-HC roles."""
    ref_roles = load_ref_master_roles()
    dedicated = {normalize_role(r) for r in dedicated_master_roles()}
    return [role for role in ref_roles if normalize_role(role) not in dedicated]


def deployed_sheet_name_for_role(role: str) -> str:
    role = _clean_space(role)
    name = f"{role} - Deployed HC"
    if len(name) > 31:
        raise ReportError(
            f"Role '{role}' is too long for the required Excel sheet name '{name}'. "
            "Excel sheet names are limited to 31 characters."
        )
    return name


def _master_sheet_lookup() -> dict[str, str]:
    if not MASTER_HC_PATH.exists():
        return {}
    wb = load_workbook(MASTER_HC_PATH, read_only=True, data_only=False)
    try:
        return {_clean_space(s).casefold(): s for s in wb.sheetnames}
    finally:
        wb.close()


def master_role_available(role: str) -> bool:
    try:
        expected = deployed_sheet_name_for_role(role)
    except ReportError:
        return False
    return _clean_space(expected).casefold() in _master_sheet_lookup()


def validate_master_role(role: str) -> str:
    expected = deployed_sheet_name_for_role(role)
    lookup = _master_sheet_lookup()
    actual = lookup.get(_clean_space(expected).casefold())
    if not MASTER_HC_PATH.exists():
        raise ReportError(
            f"Deployed HC Master.xlsx was not found in the project folder. "
            f"Create the sheet '{expected}' before selecting {role}."
        )
    if not actual:
        raise ReportError(
            f"Deployed HC sheet not found for role '{role}'.\n\n"
            f"Expected sheet: {expected}\n"
            f"Workbook: {MASTER_HC_PATH.name}"
        )
    return actual


def _find_col(columns: Iterable[object], aliases: Iterable[str], label: str, required: bool = True) -> str | None:
    targets = {_norm_header(a) for a in aliases}
    for c in columns:
        if _norm_header(c) in targets:
            return str(c)
    if required:
        raise ReportError(f"Required column not found: {label}")
    return None


def _unique_headers(values: Iterable[object]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, value in enumerate(values, start=1):
        base = _cell_text(value).lstrip("\ufeff") or f"Column {i}"
        key = base.casefold()
        seen[key] = seen.get(key, 0) + 1
        out.append(base if seen[key] == 1 else f"{base}_{seen[key]}")
    return out


def _frame_columns(df: object) -> list[str]:
    return list(df.columns)


def _frame_height(df: object) -> int:
    return int(df.height if pl is not None and isinstance(df, pl.DataFrame) else len(df))


def _frame_col(df: object, col: str) -> list[str]:
    if pl is not None and isinstance(df, pl.DataFrame):
        return [_cell_text(x) for x in df.get_column(col).to_list()]
    return [_cell_text(x) for x in df[col].tolist()]


def _frame_iter_rows(df: object, columns: list[str] | None = None) -> Iterator[tuple]:
    cols = columns or _frame_columns(df)
    if pl is not None and isinstance(df, pl.DataFrame):
        selected = df.select(cols)
        yield from selected.iter_rows(named=False)
    else:
        yield from df.loc[:, cols].itertuples(index=False, name=None)


def _make_frame(headers: list[str], rows: Iterable[Iterable[object]]) -> object:
    # Explicitly string-typing everything prevents identifier corruption and mixed-type crashes.
    if pl is not None:
        buckets = [[] for _ in headers]
        width = len(headers)
        for row in rows:
            rr = list(row)
            for i in range(width):
                buckets[i].append(_cell_text(rr[i] if i < len(rr) else ""))
        return pl.DataFrame([pl.Series(h, vals, dtype=pl.String, strict=False) for h, vals in zip(headers, buckets)])
    if pd is None:
        raise ReportError("Polars is required. Please run RUN_ME.bat to complete the one-time setup.")
    return pd.DataFrame(([ _cell_text(x) for x in row ] for row in rows), columns=headers, dtype=str).fillna("")


def _read_flat(path: str) -> object:
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ReportError(f"Unsupported file type: {ext}. Use CSV or Excel.")
    if ext == ".csv":
        if pl is not None:
            try:
                df = pl.read_csv(
                    path, infer_schema_length=0, try_parse_dates=False, ignore_errors=True,
                    truncate_ragged_lines=True, encoding="utf8-lossy",
                )
                rename = {c: _cell_text(c).lstrip("\ufeff") for c in df.columns}
                df = df.rename(rename)
                return df.with_columns([pl.col(c).cast(pl.String, strict=False).fill_null("") for c in df.columns])
            except Exception as exc:
                raise ReportError(f"Could not read {Path(path).name}: {exc}") from exc
        if pd is None:
            raise ReportError("Polars is required. Please run RUN_ME.bat to complete the one-time setup.")
        try:
            return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", low_memory=False)
        except Exception as exc:
            raise ReportError(f"Could not read {Path(path).name}: {exc}") from exc

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = _unique_headers(next(rows))
        except StopIteration:
            wb.close()
            return _make_frame([], [])
        df = _make_frame(headers, rows)
        wb.close()
        return df
    except Exception as exc:
        raise ReportError(f"Could not read {Path(path).name}: {exc}") from exc


def _align_frames(d1: object, d2: object) -> tuple[object, list[str]]:
    cols1 = _frame_columns(d1)
    cols2 = _frame_columns(d2)
    result_cols = list(cols1)
    existing = {_norm_header(c) for c in result_cols}
    for c in cols2:
        if _norm_header(c) not in existing:
            result_cols.append(c)
            existing.add(_norm_header(c))

    def rows_for(df: object):
        source = {_norm_header(c): c for c in _frame_columns(df)}
        source_cols = [source.get(_norm_header(c)) for c in result_cols]
        # iterate raw rows as dictionaries only over available columns
        if pl is not None and isinstance(df, pl.DataFrame):
            for row in df.iter_rows(named=True):
                yield tuple(_cell_text(row.get(src, "")) if src else "" for src in source_cols)
        else:
            for _, row in df.iterrows():
                yield tuple(_cell_text(row[src]) if src else "" for src in source_cols)

    if pl is not None:
        a = _make_frame(result_cols, rows_for(d1))
        b = _make_frame(result_cols, rows_for(d2))
        return pl.concat([a, b], how="vertical"), result_cols
    if pd is None:
        raise ReportError("Polars is required.")
    a = pd.DataFrame(rows_for(d1), columns=result_cols, dtype=str)
    b = pd.DataFrame(rows_for(d2), columns=result_cols, dtype=str)
    return pd.concat([a, b], ignore_index=True), result_cols


def load_users(path1: str, path2: str) -> UserData:
    d1 = _read_flat(path1)
    d2 = _read_flat(path2)
    if _frame_height(d1) == 0 and _frame_height(d2) == 0:
        raise ReportError("Both All User Data files are empty.")
    frame, columns = _align_frames(d1, d2)
    required = {
        "id": _find_col(columns, ["ID", "User ID", "user_id"], "ID"),
        "role": _find_col(columns, ["Role", "user_role"], "Role"),
        "floater": _find_col(columns, ["Floater"], "Floater", required=False),
        "status": _find_col(columns, ["Status"], "Status"),
        "facility": _find_col(columns, ["Registered Facility", "Facility Name", "facility_name"], "Registered Facility"),
        "onboarded_date": _find_col(columns, ["Onboarded Date", "Onboarding Date"], "Onboarded Date"),
    }
    return UserData(frame, columns, required)


def _parse_header_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = _cell_text(value)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_value_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = _cell_text(value)
    if not s or s in {"-", "0", "None"}:
        return None
    for fmt in (
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y",
        "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _read_attendance_csv(path: str) -> tuple[list[str], list[str], object]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        try:
            h1 = next(reader); h2 = next(reader)
        except StopIteration as exc:
            raise ReportError("Attendance Base must contain two header rows.") from exc
    if pl is not None:
        raw = pl.read_csv(
            path, has_header=False, skip_rows=2, infer_schema_length=0, try_parse_dates=False,
            ignore_errors=True, truncate_ragged_lines=True, encoding="utf8-lossy",
        )
        raw = raw.with_columns([pl.col(c).cast(pl.String, strict=False).fill_null("") for c in raw.columns])
        width = raw.width
    else:
        if pd is None:
            raise ReportError("Polars is required.")
        raw = pd.read_csv(path, header=None, skiprows=2, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        raw.columns = [f"column_{i+1}" for i in range(raw.shape[1])]
        width = raw.shape[1]
    h1 = list(h1) + [""] * max(0, width - len(h1))
    h2 = list(h2) + [""] * max(0, width - len(h2))
    return h1[:width], h2[:width], raw


def _read_attendance_excel(path: str) -> tuple[list[str], list[str], object]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        h1 = list(next(rows)); h2 = list(next(rows))
    except StopIteration as exc:
        wb.close(); raise ReportError("Attendance Base must contain two header rows.") from exc
    width = max(len(h1), len(h2))
    headers = [f"column_{i+1}" for i in range(width)]
    df = _make_frame(headers, rows)
    wb.close()
    h1 += [""] * (width-len(h1)); h2 += [""] * (width-len(h2))
    return h1, h2, df


def load_attendance(path: str) -> AttendanceData:
    ext = Path(path).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ReportError(f"Unsupported Attendance file type: {ext}.")
    h1, h2, raw = _read_attendance_csv(path) if ext == ".csv" else _read_attendance_excel(path)
    raw_cols = _frame_columns(raw)
    if _frame_height(raw) == 0:
        raise ReportError("Attendance Base is empty.")

    aliases = {
        "user_id": ["user_id", "User ID", "ID", "Employee Code"],
        "user_name": ["user_name", "User Name", "Name"],
        "role": ["role", "Role"],
        "facility_name": ["facility_name", "Facility Name", "Registered Facility", "Facility"],
        "stint_id": ["stint_id", "Stint ID"],
        "last_present_day": ["last_present_day", "Last Present Day"],
    }
    identity_map: dict[str, str] = {}
    recognized_nondate_indices: set[int] = set()
    for desired, names in aliases.items():
        targets = {_norm_header(x) for x in names}
        found_idx = None
        for i, top in enumerate(h1):
            if _norm_header(top) in targets:
                found_idx = i; break
        if found_idx is None:
            if desired in {"user_id", "role", "facility_name"}:
                raise ReportError(f"Required Attendance column not found: {desired}")
            continue
        identity_map[desired] = raw_cols[found_idx]
        recognized_nondate_indices.add(found_idx)

    blocks_idx: dict[date, dict[str, int]] = {}
    current_date: date | None = None
    date_col_indices: set[int] = set()
    sub_aliases = {
        "shift": {"shift"},
        "total_hours": {"totalhours", "totalhour", "totalhrs"},
        "working_hours": {"workinghours", "workinghour", "workinghrs"},
        "punch_status": {"punchstatus", "punch", "status"},
    }
    for i, (top, sub) in enumerate(zip(h1, h2)):
        d = _parse_header_date(top)
        if d:
            current_date = d
        sn = _norm_header(sub)
        canonical = next((name for name, aset in sub_aliases.items() if sn in aset), None)
        if current_date is not None and canonical is not None:
            blocks_idx.setdefault(current_date, {})[canonical] = i
            date_col_indices.add(i)

    dates = sorted(blocks_idx)
    if not dates:
        raise ReportError("No attendance date blocks detected.")
    # Require shift and punch_status, which drive all report metrics.
    for d in dates:
        if "shift" not in blocks_idx[d] or "punch_status" not in blocks_idx[d]:
            raise ReportError(f"Attendance date {d:%d-%b-%Y} is missing shift or punch_status.")

    # Preserve all non-date input columns, then normalize each date block to the canonical 4-column order.
    nondate_indices = [i for i in range(len(raw_cols)) if i not in date_col_indices and not _parse_header_date(h1[i])]
    output_columns: list[str] = []
    identity_columns: list[str] = []
    internal_to_output: dict[str, str] = {}
    seen: dict[str, int] = {}
    for i in nondate_indices:
        label = _cell_text(h1[i]) or f"Column {i+1}"
        key = label.casefold(); seen[key] = seen.get(key,0)+1
        out = label if seen[key]==1 else f"{label}_{seen[key]}"
        output_columns.append(out); identity_columns.append(out); internal_to_output[raw_cols[i]] = out

    blocks: dict[date, dict[str, str]] = {}
    for d in dates:
        blocks[d] = {}
        for sub in ATTENDANCE_SUBHEADERS:
            out = f"{d.isoformat()}|{sub}"
            output_columns.append(out); blocks[d][sub] = out
            idx = blocks_idx[d].get(sub)
            if idx is not None:
                internal_to_output[raw_cols[idx]] = out

    # Build a normalized frame in output order.
    if pl is not None and isinstance(raw, pl.DataFrame):
        exprs=[]
        # map output -> raw source
        rev={v:k for k,v in internal_to_output.items()}
        for out in output_columns:
            src=rev.get(out)
            exprs.append(pl.col(src).cast(pl.String, strict=False).fill_null("").alias(out) if src else pl.lit("").alias(out))
        frame=raw.select(exprs)
    else:
        rev={v:k for k,v in internal_to_output.items()}
        data={out: (raw[rev[out]].astype(str) if out in rev else "") for out in output_columns}
        frame=pd.DataFrame(data).fillna("")

    output_identity_map={}
    for desired, src in identity_map.items():
        out=internal_to_output.get(src)
        if out: output_identity_map[desired]=out
    return AttendanceData(frame, output_columns, identity_columns, output_identity_map, blocks, dates)


def _validate_window(start: date, end: date, attendance: AttendanceData) -> list[date]:
    if start > end:
        raise ReportError("From Date cannot be after To Date.")
    if (start.year, start.month) != (end.year, end.month):
        raise ReportError("From Date and To Date must be within the same calendar month.")
    max_end = date.today() - timedelta(days=1)
    if end > max_end:
        raise ReportError(f"Attendance Adherence reports can run only through D-1 ({max_end:%d-%b-%Y}).")
    dates=[]; d=start
    missing=[]
    available=set(attendance.dates)
    while d<=end:
        dates.append(d)
        if d not in available: missing.append(d)
        d += timedelta(days=1)
    if missing:
        text=", ".join(x.strftime("%d-%b") for x in missing[:6])
        if len(missing)>6: text += "..."
        raise ReportError(f"Attendance Base does not contain the selected date(s): {text}")
    return dates


def _display_sites(users: UserData, attendance: AttendanceData) -> tuple[list[str], list[str]]:
    canonical={normalize_site(x): x for x in DEFAULT_DC_SITES+DEFAULT_XD_SITES}
    candidates: dict[str,str]={}
    for v in _frame_col(users.frame, users.required["facility"]):
        label=_clean_space(v)
        if label.upper().endswith(" DC") or label.upper().endswith(" XD"):
            n=normalize_site(label); candidates.setdefault(n, canonical.get(n,label))
    att_fac=attendance.identity_map["facility_name"]
    for v in _frame_col(attendance.frame, att_fac):
        label=_clean_space(v)
        if label.upper().endswith(" DC") or label.upper().endswith(" XD"):
            n=normalize_site(label); candidates.setdefault(n, canonical.get(n,label))

    def ordered(kind: str, defaults: list[str]):
        out=[]; used=set()
        for d in defaults:
            n=normalize_site(d)
            if n in candidates and n not in used:
                out.append(d); used.add(n)
        extras=sorted((label for n,label in candidates.items() if n not in used and label.upper().endswith(kind)), key=str.casefold)
        out.extend(extras)
        return out
    return ordered(" DC",DEFAULT_DC_SITES), ordered(" XD",DEFAULT_XD_SITES)


def _attendance_roles(attendance: AttendanceData) -> list[str]:
    values = _frame_col(attendance.frame, attendance.identity_map["role"])
    found: dict[str, str] = {}
    preferred = {normalize_role(r): r for r in DEFAULT_ROLE_REPORTS}
    for value in values:
        label = _clean_space(value)
        if not label or label.casefold() in {"-", "0", "na", "n/a", "none"}:
            continue
        key = normalize_role(label)
        found.setdefault(key, preferred.get(key, label))
    return sorted(found.values(), key=str.casefold)


def _build_report_specs(selected_reports: list[str] | None) -> list[ReportSpec]:
    selected = list(selected_reports or [MAIN_REPORT_TOKEN, *dedicated_master_roles()])
    specs: list[ReportSpec] = []
    seen: set[str] = set()
    for token in selected:
        if token == MAIN_REPORT_TOKEN:
            if MAIN_REPORT_TOKEN in seen:
                continue
            seen.add(MAIN_REPORT_TOKEN)
            main_roles = effective_main_roles()
            if not main_roles:
                raise ReportError(
                    "Main Attendance was selected, but Effective REF is empty. "
                    "Add at least one non-dedicated role to REF Master.xlsx or unselect Main Attendance."
                )
            specs.append(ReportSpec(
                key="main", display_name=MAIN_REPORT_LABEL, roles=tuple(main_roles),
                deployed_sheet_name=None, hc_label="Active \nHC", is_main=True,
            ))
            continue
        role = _clean_space(token)
        if not role:
            continue
        role_key = normalize_role(role)
        key = f"role:{role_key}"
        if key in seen:
            continue
        seen.add(key)
        actual_sheet = validate_master_role(role)
        if role_key == "housekeeping":
            hc_label = "Deployed \nHC"
        else:
            hc_label = "Required HC"
        specs.append(ReportSpec(
            key=key, display_name=role, roles=(role,), deployed_sheet_name=actual_sheet,
            hc_label=hc_label, is_main=False,
        ))
    if not specs:
        raise ReportError("Select at least one report/role to generate.")
    return specs


def scan_inputs(users1: str, users2: str, attendance_path: str, start: date, end: date, progress: ProgressCallback | None=None) -> ScanResult:
    _progress(progress, 10, "Reading All User Data files...")
    users=load_users(users1,users2)
    _progress(progress, 45, "Reading Attendance Base...")
    att=load_attendance(attendance_path)
    _validate_window(start,end,att)
    _progress(progress, 72, "Detecting DC, XD and roles...")
    dc,xd=_display_sites(users,att)
    roles=_attendance_roles(att)
    attendance_by_key={normalize_role(r):r for r in roles}
    master_roles=dedicated_master_roles()
    master_by_key={normalize_role(r):r for r in master_roles}
    dedicated=[attendance_by_key[k] for k in attendance_by_key if k in master_by_key]
    master_only=[master_by_key[k] for k in master_by_key if k not in attendance_by_key]
    _progress(progress,100,"Input scan complete.")
    return ScanResult(
        dc,xd,roles,dedicated,master_only,
        _frame_height(users.frame),_frame_height(att.frame),att.dates[0],att.dates[-1]
    )


def _read_csv_asset(name: str) -> list[list[str]]:
    with (ASSET_DIR/name).open("r",encoding="utf-8-sig",newline="") as f:
        return [row for row in csv.reader(f)]


def _floater_set() -> set[str]:
    rows=_read_csv_asset("floater.csv")
    return {_clean_space(r[0]).casefold() for r in rows[1:] if r and _clean_space(r[0])}


def _deployed_map(spec: ReportSpec) -> tuple[dict[str, float], str]:
    if spec.is_main or not spec.deployed_sheet_name:
        return {}, ""
    if not MASTER_HC_PATH.exists():
        raise ReportError(f"{MASTER_HC_PATH.name} is missing from the project folder.")
    wb = load_workbook(MASTER_HC_PATH, read_only=True, data_only=True)
    try:
        if spec.deployed_sheet_name not in wb.sheetnames:
            raise ReportError(f"Required master sheet not found: {spec.deployed_sheet_name}")
        ws = wb[spec.deployed_sheet_name]
        max_col = ws.max_column
        if max_col < 2:
            raise ReportError(
                f"Master sheet '{spec.deployed_sheet_name}' must have Site Name in the first column "
                "and deployed headcount in the last column."
            )
        result: dict[str, float] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            site = _clean_space(row[0] if len(row) else "")
            if not site:
                continue
            n = normalize_site(site)
            if n in result:
                continue  # preserve first-match/XLOOKUP semantics
            raw = row[max_col-1] if len(row) >= max_col else None
            try:
                result[n] = float(raw) if raw not in (None, "") else 0.0
            except Exception:
                result[n] = 0.0
        return result, xl_col_to_name(max_col-1)
    finally:
        wb.close()


def _color_hex(color) -> str | None:
    if color is None:
        return None
    try:
        if color.type == "rgb" and color.rgb:
            rgb = color.rgb[-6:]
            return f"#{rgb}"
    except Exception:
        pass
    return None


def _border_style_code(style: str | None) -> int | None:
    return {
        "thin":1, "medium":2, "dashed":3, "dotted":4, "thick":5, "double":6,
        "hair":7, "mediumDashed":8, "dashDot":9, "mediumDashDot":10,
        "dashDotDot":11, "mediumDashDotDot":12, "slantDashDot":13,
    }.get(style or "")


def _xlsxwriter_format_from_openpyxl(wb, cell, cache: dict[int, object]):
    sid = int(cell.style_id or 0)
    if sid in cache:
        return cache[sid]
    props = {}
    f = cell.font
    if f:
        if f.name: props["font_name"] = f.name
        if f.sz: props["font_size"] = float(f.sz)
        if f.bold: props["bold"] = True
        if f.italic: props["italic"] = True
        if f.underline: props["underline"] = 1
        fc = _color_hex(f.color)
        if fc: props["font_color"] = fc
    fill = cell.fill
    if fill and fill.fill_type == "solid":
        bg = _color_hex(fill.fgColor)
        if bg:
            props["bg_color"] = bg
            props["pattern"] = 1
    align = cell.alignment
    if align:
        if align.horizontal: props["align"] = align.horizontal
        if align.vertical: props["valign"] = align.vertical
        if align.wrap_text: props["text_wrap"] = True
        if align.text_rotation: props["rotation"] = int(align.text_rotation)
    if cell.number_format and cell.number_format != "General":
        props["num_format"] = cell.number_format
    border = cell.border
    if border:
        for side_name in ("left","right","top","bottom"):
            side = getattr(border, side_name)
            code = _border_style_code(side.style if side else None)
            if code:
                props[side_name] = code
                color = _color_hex(side.color if side else None)
                if color:
                    props[f"{side_name}_color"] = color
    fmt = wb.add_format(props)
    cache[sid] = fmt
    return fmt


def _write_deployed_sheet_from_master(wb, master_sheet_name: str):
    src_wb = load_workbook(MASTER_HC_PATH, read_only=False, data_only=False)
    try:
        src = src_wb[master_sheet_name]
        dst = wb.add_worksheet("Deployed HC")
        fmt_cache: dict[int, object] = {}
        merged_nonanchors: set[tuple[int,int]] = set()
        for rng in src.merged_cells.ranges:
            min_col, min_row, max_col, max_row = rng.bounds
            for rr in range(min_row-1, max_row):
                for cc in range(min_col-1, max_col):
                    if not (rr == min_row-1 and cc == min_col-1):
                        merged_nonanchors.add((rr,cc))
        for r in range(1, src.max_row+1):
            rd = src.row_dimensions[r]
            if rd.height is not None or rd.hidden:
                dst.set_row(r-1, rd.height, None, {"hidden": bool(rd.hidden)})
            for c in range(1, src.max_column+1):
                if (r-1,c-1) in merged_nonanchors:
                    continue
                cell = src.cell(r,c)
                fmt = _xlsxwriter_format_from_openpyxl(wb, cell, fmt_cache)
                value = cell.value
                if isinstance(value, datetime):
                    dst.write_datetime(r-1,c-1,value,fmt)
                elif isinstance(value, date):
                    dst.write_datetime(r-1,c-1,datetime.combine(value,datetime.min.time()),fmt)
                elif isinstance(value,str) and value.startswith("="):
                    dst.write_formula(r-1,c-1,value,fmt)
                elif value is None:
                    if cell.has_style:
                        dst.write_blank(r-1,c-1,None,fmt)
                else:
                    dst.write(r-1,c-1,value,fmt)
        from openpyxl.utils import get_column_letter
        for c in range(1, src.max_column+1):
            dim = src.column_dimensions.get(get_column_letter(c))
            if dim and (dim.width is not None or dim.hidden):
                dst.set_column(c-1,c-1,dim.width, None, {"hidden": bool(dim.hidden)})
        for rng in src.merged_cells.ranges:
            min_col, min_row, max_col, max_row = rng.bounds
            anchor = src.cell(min_row,min_col)
            fmt = _xlsxwriter_format_from_openpyxl(wb, anchor, fmt_cache)
            dst.merge_range(min_row-1,min_col-1,max_row-1,max_col-1,anchor.value,fmt)
        if src.freeze_panes:
            coord = src.freeze_panes.coordinate if hasattr(src.freeze_panes,"coordinate") else str(src.freeze_panes)
            from openpyxl.utils.cell import coordinate_to_tuple
            rr,cc = coordinate_to_tuple(coord)
            dst.freeze_panes(rr-1,cc-1)
        if src.auto_filter and src.auto_filter.ref:
            try: dst.autofilter(src.auto_filter.ref)
            except Exception: pass
        return dst
    finally:
        src_wb.close()


def _role_key(v: object) -> str:
    return normalize_role(v)


def _floater_from_user(raw_floater: str, user_id: str, floater_ids: set[str]) -> bool | str:
    s=_clean_space(raw_floater).casefold()
    if s=="yes": return True
    if s=="no": return False
    if s=="-": return "-"
    return _clean_space(user_id).casefold() in floater_ids


def _build_caches(users: UserData, att: AttendanceData, report_dates: list[date], specs: list[ReportSpec], progress: ProgressCallback | None=None):
    floater_ids=_floater_set()
    role_sets={spec.key:{_role_key(x) for x in spec.roles} for spec in specs}
    spec_by_key={spec.key:spec for spec in specs}

    uid=_frame_col(users.frame, users.required["id"])
    urole=_frame_col(users.frame, users.required["role"])
    uflo=_frame_col(users.frame, users.required["floater"]) if users.required["floater"] else [""]*_frame_height(users.frame)
    ustatus=_frame_col(users.frame, users.required["status"])
    ufac=_frame_col(users.frame, users.required["facility"])
    uonb=_frame_col(users.frame, users.required["onboarded_date"])

    user_raw_floater: dict[str,str]={}
    for i,ident in enumerate(uid):
        k=_clean_space(ident).casefold()
        if k and k not in user_raw_floater: user_raw_floater[k]=uflo[i]

    active_dates={spec.key:{} for spec in specs}
    for i in range(len(uid)):
        if _clean_space(ustatus[i]).casefold() != "onboarded - active": continue
        d=_parse_value_date(uonb[i])
        if d is None: continue
        site=normalize_site(ufac[i])
        if not (site and (_clean_space(ufac[i]).upper().endswith(" DC") or _clean_space(ufac[i]).upper().endswith(" XD"))): continue
        role=_role_key(urole[i])
        for spec in specs:
            if role not in role_sets[spec.key]: continue
            if spec.is_main:
                fl=_floater_from_user(uflo[i],uid[i],floater_ids)
                if fl is not False: continue
            else:
                if _clean_space(uid[i]).casefold() in floater_ids: continue
            active_dates[spec.key].setdefault(site,[]).append(d)
    for key in active_dates:
        for site in active_dates[key]: active_dates[key][site].sort()

    _progress(progress,18,"Building attendance caches...")
    a_uid=_frame_col(att.frame,att.identity_map["user_id"])
    a_role=_frame_col(att.frame,att.identity_map["role"])
    a_fac=_frame_col(att.frame,att.identity_map["facility_name"])
    shift_cols={d:_frame_col(att.frame,att.blocks[d]["shift"]) for d in report_dates}
    punch_cols={d:_frame_col(att.frame,att.blocks[d]["punch_status"]) for d in report_dates}
    metrics={spec.key:{d:{} for d in report_dates} for spec in specs}
    for i in range(len(a_uid)):
        ident=_clean_space(a_uid[i]); ident_n=ident.casefold()
        rawflo=user_raw_floater.get(ident_n,"")
        fl=_floater_from_user(rawflo,ident,floater_ids)
        if fl is not False: continue
        site=normalize_site(a_fac[i])
        if not site: continue
        role=_role_key(a_role[i])
        matched=[spec for spec in specs if role in role_sets[spec.key]]
        if not matched: continue
        for d in report_dates:
            shift=_clean_space(shift_cols[d][i]).upper()
            punch=_clean_space(punch_cols[d][i]).upper()
            for spec in matched:
                arr=metrics[spec.key][d].setdefault(site,[0,0,0,0,0])
                if shift in {"P1","P2","P3"}: arr[0]+=1
                if punch=="SP": arr[1]+=1
                elif punch=="MIS": arr[2]+=1
                elif punch=="HF": arr[3]+=1
                elif punch=="P": arr[4]+=1
    return floater_ids,user_raw_floater,active_dates,metrics


def _site_order(selected: list[str], defaults: list[str]) -> list[str]:
    selected_map={normalize_site(s):s for s in selected}
    out=[]; used=set()
    for d in defaults:
        n=normalize_site(d)
        if n in selected_map and n not in used:
            out.append(d); used.add(n)
    for n,label in sorted(selected_map.items(),key=lambda x:x[1].casefold()):
        if n not in used: out.append(label); used.add(n)
    return out


def _active_count(active_dates: dict[str,list[date]], site: str, d: date) -> int:
    arr=active_dates.get(normalize_site(site),[])
    return bisect_right(arr,d)


def _summary_values(spec: ReportSpec, sites: list[str], dates: list[date], active_dates, att_metrics, deployed: dict[str,float]):
    data={}
    for d in dates:
        data[d]={}
        for site in sites:
            n=normalize_site(site)
            fallback=_active_count(active_dates[spec.key],site,d)
            hc = deployed[n] if (not spec.is_main and n in deployed) else fallback
            hc_num=float(hc or 0)
            total,sp,mis,hf,p = att_metrics[spec.key][d].get(n,[0,0,0,0,0])
            vals=[hc_num,total,(total/hc_num if hc_num else 0),sp,mis,hf,p,
                  ((mis+hf+p)/total if total else 0),(mis/total if total else 0),
                  (hf/total if total else 0),(p/total if total else 0)]
            data[d][site]=vals
    return data


def _ordinal_day(d: date) -> str:
    n=d.day
    if 10<=n%100<=20: suf="th"
    else: suf={1:"st",2:"nd",3:"rd"}.get(n%10,"th")
    return f"{n}{suf} {d:%b}"


def _excel_date_format(workbook):
    return workbook.add_format({"num_format":"dd-mm-yyyy"})


def _date_like_header(name: str) -> bool:
    n=_norm_header(name)
    return ("date" in n or n in {"doj","exitdate","dateofbirth","registereddate","onboardeddate","rejectiondate"})


def _to_number_or_text(value: object):
    s=_cell_text(value)
    if s=="": return ""
    try:
        if re.fullmatch(r"-?\d+",s) and len(s)<15: return int(s)
        if re.fullmatch(r"-?\d+\.\d+",s): return float(s)
    except Exception: pass
    return s


def _summary_formula_strings(spec: ReportSpec, row_excel: int, date_cell: str, au_letters: dict[str,str], deployed_value_col: str | None):
    fac=au_letters["facility"]; status=au_letters["status"]; onb=au_letters["onboarded_date"]
    fallback=(f"COUNTIFS('All User data'!${fac}:${fac},$A{row_excel},'All User data'!$A:$A,\"Yes\","
              f"'All User data'!${status}:${status},\"Onboarded - Active\",'All User data'!$B:$B,FALSE(),"
              f"'All User data'!${onb}:${onb},\"<=\"&{date_cell})")
    if spec.is_main:
        active="="+fallback
    else:
        active=(f"=IFNA(INDEX('Deployed HC'!${deployed_value_col}:${deployed_value_col},"
                f"MATCH($A{row_excel},'Deployed HC'!$A:$A,0)),{fallback})")

    common=("INDEX('Attendance Base'!$1:$1048576,0,MATCH(\"To be considered\",'Attendance Base'!$1:$1,0)),\"Yes\","
            "INDEX('Attendance Base'!$1:$1048576,0,MATCH(\"To be considered(Values - TRUE/FALSE)\",'Attendance Base'!$1:$1,0)),FALSE(),"
            f"INDEX('Attendance Base'!$1:$1048576,0,MATCH(\"facility_name\",'Attendance Base'!$1:$1,0)),$A{row_excel}")
    total=(f"=COUNTIFS(INDEX('Attendance Base'!$1:$1048576,0,MATCH({date_cell},'Attendance Base'!$1:$1,0)),\"P1\",{common})"
           f"+COUNTIFS(INDEX('Attendance Base'!$1:$1048576,0,MATCH({date_cell},'Attendance Base'!$1:$1,0)),\"P2\",{common})"
           f"+COUNTIFS(INDEX('Attendance Base'!$1:$1048576,0,MATCH({date_cell},'Attendance Base'!$1:$1,0)),\"P3\",{common})")

    def punch(status_code: str):
        return ("=COUNTIFS(INDEX('Attendance Base'!$1:$1048576,0,MATCH("+date_cell+",'Attendance Base'!$1:$1,0)-1+"
                "MATCH(\"punch_status\",INDEX('Attendance Base'!$2:$2,1,MATCH("+date_cell+",'Attendance Base'!$1:$1,0)):"
                "INDEX('Attendance Base'!$2:$2,1,COLUMNS('Attendance Base'!$2:$2)),0)),\""+status_code+"\","+common+")")
    return active,total,punch("SP"),punch("MIS"),punch("HF"),punch("P")


def _add_formats(wb):
    border=1
    header=wb.add_format({"bold":True,"font_name":"Calibri","font_size":11,"bg_color":"#DBE9F7","border":border,"align":"center","valign":"vcenter","text_wrap":True})
    header_metric=wb.add_format({"bold":True,"font_name":"Calibri","font_size":10,"bg_color":"#DBE9F7","border":border,"align":"center","valign":"vcenter","text_wrap":True})
    site=wb.add_format({"font_name":"Calibri","font_size":11,"border":border,"align":"left","valign":"vcenter"})
    num=wb.add_format({"font_name":"Calibri","font_size":11,"border":border,"align":"center","num_format":"#,##0"})
    pct=wb.add_format({"font_name":"Calibri","font_size":11,"border":border,"align":"center","num_format":"0%"})
    total_site=wb.add_format({"bold":True,"font_name":"Calibri","font_size":11,"bg_color":"#D9D9D9","border":border,"align":"left"})
    total_num=wb.add_format({"bold":True,"font_name":"Calibri","font_size":11,"bg_color":"#D9D9D9","border":border,"align":"center","num_format":"#,##0"})
    total_pct=wb.add_format({"bold":True,"font_name":"Calibri","font_size":11,"bg_color":"#D9D9D9","border":border,"align":"center","num_format":"0%"})
    base_header=wb.add_format({"bold":False,"font_name":"Calibri","font_size":11})
    helper_header=wb.add_format({"bold":False,"font_name":"Calibri","font_size":11})
    dark_header=wb.add_format({"bold":True,"font_name":"Calibri","font_size":10,"font_color":"#FFFFFF","bg_color":"#0C343D","border":1})
    small=wb.add_format({"font_name":"Calibri","font_size":10})
    return locals()


def _write_summary_sheet(wb, name: str, spec: ReportSpec, sites: list[str], dates: list[date], values, au_letters: dict[str,str], deployed_value_col: str | None):
    ws=wb.add_worksheet(name)
    fmt=_add_formats(wb)
    ws.freeze_panes(2,1)
    ws.set_row(0,15); ws.set_row(1,54)
    ws.set_column(0,0,16.33)
    # Write rows in strict order; merged-range metadata is injected after close.
    ws.write(0,0,"Site",fmt["header"])
    headers=[spec.hc_label]+METRIC_HEADERS[1:]
    site_start=2; total_row=site_start+len(sites)

    date_header_fmt=wb.add_format({"bold":True,"font_name":"Calibri","font_size":11,"bg_color":"#DBE9F7","border":1,"align":"center","valign":"vcenter","num_format":"dd-mm-yyyy"})
    # Constant-memory worksheets must be written in strict row order. Write/merge the complete
    # date header row first, then write the complete metric row.
    for di,d in enumerate(dates):
        c=1+di*11
        hidden=di < max(0,len(dates)-2)
        for j,w in enumerate(METRIC_WIDTHS):
            ws.set_column(c+j,c+j,w,None,{"hidden":hidden})
        # XlsxWriter's constant-memory mode does not retain horizontal merged ranges.
        # Write the header cells now and inject the merges into the final XLSX package after close.
        ws.write_datetime(0,c,datetime.combine(d,datetime.min.time()),date_header_fmt)
        for j in range(1,11):
            ws.write_blank(0,c+j,None,fmt["header"])
    ws.write_blank(1,0,None,fmt["header"])
    for di,d in enumerate(dates):
        c=1+di*11
        for j,h in enumerate(headers):
            ws.write(1,c+j,h,fmt["header_metric"])

    for i,site_name in enumerate(sites):
        r=site_start+i; excel_r=r+1
        ws.write(r,0,site_name,fmt["site"])
        for di,d in enumerate(dates):
            c=1+di*11; date_cell=f"{xl_col_to_name(c)}$1"
            vals=values[d][site_name]
            active,total,sp,mis,hf,p=_summary_formula_strings(spec,excel_r,date_cell,au_letters,deployed_value_col)
            ws.write_formula(r,c,active,fmt["num"],vals[0])
            ws.write_formula(r,c+1,total,fmt["num"],vals[1])
            ws.write_formula(r,c+2,f"=IFERROR({xl_col_to_name(c+1)}{excel_r}/{xl_col_to_name(c)}{excel_r},0)",fmt["pct"],vals[2])
            ws.write_formula(r,c+3,sp,fmt["num"],vals[3])
            ws.write_formula(r,c+4,mis,fmt["num"],vals[4])
            ws.write_formula(r,c+5,hf,fmt["num"],vals[5])
            ws.write_formula(r,c+6,p,fmt["num"],vals[6])
            ws.write_formula(r,c+7,f"=IFERROR(SUM({xl_col_to_name(c+4)}{excel_r}:{xl_col_to_name(c+6)}{excel_r})/{xl_col_to_name(c+1)}{excel_r},0)",fmt["pct"],vals[7])
            ws.write_formula(r,c+8,f"=IFERROR({xl_col_to_name(c+4)}{excel_r}/{xl_col_to_name(c+1)}{excel_r},0)",fmt["pct"],vals[8])
            ws.write_formula(r,c+9,f"=IFERROR({xl_col_to_name(c+5)}{excel_r}/{xl_col_to_name(c+1)}{excel_r},0)",fmt["pct"],vals[9])
            ws.write_formula(r,c+10,f"=IFERROR({xl_col_to_name(c+6)}{excel_r}/{xl_col_to_name(c+1)}{excel_r},0)",fmt["pct"],vals[10])

    total_excel=total_row+1
    ws.write(total_row,0,"Grand Total",fmt["total_site"])
    for di,d in enumerate(dates):
        c=1+di*11
        cached=[sum(values[d][s][j] for s in sites) for j in [0,1,3,4,5,6]]
        hc,total,sp,mis,hf,p=cached
        calc=[hc,total,(total/hc if hc else 0),sp,mis,hf,p,((mis+hf+p)/total if total else 0),(mis/total if total else 0),(hf/total if total else 0),(p/total if total else 0)]
        first=3; last=2+len(sites)
        for j,idx in [(0,0),(1,1),(3,3),(4,4),(5,5),(6,6)]:
            col=xl_col_to_name(c+j)
            ws.write_formula(total_row,c+j,f"=SUM({col}{first}:{col}{last})",fmt["total_num"],calc[idx])
        ws.write_formula(total_row,c+2,f"=IFERROR({xl_col_to_name(c+1)}{total_excel}/{xl_col_to_name(c)}{total_excel},0)",fmt["total_pct"],calc[2])
        ws.write_formula(total_row,c+7,f"=IFERROR(SUM({xl_col_to_name(c+4)}{total_excel}:{xl_col_to_name(c+6)}{total_excel})/{xl_col_to_name(c+1)}{total_excel},0)",fmt["total_pct"],calc[7])
        ws.write_formula(total_row,c+8,f"=IFERROR({xl_col_to_name(c+4)}{total_excel}/{xl_col_to_name(c+1)}{total_excel},0)",fmt["total_pct"],calc[8])
        ws.write_formula(total_row,c+9,f"=IFERROR({xl_col_to_name(c+5)}{total_excel}/{xl_col_to_name(c+1)}{total_excel},0)",fmt["total_pct"],calc[9])
        ws.write_formula(total_row,c+10,f"=IFERROR({xl_col_to_name(c+6)}{total_excel}/{xl_col_to_name(c+1)}{total_excel},0)",fmt["total_pct"],calc[10])
        # Match the reference report's directional colour scales.
        for offset,reverse in [(2,False),(7,False),(8,True),(9,True),(10,False)]:
            rng=xlsxwriter.utility.xl_range(site_start,c+offset,total_row-1,c+offset)
            colors=("#63BE7B","#FFEB84","#F8696B") if reverse else ("#F8696B","#FFEB84","#63BE7B")
            ws.conditional_format(rng,{"type":"3_color_scale","min_type":"min","mid_type":"percentile","mid_value":50,"max_type":"max","min_color":colors[0],"mid_color":colors[1],"max_color":colors[2]})
    return ws


def _write_ref_sheet(wb, roles: list[str]):
    ws=wb.add_worksheet("Ref")
    ws.set_column(0,0,24)
    for i,r in enumerate(roles): ws.write(i,0,r)


def _write_deployed_sheet(wb, rows: list[list[str]]):
    ws=wb.add_worksheet("Deployed HC")
    fmt=_add_formats(wb)
    if not rows: return
    for c,h in enumerate(rows[0]): ws.write(0,c,h,fmt["dark_header"])
    for r,row in enumerate(rows[1:],1):
        for c,v in enumerate(row): ws.write(r,c,_to_number_or_text(v),fmt["small"])
    for c in range(len(rows[0])): ws.set_column(c,c,max(12,min(30,max(len(_cell_text(x[c])) if c<len(x) else 0 for x in rows[:60])+2)))
    ws.freeze_panes(1,0)


def _write_floater_sheet(wb, rows: list[list[str]]):
    ws=wb.add_worksheet("Floater")
    fmt=_add_formats(wb)
    for c,h in enumerate(rows[0]): ws.write(0,c,h,fmt["dark_header"])
    for r,row in enumerate(rows[1:],1):
        for c,v in enumerate(row): ws.write(r,c,v,fmt["small"])
    ws.set_column(0,0,18); ws.set_column(1,1,28); ws.set_column(2,2,20)
    ws.freeze_panes(1,0)


def _write_all_users(wb, users: UserData, spec: ReportSpec, floater_ids: set[str]):
    """Create only the All User sheet skeleton; rows are streamed into XLSX after close."""
    ws=wb.add_worksheet("All User data")
    fmt=_add_formats(wb)
    raw_cols=users.columns
    ws.write(0,0,"To be considered",fmt["helper_header"])
    ws.write(0,1,"To be considered(Values - TRUE/FALSE)",fmt["helper_header"])
    for j,h in enumerate(raw_cols,2): ws.write(0,j,h,fmt["base_header"])
    ws.freeze_panes(1,2)
    ws.set_column(0,0,16); ws.set_column(1,1,36)
    for j,h in enumerate(raw_cols,2): ws.set_column(j,j,min(32,max(10,len(h)+2)))
    # Register a date style in the worksheet XML so the streaming pass can reuse it.
    date_cols=[2+i for i,h in enumerate(raw_cols) if _date_like_header(h)]
    if date_cols:
        ws.write_datetime(1,date_cols[0],datetime(2000,1,1),_excel_date_format(wb))
    col_index={c:2+i for i,c in enumerate(raw_cols)}
    return {k:xl_col_to_name(col_index[v]) for k,v in users.required.items() if v is not None}


def _write_attendance(wb, att: AttendanceData, users: UserData, spec: ReportSpec, floater_ids: set[str], user_raw_floater: dict[str,str]):
    """Create only the Attendance sheet skeleton; rows are streamed into XLSX after close."""
    ws=wb.add_worksheet("Attendance Base")
    fmt=_add_formats(wb)
    ws.write(0,0,"To be considered",fmt["helper_header"])
    ws.write(0,1,"To be considered(Values - TRUE/FALSE)",fmt["helper_header"])
    physical=[]
    for h in att.identity_columns: physical.append((h,None,None))
    for d in att.dates:
        for si,sub in enumerate(ATTENDANCE_SUBHEADERS): physical.append((att.blocks[d][sub],d if si==0 else None,sub))
    date_fmt=_excel_date_format(wb)
    for j,(internal,top,sub) in enumerate(physical,2):
        if top is not None:
            ws.write_datetime(0,j,datetime.combine(top,datetime.min.time()),date_fmt)
        else:
            ws.write(0,j,internal if sub is None else "",fmt["base_header"])
        if sub is not None: ws.write(1,j,sub,fmt["base_header"])
    ws.freeze_panes(2,2)
    ws.set_column(0,0,16); ws.set_column(1,1,36)
    for j,(internal,_,sub) in enumerate(physical,2):
        ws.set_column(j,j,16 if sub is None else (10 if sub=="shift" else 13))


def _xml_text(value: object) -> str:
    from xml.sax.saxutils import escape as xml_escape
    text=_cell_text(value)
    if not text: return ""
    text=re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]","",text)
    if len(text)>32767: text=text[:32767]
    return xml_escape(text)


def _inline_xml(ref: str, value: object, style: str | None=None) -> str:
    text=_xml_text(value)
    if not text: return ""
    sattr=f' s="{style}"' if style else ""
    preserve=' xml:space="preserve"' if text[:1].isspace() or text[-1:].isspace() else ""
    return f'<c r="{ref}"{sattr} t="inlineStr"><is><t{preserve}>{text}</t></is></c>'


def _number_xml(ref: str, value: int|float, style: str|None=None) -> str:
    sattr=f' s="{style}"' if style else ""
    return f'<c r="{ref}"{sattr}><v>{value}</v></c>'


def _formula_xml(ref: str, formula: str, cached: object) -> str:
    from xml.sax.saxutils import escape as xml_escape
    f=xml_escape(formula[1:] if formula.startswith("=") else formula)
    if isinstance(cached,bool):
        return f'<c r="{ref}" t="b"><f>{f}</f><v>{1 if cached else 0}</v></c>'
    if isinstance(cached,(int,float)) and not isinstance(cached,bool):
        return f'<c r="{ref}"><f>{f}</f><v>{cached}</v></c>'
    val=_xml_text(cached)
    return f'<c r="{ref}" t="str"><f>{f}</f><v>{val}</v></c>'


def _excel_serial(d: date) -> int:
    return (d-date(1899,12,30)).days


def _split_sheetdata(xml: bytes) -> tuple[bytes,bytes]:
    a=xml.index(b"<sheetData>"); b=xml.index(b"</sheetData>")+len(b"</sheetData>")
    return xml[:a],xml[b:]


def _stream_all_user_xml(path: Path, prefix: bytes, suffix: bytes, users: UserData, spec: ReportSpec,
                         floater_ids:set[str], date_style:str|None, progress:ProgressCallback|None, base_pct:int) -> None:
    raw_cols=users.columns
    physical_count=2+len(raw_cols)
    letters=[xl_col_to_name(i) for i in range(physical_count)]
    max_col=letters[-1]
    col_idx={c:2+i for i,c in enumerate(raw_cols)}
    id_idx=col_idx[users.required["id"]]; role_idx=col_idx[users.required["role"]]
    floater_idx=col_idx[users.required["floater"]] if users.required["floater"] else None
    uid_vals=_frame_col(users.frame,users.required["id"]); role_vals=_frame_col(users.frame,users.required["role"])
    source_total=_frame_height(users.frame)
    flo_vals=_frame_col(users.frame,users.required["floater"]) if users.required["floater"] else [""]*source_total
    role_set={_role_key(x) for x in spec.roles}
    # Main Attendance intentionally keeps the complete source data. Dedicated role
    # reports are cleaned here so their embedded base sheet contains only that role.
    # This is done during streaming (rather than copying/filtering the whole frame)
    # to keep memory use low on large source files.
    total=(source_total if spec.is_main else sum(1 for v in role_vals if _role_key(v) in role_set))
    prefix=re.sub(rb'<dimension ref="[^"]*"\s*/>',f'<dimension ref="A1:{max_col}{total+1}"/>'.encode(),prefix,count=1)
    date_cols={i for i,h in enumerate(raw_cols) if _date_like_header(h)}
    with open(path,"wb",buffering=1024*1024) as f:
        f.write(prefix); f.write(b"<sheetData>")
        cells=[_inline_xml("A1","To be considered"),_inline_xml("B1","To be considered(Values - TRUE/FALSE)")]
        cells += [_inline_xml(f"{letters[2+i]}1",h) for i,h in enumerate(raw_cols)]
        f.write((f'<row r="1">{"".join(cells)}</row>').encode())
        out_row=2
        for idx,row in enumerate(_frame_iter_rows(users.frame,raw_cols)):
            if not spec.is_main and _role_key(role_vals[idx]) not in role_set:
                continue
            i=out_row; out_row+=1
            role_cell=f"{letters[role_idx]}{i}"; id_cell=f"{letters[id_idx]}{i}"
            consider=_role_key(role_vals[idx]) in role_set
            fa=f'=IF(COUNTIF(Ref!$A:$A,{role_cell})>0,"Yes","No")'
            if spec.is_main:
                if floater_idx is not None:
                    fb=(f'=IF(INDEX($1:$1048576,ROW(),MATCH("Floater",$1:$1,0))="Yes",TRUE,'
                        f'IF(INDEX($1:$1048576,ROW(),MATCH("Floater",$1:$1,0))="No",FALSE,'
                        f'IF(INDEX($1:$1048576,ROW(),MATCH("Floater",$1:$1,0))="-","-",COUNTIF(Floater!$A:$A,{id_cell})>0)))')
                else: fb=f'=COUNTIF(Floater!$A:$A,{id_cell})>0'
                cached_b=_floater_from_user(flo_vals[idx],uid_vals[idx],floater_ids)
            else:
                fb=f'=COUNTIF(Floater!$A:$A,{id_cell})>0'; cached_b=_clean_space(uid_vals[idx]).casefold() in floater_ids
            parts=[_formula_xml(f"A{i}",fa,"Yes" if consider else "No"),_formula_xml(f"B{i}",fb,cached_b)]
            for j,v in enumerate(row):
                ref=f"{letters[2+j]}{i}"
                if j in date_cols:
                    d=_parse_value_date(v)
                    if d: parts.append(_number_xml(ref,_excel_serial(d),date_style))
                    elif _cell_text(v): parts.append(_inline_xml(ref,v))
                elif _cell_text(v): parts.append(_inline_xml(ref,v))
            f.write((f'<row r="{i}">{"".join(parts)}</row>').encode())
            written=i-1
            if progress and written%15000==0:
                _progress(progress,base_pct+4+int(5*written/max(total,1)),f"{spec.display_name}: streaming All User data... {written:,}/{total:,}")
        f.write(b"</sheetData>"); f.write(suffix)


def _stream_attendance_xml(path: Path,prefix:bytes,suffix:bytes,att:AttendanceData,users:UserData,spec:ReportSpec,
                           floater_ids:set[str],user_raw_floater:dict[str,str],date_style:str|None,
                           progress:ProgressCallback|None,base_pct:int)->None:
    physical=[]
    for h in att.identity_columns: physical.append((h,None,None))
    for d in att.dates:
        for si,sub in enumerate(ATTENDANCE_SUBHEADERS): physical.append((att.blocks[d][sub],d if si==0 else None,sub))
    count=2+len(physical); letters=[xl_col_to_name(i) for i in range(count)]
    col_idx={internal:2+i for i,(internal,_,_) in enumerate(physical)}
    user_idx=col_idx[att.identity_map["user_id"]]; role_idx=col_idx[att.identity_map["role"]]
    all_id_idx=2+users.columns.index(users.required["id"]); all_id_letter=xl_col_to_name(all_id_idx)
    uid_vals=_frame_col(att.frame,att.identity_map["user_id"]); role_vals=_frame_col(att.frame,att.identity_map["role"])
    role_set={_role_key(x) for x in spec.roles}; phys_cols=[x[0] for x in physical]
    source_total=_frame_height(att.frame)
    # As with All User data, retain the full base only for Main Attendance. A
    # dedicated role workbook keeps only rows matching that report's role.
    total=(source_total if spec.is_main else sum(1 for v in role_vals if _role_key(v) in role_set))
    prefix=re.sub(rb'<dimension ref="[^"]*"\s*/>',f'<dimension ref="A1:{letters[-1]}{total+2}"/>'.encode(),prefix,count=1)
    with open(path,"wb",buffering=1024*1024) as f:
        f.write(prefix); f.write(b"<sheetData>")
        row1=[_inline_xml("A1","To be considered"),_inline_xml("B1","To be considered(Values - TRUE/FALSE)")]
        for j,(internal,top,sub) in enumerate(physical,2):
            ref=f"{letters[j]}1"
            if top is not None: row1.append(_number_xml(ref,_excel_serial(top),date_style))
            elif sub is None: row1.append(_inline_xml(ref,internal))
        f.write((f'<row r="1">{"".join(row1)}</row>').encode())
        row2=[]
        for j,(_,_,sub) in enumerate(physical,2):
            if sub is not None: row2.append(_inline_xml(f"{letters[j]}2",sub))
        f.write((f'<row r="2">{"".join(row2)}</row>').encode())
        out_row=3
        for idx,row in enumerate(_frame_iter_rows(att.frame,phys_cols)):
            if not spec.is_main and _role_key(role_vals[idx]) not in role_set:
                continue
            i=out_row; out_row+=1
            role_cell=f"{letters[role_idx]}{i}"; user_cell=f"{letters[user_idx]}{i}"
            consider=_role_key(role_vals[idx]) in role_set
            fa=f'=IF(COUNTIF(Ref!$A:$A,{role_cell})>0,"Yes","No")'
            if users.required.get("floater"):
                fb=(f'=IF(INDEX(\'All User data\'!$1:$1048576,MATCH({user_cell},\'All User data\'!${all_id_letter}:${all_id_letter},0),MATCH("Floater",\'All User data\'!$1:$1,0))="Yes",TRUE,'
                    f'IF(INDEX(\'All User data\'!$1:$1048576,MATCH({user_cell},\'All User data\'!${all_id_letter}:${all_id_letter},0),MATCH("Floater",\'All User data\'!$1:$1,0))="No",FALSE,'
                    f'IF(INDEX(\'All User data\'!$1:$1048576,MATCH({user_cell},\'All User data\'!${all_id_letter}:${all_id_letter},0),MATCH("Floater",\'All User data\'!$1:$1,0))="-","-",COUNTIF(Floater!$A:$A,{user_cell})>0)))')
            else: fb=f'=COUNTIF(Floater!$A:$A,{user_cell})>0'
            rawflo=user_raw_floater.get(_clean_space(uid_vals[idx]).casefold(),""); cached_b=_floater_from_user(rawflo,uid_vals[idx],floater_ids)
            parts=[_formula_xml(f"A{i}",fa,"Yes" if consider else "No"),_formula_xml(f"B{i}",fb,cached_b)]
            for j,v in enumerate(row,2):
                if _cell_text(v): parts.append(_inline_xml(f"{letters[j]}{i}",v))
            f.write((f'<row r="{i}">{"".join(parts)}</row>').encode())
            written=i-2
            if progress and written%8000==0:
                _progress(progress,base_pct+10+int(9*written/max(total,1)),f"{spec.display_name}: streaming Attendance Base... {written:,}/{total:,}")
        f.write(b"</sheetData>"); f.write(suffix)


def _inject_large_base_sheets(workbook_path:Path,users:UserData,att:AttendanceData,spec:ReportSpec,floater_ids:set[str],
                              user_raw_floater:dict[str,str],progress:ProgressCallback|None,base_pct:int,date_count:int)->None:
    import xml.etree.ElementTree as ET
    replacement=workbook_path.with_name(workbook_path.stem+".streamed.xlsx")
    with zipfile.ZipFile(workbook_path,"r") as zin:
        wbroot=ET.fromstring(zin.read("xl/workbook.xml")); ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        relroot=ET.fromstring(zin.read("xl/_rels/workbook.xml.rels")); relmap={e.attrib["Id"]:e.attrib["Target"] for e in relroot}
        targets={}
        summary_targets=set()
        for sh in wbroot.find("m:sheets",ns):
            if sh.attrib["name"] in {"All User data","Attendance Base","Attendance Adherence - DC","Attendance Adherence - XD"}:
                rid=sh.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]; t=relmap[rid]
                target=("xl/"+t) if not t.startswith("/") else t.lstrip("/")
                targets[sh.attrib["name"]]=target
                if sh.attrib["name"] in {"Attendance Adherence - DC","Attendance Adherence - XD"}: summary_targets.add(target)
        au_xml=zin.read(targets["All User data"]); at_xml=zin.read(targets["Attendance Base"])
        au_pre,au_suf=_split_sheetdata(au_xml); at_pre,at_suf=_split_sheetdata(at_xml)
        # Capture the registered date formats from skeleton cells.
        m=re.search(rb'<c r="[A-Z]+2"(?: s="(\d+)")?',au_xml); au_date_style=m.group(1).decode() if m and m.group(1) else None
        m=re.search(rb'<c r="[A-Z]+1" s="(\d+)"[^>]*><v>\d+</v></c>',at_xml); at_date_style=m.group(1).decode() if m else None
        with tempfile.TemporaryDirectory(prefix="aa_stream_") as td:
            td=Path(td); au_file=td/"alluser.xml"; at_file=td/"attendance.xml"
            _stream_all_user_xml(au_file,au_pre,au_suf,users,spec,floater_ids,au_date_style,progress,base_pct)
            _stream_attendance_xml(at_file,at_pre,at_suf,att,users,spec,floater_ids,user_raw_floater,at_date_style,progress,base_pct)
            with zipfile.ZipFile(replacement,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=1) as zout:
                merge_refs=["A1:A2"]+[f"{xl_col_to_name(1+i*11)}1:{xl_col_to_name(1+i*11+10)}1" for i in range(date_count)]
                for info in zin.infolist():
                    if info.filename==targets["All User data"]:
                        zout.write(au_file,info.filename)
                    elif info.filename==targets["Attendance Base"]:
                        zout.write(at_file,info.filename)
                    elif info.filename in summary_targets:
                        text=zin.read(info.filename).decode("utf-8")
                        existing=set(re.findall(r'<mergeCell ref="([A-Z]+\d+:[A-Z]+\d+)"',text))
                        add=[r for r in merge_refs if r not in existing]
                        if add:
                            additions="".join(f'<mergeCell ref="{r}"/>' for r in add)
                            m2=re.search(r'<mergeCells count="(\d+)">(.*?)</mergeCells>',text,flags=re.S)
                            if m2:
                                repl=f'<mergeCells count="{int(m2.group(1))+len(add)}">{m2.group(2)}{additions}</mergeCells>'
                                text=text[:m2.start()]+repl+text[m2.end():]
                            else:
                                block=f'<mergeCells count="{len(add)}">{additions}</mergeCells>'
                                text=text.replace("</sheetData>","</sheetData>"+block,1)
                        zout.writestr(info,text.encode("utf-8"))
                    else:
                        zout.writestr(info,zin.read(info.filename))
    os.replace(replacement,workbook_path)


def _inject_summary_date_merges(path: Path, date_count: int) -> None:
    if date_count <= 0:
        return
    import xml.etree.ElementTree as ET
    refs=["A1:A2"]
    for i in range(date_count):
        start=1+i*11
        refs.append(f"{xl_col_to_name(start)}1:{xl_col_to_name(start+10)}1")
    replacement=path.with_name(path.stem+".merged.xlsx")
    with zipfile.ZipFile(path,"r") as zin:
        wb=ET.fromstring(zin.read("xl/workbook.xml"))
        ns={"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        relroot=ET.fromstring(zin.read("xl/_rels/workbook.xml.rels")); relmap={e.attrib["Id"]:e.attrib["Target"] for e in relroot}
        targets=set()
        for sh in wb.find("m:sheets",ns):
            if sh.attrib["name"] in {"Attendance Adherence - DC","Attendance Adherence - XD"}:
                rid=sh.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                t=relmap[rid]
                targets.add(("xl/"+t) if not t.startswith("/") else t.lstrip("/"))
        with zipfile.ZipFile(replacement,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=1) as zout:
            for info in zin.infolist():
                data=zin.read(info.filename)
                if info.filename in targets:
                    text=data.decode("utf-8")
                    existing_refs=set(re.findall(r'<mergeCell ref="([A-Z]+\d+:[A-Z]+\d+)"',text))
                    add=[r for r in refs if r not in existing_refs]
                    if add:
                        m=re.search(r'<mergeCells count="(\d+)">(.*?)</mergeCells>',text,flags=re.S)
                        additions="".join(f'<mergeCell ref="{r}"/>' for r in add)
                        if m:
                            new_count=int(m.group(1))+len(add)
                            repl=f'<mergeCells count="{new_count}">{m.group(2)}{additions}</mergeCells>'
                            text=text[:m.start()]+repl+text[m.end():]
                        else:
                            block=f'<mergeCells count="{len(add)}">{additions}</mergeCells>'
                            text=text.replace("</sheetData>","</sheetData>"+block,1)
                    data=text.encode("utf-8")
                zout.writestr(info,data)
    os.replace(replacement,path)


def _safe_filename_component(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]+', "_", _clean_space(value)).strip(" .")
    return cleaned or "Role"


def _generate_one(path: Path, spec: ReportSpec, users: UserData, att: AttendanceData, dates: list[date], dc_sites: list[str], xd_sites: list[str], floater_ids, user_raw_floater, active_dates, metrics, progress: ProgressCallback | None, base_pct: int):
    if spec.is_main:
        deployed_map, dep_value_col = {}, None
    else:
        deployed_map, dep_value_col = _deployed_map(spec)
    all_sites=dc_sites+xd_sites
    vals=_summary_values(spec,all_sites,dates,active_dates,metrics,deployed_map)

    # Display-only ordering: last report date's 8 Hrs+ Adherence, low to high.
    # Grand Total is written separately and never participates in sorting.
    last_report_date = dates[-1]
    dc_output_sites = sorted(dc_sites, key=lambda site: vals[last_report_date][site][10])
    xd_output_sites = sorted(xd_sites, key=lambda site: vals[last_report_date][site][10])

    wb=xlsxwriter.Workbook(path,{"constant_memory":True,"use_future_functions":True})
    wb.set_calc_mode("auto")
    try:
        raw_col_idx={c:2+i for i,c in enumerate(users.columns)}
        au_letters={k:xl_col_to_name(raw_col_idx[v]) for k,v in users.required.items() if v is not None}
        _write_summary_sheet(wb,"Attendance Adherence - DC",spec,dc_output_sites,dates,vals,au_letters,dep_value_col)
        _write_summary_sheet(wb,"Attendance Adherence - XD",spec,xd_output_sites,dates,vals,au_letters,dep_value_col)
        _write_ref_sheet(wb,list(spec.roles))
        if not spec.is_main and spec.deployed_sheet_name:
            _write_deployed_sheet_from_master(wb,spec.deployed_sheet_name)
        _write_floater_sheet(wb,_read_csv_asset("floater.csv"))
        _progress(progress,base_pct+4,f"{spec.display_name}: writing All User data...")
        _write_all_users(wb,users,spec,floater_ids)
        _progress(progress,base_pct+10,f"{spec.display_name}: writing Attendance Base...")
        _write_attendance(wb,att,users,spec,floater_ids,user_raw_floater)
    finally:
        wb.close()
    _inject_large_base_sheets(path,users,att,spec,floater_ids,user_raw_floater,progress,base_pct,len(dates))


def generate_reports(inputs: ReportInputs, progress: ProgressCallback | None=None) -> list[Path]:
    out=Path(inputs.output_folder); out.mkdir(parents=True,exist_ok=True)
    _progress(progress,3,"Loading All User Data with Polars...")
    users=load_users(inputs.users_path_1,inputs.users_path_2)
    _progress(progress,10,"Loading Attendance Base...")
    att=load_attendance(inputs.attendance_path)
    dates=_validate_window(inputs.report_start,inputs.report_end,att)
    dc=_site_order(inputs.selected_dc,DEFAULT_DC_SITES); xd=_site_order(inputs.selected_xd,DEFAULT_XD_SITES)
    if not dc and not xd: raise ReportError("Select at least one DC or XD facility.")
    specs=_build_report_specs(inputs.selected_reports)
    _progress(progress,14,"Building role, floater and attendance caches...")
    floater_ids,user_raw_floater,active_dates,metrics=_build_caches(users,att,dates,specs,progress)
    stamp=_ordinal_day(date.today())
    outputs: list[tuple[ReportSpec,Path]]=[]
    for spec in specs:
        if spec.is_main:
            name=f"Attendance Adherence Report - {stamp}.xlsx"
        else:
            name=f"Attendance Adherence Report - {_safe_filename_component(spec.display_name)} - {stamp}.xlsx"
        outputs.append((spec,out/name))

    n=max(1,len(outputs))
    start_pct=22
    available=75
    step=max(18, available//n)
    for idx,(spec,path) in enumerate(outputs):
        base=min(90,start_pct+idx*step)
        _progress(progress,base,f"Creating {spec.display_name} report...")
        _generate_one(path,spec,users,att,dates,dc,xd,floater_ids,user_raw_floater,active_dates,metrics,progress,base)
        _progress(progress,min(98,base+step-2),f"{spec.display_name} report saved.")
    _progress(progress,100,f"{len(outputs)} report(s) generated successfully.")
    return [p for _,p in outputs]

